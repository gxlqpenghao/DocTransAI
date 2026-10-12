from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from doctransai.models.render_plan import ImagePlacement, RenderPlan, TargetRef, TextFill
from doctransai.renderers import RenderDiagnostic, RenderResult, resolve_image_placement


class XlsxRenderer:
    format_key = "xlsx"

    def render(
        self,
        plan: RenderPlan,
        *,
        bundle_dir: str | Path | None = None,
    ) -> RenderResult:
        workbook = load_workbook(plan.template.template_path)
        diagnostics: list[RenderDiagnostic] = []
        applied_text_slots: list[str] = []
        applied_image_slots: list[str] = []

        for text_fill in plan.text_fills:
            if text_fill.target.target_kind != "cell":
                diagnostics.append(_unsupported_target_diagnostic(text_fill.slot_id, text_fill.target))
                continue

            target = _resolve_target_cell(workbook, text_fill.target, slot_id=text_fill.slot_id)
            if isinstance(target, RenderDiagnostic):
                diagnostics.append(target)
                continue

            worksheet, cell_ref = target
            worksheet[cell_ref] = text_fill.value
            applied_text_slots.append(text_fill.slot_id)

        for placement in plan.image_placements:
            if placement.target.target_kind != "cell":
                diagnostics.append(
                    _unsupported_target_diagnostic(
                        placement.slot_id,
                        placement.target,
                        asset_id=placement.asset_id,
                        source_ref=placement.source_ref,
                    )
                )
                continue

            target = _resolve_target_cell(workbook, placement.target, slot_id=placement.slot_id, asset_id=placement.asset_id)
            if isinstance(target, RenderDiagnostic):
                diagnostics.append(target)
                continue

            worksheet, cell_ref = target
            target_width, target_height = _target_cell_box_pixels(worksheet, cell_ref)
            resolved = resolve_image_placement(
                placement,
                bundle_dir=bundle_dir,
                source_bundle_path=plan.source_bundle_path,
                target_width=target_width,
                target_height=target_height,
            )
            if resolved.diagnostics:
                diagnostics.extend(resolved.diagnostics)
                continue
            if resolved.asset_path is None:
                continue

            image = XLImage(str(resolved.asset_path))
            if resolved.render_size is not None:
                image.width = max(1, int(round(resolved.render_size.width)))
                image.height = max(1, int(round(resolved.render_size.height)))
            worksheet.add_image(image, cell_ref)
            applied_image_slots.append(placement.slot_id)

        output_path = Path(plan.output.output_path).expanduser()
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)

        return RenderResult(
            output_path=str(output_path),
            applied_text_slots=applied_text_slots,
            applied_image_slots=applied_image_slots,
            diagnostics=diagnostics,
        )


def _resolve_target_cell(
    workbook: Workbook,
    target: TargetRef,
    *,
    slot_id: str,
    asset_id: str | None = None,
) -> tuple[Worksheet, str] | RenderDiagnostic:
    sheet_name = target.sheet_name or workbook.sheetnames[0]
    cell_ref = target.cell_ref
    if not cell_ref:
        return RenderDiagnostic(
            severity="error",
            code="invalid_target",
            message=f"XLSX cell target for slot {slot_id} is missing cell_ref.",
            slot_id=slot_id,
            asset_id=asset_id,
        )
    if sheet_name not in workbook.sheetnames:
        return RenderDiagnostic(
            severity="error",
            code="missing_sheet",
            message=f"Worksheet '{sheet_name}' was not found for slot {slot_id}.",
            slot_id=slot_id,
            asset_id=asset_id,
        )
    return workbook[sheet_name], cell_ref


def _target_cell_box_pixels(worksheet: Worksheet, cell_ref: str) -> tuple[float | None, float | None]:
    row_index, column_index = coordinate_to_tuple(cell_ref)
    min_column = max_column = column_index
    min_row = max_row = row_index

    for merged_range in worksheet.merged_cells.ranges:
        if cell_ref in merged_range:
            min_column = merged_range.min_col
            max_column = merged_range.max_col
            min_row = merged_range.min_row
            max_row = merged_range.max_row
            break

    width = sum(_column_width_pixels(worksheet, index) for index in range(min_column, max_column + 1))
    height = sum(_row_height_pixels(worksheet, index) for index in range(min_row, max_row + 1))
    return max(1.0, width), max(1.0, height)


def _column_width_pixels(worksheet: Worksheet, column_index: int) -> float:
    column_letter = get_column_letter(column_index)
    column_dimension = worksheet.column_dimensions.get(column_letter)
    width = column_dimension.width if column_dimension is not None and column_dimension.width is not None else None
    if width is None:
        width = worksheet.sheet_format.defaultColWidth or 8.43
    return float(width) * 7.0 + 5.0


def _row_height_pixels(worksheet: Worksheet, row_index: int) -> float:
    row_dimension = worksheet.row_dimensions.get(row_index)
    height = row_dimension.height if row_dimension is not None and row_dimension.height is not None else None
    if height is None:
        height = worksheet.sheet_format.defaultRowHeight or 15.0
    return float(height) * 96.0 / 72.0


def _unsupported_target_diagnostic(
    slot_id: str,
    target: TargetRef,
    *,
    asset_id: str | None = None,
    source_ref: str | None = None,
) -> RenderDiagnostic:
    return RenderDiagnostic(
        severity="error",
        code="unsupported_target_kind",
        message=f"XLSX renderer does not support target_kind '{target.target_kind}' for slot {slot_id}.",
        slot_id=slot_id,
        asset_id=asset_id,
        source_ref=source_ref,
    )
