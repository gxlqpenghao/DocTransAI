from __future__ import annotations

from copy import deepcopy
from dataclasses import dataclass
import re
from pathlib import Path
import tempfile
import zipfile
import xml.etree.ElementTree as ET

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.shared import Emu
from docx.table import _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from doctransai.models.mapping import UnresolvedSlot
from doctransai.models.render_plan import ImagePlacement, RenderPlan, TargetRef, TextFill
from doctransai.renderers.docx_package import prune_unreferenced_docx_media_parts
from doctransai.renderers.image_placement import resolve_image_placement
from doctransai.renderers.models import RenderDiagnostic, RenderResult


_PLACEHOLDER_PATTERN = re.compile(r"\{[^{}\r\n]+\}")
_SETTINGS_TAG = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}settings"
_UPDATE_FIELDS_TAG = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}updateFields"
_PARAGRAPH_ALLOWED_DIRECT_CHILDREN = {
    "pPr",
    "r",
    "bookmarkStart",
    "bookmarkEnd",
    "proofErr",
    "commentRangeStart",
    "commentRangeEnd",
    "moveFromRangeStart",
    "moveFromRangeEnd",
    "moveToRangeStart",
    "moveToRangeEnd",
    "customXmlInsRangeStart",
    "customXmlInsRangeEnd",
    "customXmlDelRangeStart",
    "customXmlDelRangeEnd",
    "customXmlMoveFromRangeStart",
    "customXmlMoveFromRangeEnd",
    "customXmlMoveToRangeStart",
    "customXmlMoveToRangeEnd",
    "permStart",
    "permEnd",
}
_DOCX_PARAGRAPH_IMAGE_WIDTH_RATIO = 0.92
_DOCX_PARAGRAPH_IMAGE_HEIGHT_RATIO = 0.48
_DOCX_REPEAT_REGION_IMAGE_HEIGHT_RATIO = 0.34
_DOCX_TABLE_CELL_IMAGE_WIDTH_RATIO = 0.92


@dataclass(slots=True, frozen=True)
class RunStyleSnapshot:
    style_name: str | None
    font_name: str | None
    font_size_pt: float | None
    bold: bool | None
    italic: bool | None
    underline: bool | None
    all_caps: bool | None
    small_caps: bool | None
    east_asia_font: str | None
    ascii_font: str | None
    hansi_font: str | None
    cs_font: str | None


@dataclass(slots=True)
class PlaceholderRewriteResult:
    applied: bool
    reason: str = ""
    target_run: Run | None = None


class DocxRenderer:
    format_key = "docx"

    def render(
        self,
        plan: RenderPlan,
        *,
        bundle_dir: str | Path | None = None,
    ) -> RenderResult:
        document = Document(plan.template.template_path)
        diagnostics: list[RenderDiagnostic] = []
        applied_text_slots: list[str] = []
        applied_image_slots: list[str] = []

        repeated_text_fills = [fill for fill in plan.text_fills if _is_repeated_paragraph_target(fill.target)]
        repeated_image_placements = [
            placement for placement in plan.image_placements if _is_repeated_paragraph_target(placement.target)
        ]
        regular_text_fills = [fill for fill in plan.text_fills if not _is_repeated_paragraph_target(fill.target)]
        regular_image_placements = [
            placement for placement in plan.image_placements if not _is_repeated_paragraph_target(placement.target)
        ]

        repeated_applied_text_slots, repeated_applied_image_slots, repeated_diagnostics = _apply_repeated_paragraph_regions(
            document,
            text_fills=repeated_text_fills,
            image_placements=repeated_image_placements,
            bundle_dir=bundle_dir,
            source_bundle_path=plan.source_bundle_path,
        )
        applied_text_slots.extend(repeated_applied_text_slots)
        applied_image_slots.extend(repeated_applied_image_slots)
        diagnostics.extend(repeated_diagnostics)

        for text_fill in regular_text_fills:
            paragraph = _resolve_target_paragraph(document, text_fill.target)
            if paragraph is None:
                diagnostics.append(_missing_target_diagnostic(text_fill.slot_id, text_fill.target))
                continue
            if _is_implicit_blank_target(text_fill.target):
                rewrite_result = _replace_implicit_blank(paragraph, text_fill.value, target=text_fill.target)
            elif text_fill.target.target_kind == "visual_text_anchor":
                rewrite_result = _replace_visual_text_anchor(paragraph, text_fill.value, target=text_fill.target)
            elif text_fill.target.target_kind == "table_cell_blank":
                rewrite_result = _replace_table_cell_blank(paragraph, text_fill.value)
            elif text_fill.target.target_kind == "table_cell_instruction":
                rewrite_result = _replace_table_cell_instruction(paragraph, text_fill.value)
            elif text_fill.target.target_kind == "table_cell_choice":
                rewrite_result = _replace_table_cell_choice(paragraph, text_fill.value)
            else:
                rewrite_result = _replace_placeholder_or_expand_list(
                    paragraph,
                    text_fill.value,
                    placeholder_name=text_fill.target.notes,
                )
            if not rewrite_result.applied:
                diagnostics.append(_rewrite_failure_diagnostic(text_fill.slot_id, text_fill.target, rewrite_result.reason))
                continue
            applied_text_slots.append(text_fill.slot_id)

        for image_group in _group_image_placements_by_target(regular_image_placements):
            group_applied_slots, group_diagnostics = _apply_regular_image_group(
                document,
                image_group,
                bundle_dir=bundle_dir,
                source_bundle_path=plan.source_bundle_path,
            )
            applied_image_slots.extend(group_applied_slots)
            diagnostics.extend(group_diagnostics)

        _clear_unresolved_placeholders(document, plan.unresolved_slots)
        _replace_residual_filled_placeholders(document, plan.text_fills)
        _clear_underlined_x_template_markers(document)
        output_path = Path(plan.output.output_path).expanduser()
        output_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(output_path)
        try:
            _ensure_update_fields_on_open(output_path)
            prune_unreferenced_docx_media_parts(output_path)
        except Exception as exc:  # noqa: BLE001
            diagnostics.append(
                RenderDiagnostic(
                    severity="warning",
                    code="update_fields_postprocess_failed",
                    message=f"Failed to mark DOCX fields for refresh on open: {type(exc).__name__}: {exc}",
                )
            )

        return RenderResult(
            output_path=str(output_path),
            applied_text_slots=applied_text_slots,
            applied_image_slots=applied_image_slots,
            diagnostics=diagnostics,
        )


def _apply_repeated_paragraph_regions(
    document: DocxDocument,
    *,
    text_fills: list[TextFill],
    image_placements: list[ImagePlacement],
    bundle_dir: str | Path | None,
    source_bundle_path: str,
) -> tuple[list[str], list[str], list[RenderDiagnostic]]:
    if not text_fills and not image_placements:
        return [], [], []

    groups: dict[str, dict[str, object]] = {}
    for text_fill in text_fills:
        region_id = text_fill.target.repeat_region_id
        if region_id is None:
            continue
        group = groups.setdefault(region_id, {"text_fills": [], "image_placements": []})
        group["text_fills"].append(text_fill)
    for placement in image_placements:
        region_id = placement.target.repeat_region_id
        if region_id is None:
            continue
        group = groups.setdefault(region_id, {"text_fills": [], "image_placements": []})
        group["image_placements"].append(placement)

    applied_text_slots: list[str] = []
    applied_image_slots: list[str] = []
    diagnostics: list[RenderDiagnostic] = []

    sorted_groups = sorted(
        groups.items(),
        key=lambda item: min(
            target.paragraph_index
            for target in [*[_item.target for _item in item[1]["text_fills"]], *[_item.target for _item in item[1]["image_placements"]]]
            if target.paragraph_index is not None
        ),
        reverse=True,
    )
    for _, group in sorted_groups:
        group_text_fills: list[TextFill] = group["text_fills"]
        group_image_placements: list[ImagePlacement] = group["image_placements"]
        base_paragraph_indexes = sorted(
            {
                target.paragraph_index
                for target in [
                    *[item.target for item in group_text_fills],
                    *[item.target for item in group_image_placements],
                ]
                if target.paragraph_index is not None
            }
        )
        if not base_paragraph_indexes:
            continue
        if any(index < 0 or index >= len(document.paragraphs) for index in base_paragraph_indexes):
            for item in [*group_text_fills, *group_image_placements]:
                diagnostics.append(_missing_target_diagnostic(item.slot_id, item.target))
            continue

        seed_paragraphs = [document.paragraphs[index] for index in base_paragraph_indexes]
        paragraph_position_by_index = {index: position for position, index in enumerate(base_paragraph_indexes)}
        max_instance_index = max(
            target.instance_index or 0
            for target in [*[item.target for item in group_text_fills], *[item.target for item in group_image_placements]]
        )

        paragraph_instances: dict[int, list[Paragraph]] = {0: seed_paragraphs}
        _mark_region_paragraphs_keep_together(seed_paragraphs)
        last_tail = seed_paragraphs[-1]
        for instance_index in range(1, max_instance_index + 1):
            cloned_paragraphs: list[Paragraph] = []
            for seed_paragraph in seed_paragraphs:
                cloned_paragraph = _clone_paragraph_after(seed_paragraph, after_paragraph=last_tail)
                cloned_paragraphs.append(cloned_paragraph)
                last_tail = cloned_paragraph
            _mark_region_paragraphs_keep_together(cloned_paragraphs)
            paragraph_instances[instance_index] = cloned_paragraphs

        for text_fill in sorted(group_text_fills, key=lambda item: (item.target.instance_index or 0, item.target.paragraph_index or -1)):
            paragraph = _region_instance_paragraph(
                paragraph_instances,
                paragraph_position_by_index=paragraph_position_by_index,
                target=text_fill.target,
            )
            if paragraph is None:
                diagnostics.append(_missing_target_diagnostic(text_fill.slot_id, text_fill.target))
                continue
            rewrite_result = _replace_placeholder(
                paragraph,
                text_fill.value,
                placeholder_name=text_fill.target.notes,
            )
            if not rewrite_result.applied:
                diagnostics.append(_rewrite_failure_diagnostic(text_fill.slot_id, text_fill.target, rewrite_result.reason))
                continue
            applied_text_slots.append(text_fill.slot_id)

        for placement in sorted(group_image_placements, key=lambda item: (item.target.instance_index or 0, item.target.paragraph_index or -1)):
            paragraph = _region_instance_paragraph(
                paragraph_instances,
                paragraph_position_by_index=paragraph_position_by_index,
                target=placement.target,
            )
            if paragraph is None:
                diagnostics.append(_missing_target_diagnostic(placement.slot_id, placement.target))
                continue
            resolved = resolve_image_placement(
                placement,
                bundle_dir=bundle_dir,
                source_bundle_path=source_bundle_path,
                **_docx_image_resolve_kwargs(document, paragraph, placement),
            )
            if resolved.diagnostics:
                diagnostics.extend(resolved.diagnostics)
                continue
            if resolved.asset_path is None:
                diagnostics.append(
                    RenderDiagnostic(
                        severity="error",
                        code="missing_asset",
                        message=f"Image asset is missing for slot {placement.slot_id}.",
                        slot_id=placement.slot_id,
                        asset_id=placement.asset_id,
                        source_ref=placement.source_ref,
                    )
                )
                continue
            rewrite_result = _replace_placeholder(paragraph, "", placeholder_name=placement.target.notes)
            if not rewrite_result.applied or rewrite_result.target_run is None:
                diagnostics.append(_rewrite_failure_diagnostic(placement.slot_id, placement.target, rewrite_result.reason))
                continue
            picture_kwargs = _picture_dimensions(resolved.render_size or resolved.intrinsic_size)
            rewrite_result.target_run.add_picture(str(resolved.asset_path), **picture_kwargs)
            applied_image_slots.append(placement.slot_id)

    return applied_text_slots, applied_image_slots, diagnostics


def _resolve_target_paragraph(document: DocxDocument, target: TargetRef) -> Paragraph | None:
    if _is_header_footer_target(target):
        return _resolve_header_footer_target_paragraph(document, target)

    if target.paragraph_index is not None:
        if 0 <= target.paragraph_index < len(document.paragraphs):
            return document.paragraphs[target.paragraph_index]
        return None

    if target.table_index is None or target.row_index is None or target.column_index is None:
        return None
    if not 0 <= target.table_index < len(document.tables):
        return None

    table = document.tables[target.table_index]
    if not 0 <= target.row_index < len(table.rows):
        return None
    if not 0 <= target.column_index < len(table.columns):
        return None

    cell = table.cell(target.row_index, target.column_index)
    if target.cell_paragraph_index is not None:
        if 0 <= target.cell_paragraph_index < len(cell.paragraphs):
            return cell.paragraphs[target.cell_paragraph_index]
        return None
    placeholder_paragraph = _paragraph_with_placeholder(cell)
    if placeholder_paragraph is not None:
        return placeholder_paragraph
    if cell.paragraphs:
        return cell.paragraphs[0]
    return None


def _resolve_header_footer_target_paragraph(document: DocxDocument, target: TargetRef) -> Paragraph | None:
    story = _header_footer_story(document, target)
    if story is None:
        return None

    if target.table_index is None:
        if target.paragraph_index is None:
            return None
        if 0 <= target.paragraph_index < len(story.paragraphs):
            return story.paragraphs[target.paragraph_index]
        return None

    if target.row_index is None or target.column_index is None:
        return None
    if not 0 <= target.table_index < len(story.tables):
        return None

    table = story.tables[target.table_index]
    if not 0 <= target.row_index < len(table.rows):
        return None
    if not 0 <= target.column_index < len(table.columns):
        return None

    cell = table.cell(target.row_index, target.column_index)
    if target.cell_paragraph_index is not None:
        if 0 <= target.cell_paragraph_index < len(cell.paragraphs):
            return cell.paragraphs[target.cell_paragraph_index]
        return None
    placeholder_paragraph = _paragraph_with_placeholder(cell)
    if placeholder_paragraph is not None:
        return placeholder_paragraph
    if cell.paragraphs:
        return cell.paragraphs[0]
    return None


def _header_footer_story(document: DocxDocument, target: TargetRef):
    if not document.sections:
        return None
    section_index = target.page_index if target.page_index is not None else 0
    if section_index < 0 or section_index >= len(document.sections):
        return None
    section = document.sections[section_index]
    if target.target_kind.startswith("header_"):
        return section.header
    if target.target_kind.startswith("footer_"):
        return section.footer
    return None


def _is_header_footer_target(target: TargetRef) -> bool:
    return target.target_kind.startswith(("header_", "footer_"))


def _is_implicit_blank_target(target: TargetRef) -> bool:
    return target.target_kind == "implicit_blank" or target.target_kind.endswith("_implicit_blank")


def _paragraph_with_placeholder(cell: _Cell) -> Paragraph | None:
    for paragraph in cell.paragraphs:
        if _PLACEHOLDER_PATTERN.search(paragraph.text):
            return paragraph
    return None


def _is_repeated_paragraph_target(target: TargetRef) -> bool:
    return (
        target.repeat_region_id is not None
        and target.instance_index is not None
        and target.paragraph_index is not None
        and target.table_index is None
    )


def _clone_paragraph_after(source: Paragraph, *, after_paragraph: Paragraph) -> Paragraph:
    cloned_element = deepcopy(source._p)
    after_paragraph._p.addnext(cloned_element)
    return Paragraph(cloned_element, after_paragraph._parent)


def _region_instance_paragraph(
    paragraph_instances: dict[int, list[Paragraph]],
    *,
    paragraph_position_by_index: dict[int, int],
    target: TargetRef,
) -> Paragraph | None:
    if target.instance_index is None or target.paragraph_index is None:
        return None
    paragraphs = paragraph_instances.get(target.instance_index)
    if paragraphs is None:
        return None
    position = paragraph_position_by_index.get(target.paragraph_index)
    if position is None or position >= len(paragraphs):
        return None
    return paragraphs[position]


def _mark_region_paragraphs_keep_together(paragraphs: list[Paragraph]) -> None:
    if len(paragraphs) < 2:
        return
    for paragraph in paragraphs[:-1]:
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.keep_together = True
    paragraphs[-1].paragraph_format.keep_together = True


def _group_image_placements_by_target(placements: list[ImagePlacement]) -> list[list[ImagePlacement]]:
    groups: dict[tuple[object, ...], list[ImagePlacement]] = {}
    for placement in placements:
        groups.setdefault(_image_target_key(placement.target), []).append(placement)
    return list(groups.values())


def _image_target_key(target: TargetRef) -> tuple[object, ...]:
    return (
        target.format_key,
        target.target_kind,
        target.paragraph_index,
        target.cell_paragraph_index,
        target.run_index,
        target.run_end_index,
        target.table_index,
        target.row_index,
        target.column_index,
        target.sheet_name,
        target.cell_ref,
        target.page_index,
        target.form_field_name,
        tuple(target.bbox or []),
        target.notes,
        target.repeat_region_id,
        target.instance_index,
    )


def _apply_regular_image_group(
    document: DocxDocument,
    placements: list[ImagePlacement],
    *,
    bundle_dir: str | Path | None,
    source_bundle_path: str,
) -> tuple[list[str], list[RenderDiagnostic]]:
    if not placements:
        return [], []

    paragraph = _resolve_target_paragraph(document, placements[0].target)
    diagnostics: list[RenderDiagnostic] = []
    if paragraph is None:
        return [], [_missing_target_diagnostic(placement.slot_id, placement.target) for placement in placements]

    resolved_items = []
    for placement in placements:
        resolved = resolve_image_placement(
            placement,
            bundle_dir=bundle_dir,
            source_bundle_path=source_bundle_path,
            **_docx_image_resolve_kwargs(document, paragraph, placement),
        )
        if resolved.diagnostics:
            diagnostics.extend(resolved.diagnostics)
            continue
        if resolved.asset_path is None:
            diagnostics.append(
                RenderDiagnostic(
                    severity="error",
                    code="missing_asset",
                    message=f"Image asset is missing for slot {placement.slot_id}.",
                    slot_id=placement.slot_id,
                    asset_id=placement.asset_id,
                    source_ref=placement.source_ref,
                )
            )
            continue
        resolved_items.append((placement, resolved))

    if not resolved_items:
        return [], diagnostics

    first_placement = placements[0]
    if first_placement.target.target_kind == "table_cell_blank":
        rewrite_result = _prepare_table_cell_blank(paragraph)
    else:
        rewrite_result = _replace_placeholder_for_inline_object(
            paragraph,
            placeholder_name=first_placement.target.notes,
        )
    if not rewrite_result.applied or rewrite_result.target_run is None:
        diagnostics.extend(
            _rewrite_failure_diagnostic(placement.slot_id, placement.target, rewrite_result.reason)
            for placement in placements
        )
        return [], diagnostics

    target_run = rewrite_result.target_run
    run_style = _capture_run_style(target_run)
    current_run = target_run
    applied_slots: list[str] = []
    for index, (placement, resolved) in enumerate(resolved_items):
        if index > 0:
            current_run = _insert_run_after(current_run, style=run_style)
        picture_kwargs = _picture_dimensions(resolved.render_size or resolved.intrinsic_size)
        current_run.add_picture(str(resolved.asset_path), **picture_kwargs)
        if placement.caption:
            caption_run = _insert_run_after(current_run, style=run_style)
            caption_run.add_break()
            caption_run.add_text(placement.caption)
            current_run = caption_run
        applied_slots.append(placement.slot_id)

    diagnostics.extend(_clear_image_caption_placeholders(document, placements))
    return applied_slots, diagnostics


def _clear_image_caption_placeholders(
    document: DocxDocument,
    placements: list[ImagePlacement],
) -> list[RenderDiagnostic]:
    diagnostics: list[RenderDiagnostic] = []
    cleared_targets: set[tuple[object, ...]] = set()
    for placement in placements:
        if placement.caption_target is None or placement.caption_slot_id is None:
            continue
        target_key = _image_target_key(placement.caption_target)
        if target_key in cleared_targets:
            continue
        cleared_targets.add(target_key)

        paragraph = _resolve_target_paragraph(document, placement.caption_target)
        if paragraph is None:
            diagnostics.append(_missing_target_diagnostic(placement.caption_slot_id, placement.caption_target))
            continue

        while _placeholder_match(paragraph.text, placeholder_name=placement.caption_target.notes) is not None:
            rewrite_result = _replace_placeholder(paragraph, "", placeholder_name=placement.caption_target.notes)
            if not rewrite_result.applied:
                diagnostics.append(
                    _rewrite_failure_diagnostic(
                        placement.caption_slot_id,
                        placement.caption_target,
                        rewrite_result.reason,
                    )
                )
                break
    return diagnostics


def _docx_image_resolve_kwargs(
    document: DocxDocument,
    paragraph: Paragraph,
    placement: ImagePlacement,
) -> dict[str, float]:
    if placement.sizing_policy == "original":
        return {}

    if placement.target.table_index is not None and placement.target.row_index is not None and placement.target.column_index is not None:
        table = document.tables[placement.target.table_index]
        cell = table.cell(placement.target.row_index, placement.target.column_index)
        cell_width = max(1, int(round(cell.width * _DOCX_TABLE_CELL_IMAGE_WIDTH_RATIO))) if cell.width is not None else None
        target_width = _emu_to_pixels(cell_width) if cell_width is not None else None
        if placement.sizing_policy == "max_width":
            return {"max_width": target_width} if target_width is not None else {}
        return {"target_width": target_width} if target_width is not None else {}

    content_width = _paragraph_content_width_emu(document, paragraph)
    content_height = _page_content_height_emu(document)
    target_width = _emu_to_pixels(int(round(content_width * _DOCX_PARAGRAPH_IMAGE_WIDTH_RATIO))) if content_width is not None else None
    height_ratio = (
        _DOCX_REPEAT_REGION_IMAGE_HEIGHT_RATIO if placement.target.repeat_region_id is not None else _DOCX_PARAGRAPH_IMAGE_HEIGHT_RATIO
    )
    target_height = _emu_to_pixels(int(round(content_height * height_ratio))) if content_height is not None else None

    if placement.sizing_policy == "max_width":
        return {"max_width": target_width} if target_width is not None else {}

    payload: dict[str, float] = {}
    if target_width is not None:
        payload["target_width"] = target_width
    if target_height is not None:
        payload["target_height"] = target_height
    return payload


def _paragraph_content_width_emu(document: DocxDocument, paragraph: Paragraph) -> int | None:
    if not document.sections:
        return None
    section = document.sections[0]
    width = int(section.page_width) - int(section.left_margin) - int(section.right_margin)
    left_indent = paragraph.paragraph_format.left_indent
    right_indent = paragraph.paragraph_format.right_indent
    if left_indent is not None and int(left_indent) > 0:
        width -= int(left_indent)
    if right_indent is not None and int(right_indent) > 0:
        width -= int(right_indent)
    return max(width, 1)


def _page_content_height_emu(document: DocxDocument) -> int | None:
    if not document.sections:
        return None
    section = document.sections[0]
    height = int(section.page_height) - int(section.top_margin) - int(section.bottom_margin)
    return max(height, 1)


def _emu_to_pixels(value: int | None) -> float | None:
    if value is None:
        return None
    return (float(value) / 914400.0) * 96.0


def _replace_placeholder(
    paragraph: Paragraph,
    replacement: str,
    *,
    placeholder_name: str | None = None,
) -> PlaceholderRewriteResult:
    match = _placeholder_match(paragraph.text, placeholder_name=placeholder_name)
    if match is None:
        return PlaceholderRewriteResult(False, "missing_placeholder")
    if not _paragraph_markup_is_safe(paragraph):
        return PlaceholderRewriteResult(False, "unsafe_paragraph_markup")

    spans = _run_spans(paragraph)
    overlap_indices = [
        index
        for index, (run_start, run_end) in enumerate(spans)
        if run_start < match.end() and match.start() < run_end
    ]
    if not overlap_indices:
        return PlaceholderRewriteResult(False, "missing_placeholder")

    first_index = overlap_indices[0]
    last_index = overlap_indices[-1]
    target_run = paragraph.runs[first_index]
    run_style = _capture_run_style(target_run)

    for index in overlap_indices:
        run = paragraph.runs[index]
        run_start, run_end = spans[index]
        start_in_run = max(match.start() - run_start, 0)
        end_in_run = min(match.end() - run_start, len(run.text))

        prefix = run.text[:start_in_run]
        suffix = run.text[end_in_run:]

        if index == first_index and index == last_index:
            run.text = prefix + replacement + suffix
        elif index == first_index:
            run.text = prefix + replacement
        elif index == last_index:
            run.text = suffix
        else:
            run.text = ""

    _apply_run_style(target_run, run_style)
    return PlaceholderRewriteResult(True, "applied", target_run=target_run)


def _replace_visual_text_anchor(
    paragraph: Paragraph,
    replacement: str,
    *,
    target: TargetRef,
) -> PlaceholderRewriteResult:
    label = target.notes
    if not label:
        return PlaceholderRewriteResult(False, "missing_visual_anchor")
    if not _paragraph_markup_is_safe(paragraph):
        return PlaceholderRewriteResult(False, "unsafe_paragraph_markup")

    date_unit_result = _replace_visual_date_unit_anchor(paragraph, replacement, label=label)
    if date_unit_result is not None:
        return date_unit_result

    text = paragraph.text
    label_start = text.find(label)
    if label_start < 0:
        return PlaceholderRewriteResult(False, "missing_visual_anchor")

    insert_at = label_start + len(label)
    if insert_at < len(text) and text[insert_at] in {":", "："}:
        insert_at += 1
    return _insert_text_at_paragraph_offset(paragraph, insert_at, replacement)


def _replace_visual_date_unit_anchor(
    paragraph: Paragraph,
    replacement: str,
    *,
    label: str,
) -> PlaceholderRewriteResult | None:
    if _compact_visual_anchor_text(label) not in {"年月日", "年月日时"}:
        return None
    match = re.search(r"\s*年\s*月\s*日(?:\s*时)?", paragraph.text)
    if match is None:
        return None
    return _replace_text_at_paragraph_range(paragraph, match.start(), match.end(), replacement)


def _compact_visual_anchor_text(text: str) -> str:
    return "".join(str(text).split())


def _insert_text_at_paragraph_offset(
    paragraph: Paragraph,
    offset: int,
    text: str,
) -> PlaceholderRewriteResult:
    if not paragraph.runs:
        run = paragraph.add_run(text)
        return PlaceholderRewriteResult(True, "applied", target_run=run)

    spans = _run_spans(paragraph)
    for index, (run_start, run_end) in enumerate(spans):
        is_last_run = index == len(spans) - 1
        if run_start <= offset < run_end or (is_last_run and offset == run_end):
            run = paragraph.runs[index]
            style = _capture_run_style(run)
            in_run_offset = max(0, min(offset - run_start, len(run.text)))
            run.text = run.text[:in_run_offset] + text + run.text[in_run_offset:]
            _apply_run_style(run, style)
            return PlaceholderRewriteResult(True, "applied", target_run=run)
    return PlaceholderRewriteResult(False, "missing_visual_anchor")


def _replace_text_at_paragraph_range(
    paragraph: Paragraph,
    start: int,
    end: int,
    text: str,
) -> PlaceholderRewriteResult:
    spans = _run_spans(paragraph)
    overlap_indices = [
        index
        for index, (run_start, run_end) in enumerate(spans)
        if run_start < end and start < run_end
    ]
    if not overlap_indices:
        return PlaceholderRewriteResult(False, "missing_visual_anchor")

    first_index = overlap_indices[0]
    last_index = overlap_indices[-1]
    first_run = paragraph.runs[first_index]
    run_style = _capture_run_style(first_run)
    first_start, _ = spans[first_index]
    _, last_end = spans[last_index]
    last_run = paragraph.runs[last_index]
    prefix = first_run.text[: max(start - first_start, 0)]
    suffix = last_run.text[max(end - (last_end - len(last_run.text)), 0) :]
    first_run.text = prefix + text + suffix
    _apply_run_style(first_run, run_style)
    for index in overlap_indices[1:]:
        paragraph.runs[index].text = ""
    return PlaceholderRewriteResult(True, "applied", target_run=first_run)


def _replace_placeholder_for_inline_object(
    paragraph: Paragraph,
    *,
    placeholder_name: str | None = None,
) -> PlaceholderRewriteResult:
    match = _placeholder_match(paragraph.text, placeholder_name=placeholder_name)
    if match is None:
        return PlaceholderRewriteResult(False, "missing_placeholder")
    if not _paragraph_markup_is_safe(paragraph):
        return PlaceholderRewriteResult(False, "unsafe_paragraph_markup")

    spans = _run_spans(paragraph)
    overlap_indices = [
        index
        for index, (run_start, run_end) in enumerate(spans)
        if run_start < match.end() and match.start() < run_end
    ]
    if not overlap_indices:
        return PlaceholderRewriteResult(False, "missing_placeholder")

    first_index = overlap_indices[0]
    last_index = overlap_indices[-1]
    first_run = paragraph.runs[first_index]
    run_style = _capture_run_style(first_run)

    first_start, _ = spans[first_index]
    prefix_end = max(match.start() - first_start, 0)
    prefix = first_run.text[:prefix_end]

    if first_index == last_index:
        suffix_start = min(match.end() - first_start, len(first_run.text))
        suffix = first_run.text[suffix_start:]
        first_run.text = prefix
        _apply_run_style(first_run, run_style)
        object_run = _insert_run_after(first_run, style=run_style)
        if suffix:
            _insert_run_after(object_run, text=suffix, style=run_style)
        return PlaceholderRewriteResult(True, "applied", target_run=object_run)

    first_run.text = prefix
    _apply_run_style(first_run, run_style)
    object_run = _insert_run_after(first_run, style=run_style)
    for index in overlap_indices[1:last_index]:
        paragraph.runs[index].text = ""

    last_run = paragraph.runs[last_index]
    last_start, _ = spans[last_index]
    suffix_start = min(match.end() - last_start, len(last_run.text))
    last_run.text = last_run.text[suffix_start:]
    return PlaceholderRewriteResult(True, "applied", target_run=object_run)


def _insert_run_after(
    run: Run,
    *,
    text: str = "",
    style: RunStyleSnapshot | None = None,
) -> Run:
    paragraph = run._parent
    new_run = paragraph.add_run(text)
    run._r.addnext(new_run._r)
    if style is not None:
        _apply_run_style(new_run, style)
    return new_run


def _replace_placeholder_or_expand_list(
    paragraph: Paragraph,
    replacement: str,
    *,
    placeholder_name: str | None = None,
) -> PlaceholderRewriteResult:
    items = _split_list_replacement(replacement)
    if len(items) < 2 or not _paragraph_is_list_item(paragraph):
        return _replace_placeholder(paragraph, replacement, placeholder_name=placeholder_name)

    paragraphs = [paragraph]
    last_paragraph = paragraph
    for _ in items[1:]:
        cloned = _clone_paragraph_after(paragraph, after_paragraph=last_paragraph)
        paragraphs.append(cloned)
        last_paragraph = cloned

    first_result: PlaceholderRewriteResult | None = None
    for item_paragraph, item_text in zip(paragraphs, items):
        result = _replace_placeholder(item_paragraph, item_text, placeholder_name=placeholder_name)
        if not result.applied:
            return result
        if first_result is None:
            first_result = result

    return first_result or PlaceholderRewriteResult(False, "missing_placeholder")


def _split_list_replacement(replacement: str) -> list[str]:
    text = replacement.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not text:
        return []

    if "\n" in text:
        raw_items = text.split("\n")
    else:
        raw_items = re.split(r"[；;]\s*", text)

    items = [_strip_list_item_text(item) for item in raw_items]
    return [item for item in items if item]


def _strip_list_item_text(text: str) -> str:
    item = text.strip().strip("；;")
    item = re.sub(r"^\s*(?:[（(]\s*[一二三四五六七八九十\d]+\s*[）)]|[一二三四五六七八九十]+[、.]|\d+[、.)．])\s*", "", item)
    return item.strip()


def _paragraph_is_list_item(paragraph: Paragraph) -> bool:
    ppr = paragraph._p.pPr
    if ppr is not None and ppr.numPr is not None:
        return True
    style_name = paragraph.style.name.casefold() if paragraph.style is not None and paragraph.style.name else ""
    return "list" in style_name


def _replace_implicit_blank(
    paragraph: Paragraph,
    replacement: str,
    *,
    target: TargetRef,
) -> PlaceholderRewriteResult:
    if not _paragraph_markup_is_safe(paragraph):
        return PlaceholderRewriteResult(False, "unsafe_paragraph_markup")
    if target.run_index is None:
        return PlaceholderRewriteResult(False, "missing_run_index")

    start = target.run_index
    end = target.run_end_index if target.run_end_index is not None else start
    if start < 0 or end < start or end >= len(paragraph.runs):
        return PlaceholderRewriteResult(False, "missing_implicit_blank")

    target_run = paragraph.runs[start]
    run_style = _capture_run_style(target_run)
    target_run.text = replacement
    for index in range(start + 1, end + 1):
        paragraph.runs[index].text = ""
    _apply_run_style(target_run, run_style)
    return PlaceholderRewriteResult(True, "applied", target_run=target_run)


def _replace_table_cell_blank(paragraph: Paragraph, replacement: str) -> PlaceholderRewriteResult:
    rewrite_result = _prepare_table_cell_blank(paragraph)
    if not rewrite_result.applied or rewrite_result.target_run is None:
        return rewrite_result
    rewrite_result.target_run.text = replacement
    return rewrite_result


def _replace_table_cell_instruction(paragraph: Paragraph, replacement: str) -> PlaceholderRewriteResult:
    if not _paragraph_markup_is_safe(paragraph):
        return PlaceholderRewriteResult(False, "unsafe_paragraph_markup")
    target_run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    run_style = _capture_run_style(target_run)
    for run in paragraph.runs:
        run.text = ""
    lines = _instruction_lines(replacement)
    if len(lines) <= 1:
        target_run.text = replacement
        _apply_run_style(target_run, run_style)
        return PlaceholderRewriteResult(True, "applied", target_run=target_run)

    target_run.text = lines[0]
    current_run = target_run
    for line in lines[1:]:
        next_run = _insert_run_after(current_run, style=run_style)
        next_run.add_break()
        next_run.add_text(line)
        current_run = next_run
    _apply_run_style(target_run, run_style)
    return PlaceholderRewriteResult(True, "applied", target_run=target_run)


def _instruction_lines(replacement: str) -> list[str]:
    text = replacement.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not text:
        return []
    if "\n" in text:
        return [line.strip() for line in text.split("\n") if line.strip()]

    # Preserve numbered/source-authored list structure inside instruction cells.
    text = re.sub(
        r"(?<=[：:；;])\s*(?=(?:[（(]?\d+[）)]|[一二三四五六七八九十]+[、.]|\d+[、.)．]))",
        "\n",
        text,
    )
    text = re.sub(
        r"(?<=[；;])\s*(?=(?:专利|软件著作权|标准规范|应用证明|获奖证书|论文|著作|工法|标准)[:：])",
        "\n",
        text,
    )
    lines = [line.strip() for line in text.split("\n") if line.strip()]
    return lines or [replacement]


def _replace_table_cell_choice(paragraph: Paragraph, replacement: str) -> PlaceholderRewriteResult:
    if not _paragraph_markup_is_safe(paragraph):
        return PlaceholderRewriteResult(False, "unsafe_paragraph_markup")
    if "□" not in paragraph.text:
        return PlaceholderRewriteResult(False, "missing_choice_markers")

    target_run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    run_style = _capture_run_style(target_run)
    marked_text = _mark_checkbox_choice_text(paragraph.text, replacement)
    for run in paragraph.runs:
        run.text = ""
    target_run.text = marked_text
    _apply_run_style(target_run, run_style)
    return PlaceholderRewriteResult(True, "applied", target_run=target_run)


def _mark_checkbox_choice_text(template_text: str, replacement: str) -> str:
    selected_choices = _selected_choice_map(replacement)

    def repl(match: re.Match[str]) -> str:
        raw_choice = match.group(1)
        choice = _normalize_choice_text(_choice_label_without_instruction(raw_choice))
        selected = choice in selected_choices
        rendered_choice = _fill_choice_detail(raw_choice, selected_choices.get(choice)) if selected else raw_choice
        return ("☑" if selected else "□") + rendered_choice

    return re.sub(r"□([^□]+)", repl, template_text)


def _selected_choice_map(replacement: str) -> dict[str, str | None]:
    selected: dict[str, str | None] = {}
    for raw_part in re.split(r"[；;\n]+", replacement or ""):
        part = raw_part.strip()
        if not part:
            continue
        label, detail = _split_choice_detail(part)
        normalized_label = _normalize_choice_text(label)
        if normalized_label:
            selected[normalized_label] = detail
    return selected


def _split_choice_detail(text: str) -> tuple[str, str | None]:
    if "：" in text:
        label, detail = text.split("：", 1)
        return label.strip(), detail.strip() or None
    if ":" in text:
        label, detail = text.split(":", 1)
        return label.strip(), detail.strip() or None
    return text.strip(), None


def _fill_choice_detail(raw_choice: str, detail: str | None) -> str:
    if not detail:
        return raw_choice
    return re.sub(r"[_＿]{2,}", detail, raw_choice, count=1)


def _choice_label_without_instruction(text: str) -> str:
    return re.sub(r"[（(].*?[）)]", "", text).strip("：:；;，,。.、 ")


def _normalize_choice_text(text: str) -> str:
    return re.sub(r"[\s：:；;，,。.、/（）()【】\\[\\]\\-_]+", "", text or "")


def _prepare_table_cell_blank(paragraph: Paragraph) -> PlaceholderRewriteResult:
    if not _paragraph_markup_is_safe(paragraph):
        return PlaceholderRewriteResult(False, "unsafe_paragraph_markup")
    if paragraph.text.strip():
        return PlaceholderRewriteResult(False, "non_empty_table_cell_blank")

    paragraph_style = _capture_paragraph_default_run_style(paragraph)
    target_run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    run_style = _merge_run_style(_capture_run_style(target_run), paragraph_style)
    for run in paragraph.runs:
        run.text = ""
    _apply_run_style(target_run, run_style)
    return PlaceholderRewriteResult(True, "applied", target_run=target_run)


def _clear_unresolved_placeholders(document: DocxDocument, unresolved_slots: list[UnresolvedSlot]) -> None:
    placeholder_names = [slot.canonical_key for slot in unresolved_slots if slot.canonical_key]
    if not placeholder_names:
        return

    for paragraph in _iter_docx_paragraphs(document):
        for placeholder_name in placeholder_names:
            while _placeholder_match(paragraph.text, placeholder_name=placeholder_name) is not None:
                rewrite_result = _replace_placeholder(paragraph, "", placeholder_name=placeholder_name)
                if not rewrite_result.applied:
                    break


def _replace_residual_filled_placeholders(document: DocxDocument, text_fills: list[TextFill]) -> None:
    filled_values_by_name: dict[str, tuple[float, str]] = {}
    for fill in text_fills:
        if not fill.target.allow_same_name_residual_fill:
            continue
        placeholder_name = fill.target.notes
        if not placeholder_name or not fill.value:
            continue
        current = filled_values_by_name.get(placeholder_name)
        if current is None or fill.confidence > current[0]:
            filled_values_by_name[placeholder_name] = (fill.confidence, fill.value)

    if not filled_values_by_name:
        return

    for paragraph in _iter_docx_paragraphs(document):
        for placeholder_name, (_, replacement) in filled_values_by_name.items():
            while _placeholder_match(paragraph.text, placeholder_name=placeholder_name) is not None:
                rewrite_result = _replace_placeholder(paragraph, replacement, placeholder_name=placeholder_name)
                if not rewrite_result.applied:
                    break


def _clear_underlined_x_template_markers(document: DocxDocument) -> None:
    for paragraph in _iter_docx_paragraphs(document):
        if not _paragraph_markup_is_safe(paragraph):
            continue
        for run in paragraph.runs:
            if _is_underlined_x_template_marker(run):
                run.text = ""


def _is_underlined_x_template_marker(run: Run) -> bool:
    return run.underline is True and (run.text or "").strip() == "X"


def _iter_docx_paragraphs(document: DocxDocument):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in document.sections:
        for story in (section.header, section.footer):
            yield from story.paragraphs
            for table in story.tables:
                for row in table.rows:
                    for cell in row.cells:
                        yield from cell.paragraphs


def _placeholder_match(paragraph_text: str, *, placeholder_name: str | None) -> re.Match[str] | None:
    if placeholder_name is None:
        return _PLACEHOLDER_PATTERN.search(paragraph_text)

    expected_name = placeholder_name.strip()
    for match in _PLACEHOLDER_PATTERN.finditer(paragraph_text):
        candidate_name = match.group(0)[1:-1].strip()
        if candidate_name == expected_name:
            return match
    return None


def _run_spans(paragraph: Paragraph) -> list[tuple[int, int]]:
    spans: list[tuple[int, int]] = []
    cursor = 0
    for run in paragraph.runs:
        next_cursor = cursor + len(run.text)
        spans.append((cursor, next_cursor))
        cursor = next_cursor
    return spans


def _paragraph_markup_is_safe(paragraph: Paragraph) -> bool:
    for child in list(paragraph._p):
        local_name = child.tag.split("}")[-1] if "}" in child.tag else child.tag
        if local_name not in _PARAGRAPH_ALLOWED_DIRECT_CHILDREN:
            return False
    return True


def _capture_run_style(run: Run) -> RunStyleSnapshot:
    rfonts = run._element.rPr.rFonts if run._element.rPr is not None else None
    return RunStyleSnapshot(
        style_name=run.style.name if run.style is not None else None,
        font_name=run.font.name,
        font_size_pt=run.font.size.pt if run.font.size is not None else None,
        bold=run.bold,
        italic=run.italic,
        underline=run.underline,
        all_caps=run.font.all_caps,
        small_caps=run.font.small_caps,
        east_asia_font=rfonts.get(qn("w:eastAsia")) if rfonts is not None else None,
        ascii_font=rfonts.get(qn("w:ascii")) if rfonts is not None else None,
        hansi_font=rfonts.get(qn("w:hAnsi")) if rfonts is not None else None,
        cs_font=rfonts.get(qn("w:cs")) if rfonts is not None else None,
    )


def _capture_paragraph_default_run_style(paragraph: Paragraph) -> RunStyleSnapshot:
    ppr = paragraph._p.pPr
    rpr = ppr.find(qn("w:rPr")) if ppr is not None else None
    rfonts = rpr.find(qn("w:rFonts")) if rpr is not None else None
    return RunStyleSnapshot(
        style_name=None,
        font_name=None,
        font_size_pt=_font_size_pt_from_rpr(rpr),
        bold=_toggle_value_from_rpr(rpr, "w:b"),
        italic=_toggle_value_from_rpr(rpr, "w:i"),
        underline=_underline_value_from_rpr(rpr),
        all_caps=_toggle_value_from_rpr(rpr, "w:caps"),
        small_caps=_toggle_value_from_rpr(rpr, "w:smallCaps"),
        east_asia_font=rfonts.get(qn("w:eastAsia")) if rfonts is not None else None,
        ascii_font=rfonts.get(qn("w:ascii")) if rfonts is not None else None,
        hansi_font=rfonts.get(qn("w:hAnsi")) if rfonts is not None else None,
        cs_font=rfonts.get(qn("w:cs")) if rfonts is not None else None,
    )


def _merge_run_style(primary: RunStyleSnapshot, fallback: RunStyleSnapshot) -> RunStyleSnapshot:
    return RunStyleSnapshot(
        style_name=primary.style_name or fallback.style_name,
        font_name=primary.font_name or fallback.font_name,
        font_size_pt=primary.font_size_pt if primary.font_size_pt is not None else fallback.font_size_pt,
        bold=primary.bold if primary.bold is not None else fallback.bold,
        italic=primary.italic if primary.italic is not None else fallback.italic,
        underline=primary.underline if primary.underline is not None else fallback.underline,
        all_caps=primary.all_caps if primary.all_caps is not None else fallback.all_caps,
        small_caps=primary.small_caps if primary.small_caps is not None else fallback.small_caps,
        east_asia_font=primary.east_asia_font or fallback.east_asia_font,
        ascii_font=primary.ascii_font or fallback.ascii_font,
        hansi_font=primary.hansi_font or fallback.hansi_font,
        cs_font=primary.cs_font or fallback.cs_font,
    )


def _font_size_pt_from_rpr(rpr) -> float | None:
    if rpr is None:
        return None
    size = rpr.find(qn("w:sz"))
    if size is None:
        return None
    value = size.get(qn("w:val"))
    if value is None:
        return None
    try:
        return float(value) / 2.0
    except ValueError:
        return None


def _toggle_value_from_rpr(rpr, tag: str) -> bool | None:
    if rpr is None:
        return None
    element = rpr.find(qn(tag))
    if element is None:
        return None
    value = element.get(qn("w:val"))
    if value is not None and value.lower() in {"0", "false", "off"}:
        return False
    return True


def _underline_value_from_rpr(rpr) -> bool | None:
    if rpr is None:
        return None
    element = rpr.find(qn("w:u"))
    if element is None:
        return None
    value = element.get(qn("w:val"))
    if value is not None and value.lower() in {"none", "0", "false", "off"}:
        return False
    return True


def _apply_run_style(run: Run, style: RunStyleSnapshot) -> None:
    if style.style_name:
        try:
            run.style = style.style_name
        except Exception:  # noqa: BLE001
            pass
    run.font.name = style.font_name
    if style.font_size_pt is not None:
        from docx.shared import Pt

        run.font.size = Pt(style.font_size_pt)
    run.bold = style.bold
    run.italic = style.italic
    run.underline = style.underline
    run.font.all_caps = style.all_caps
    run.font.small_caps = style.small_caps

    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        from docx.oxml import OxmlElement

        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    _set_rfont_if_value(rfonts, "w:eastAsia", style.east_asia_font)
    _set_rfont_if_value(rfonts, "w:ascii", style.ascii_font)
    _set_rfont_if_value(rfonts, "w:hAnsi", style.hansi_font)
    _set_rfont_if_value(rfonts, "w:cs", style.cs_font)


def _set_rfont_if_value(rfonts, attr_name: str, value: str | None) -> None:
    if value:
        rfonts.set(qn(attr_name), value)


def _picture_dimensions(size) -> dict[str, Emu]:
    if size is None:
        return {}
    return {
        "width": _emu_from_pixels(size.width),
        "height": _emu_from_pixels(size.height),
    }


def _emu_from_pixels(value: float) -> Emu:
    pixels = max(1.0, float(value))
    return Emu(int(round((pixels / 96.0) * 914400)))


def _ensure_update_fields_on_open(docx_path: Path) -> None:
    with zipfile.ZipFile(docx_path, "r") as archive:
        names = archive.namelist()
        settings_bytes = archive.read("word/settings.xml") if "word/settings.xml" in names else None
        contents = {name: archive.read(name) for name in names if name != "word/settings.xml"}

    if settings_bytes is None:
        settings_root = ET.Element(_SETTINGS_TAG)
    else:
        settings_root = ET.fromstring(settings_bytes)

    update_fields = settings_root.find(_UPDATE_FIELDS_TAG)
    if update_fields is None:
        update_fields = ET.SubElement(settings_root, _UPDATE_FIELDS_TAG)
    update_fields.set(qn("w:val"), "true")

    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx", dir=docx_path.parent) as handle:
        temp_path = Path(handle.name)

    try:
        with zipfile.ZipFile(temp_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for name, payload in contents.items():
                archive.writestr(name, payload)
            archive.writestr("word/settings.xml", ET.tostring(settings_root, encoding="utf-8", xml_declaration=True))
        temp_path.replace(docx_path)
    finally:
        if temp_path.exists():
            temp_path.unlink(missing_ok=True)


def _missing_target_diagnostic(slot_id: str, target: TargetRef) -> RenderDiagnostic:
    return RenderDiagnostic(
        severity="error",
        code="missing_target",
        message=f"DOCX target could not be resolved for slot {slot_id}: {_target_label(target)}.",
        slot_id=slot_id,
    )


def _missing_placeholder_diagnostic(slot_id: str, target: TargetRef) -> RenderDiagnostic:
    return RenderDiagnostic(
        severity="error",
        code="missing_placeholder",
        message=f"No DOCX placeholder was found at target {_target_label(target)} for slot {slot_id}.",
        slot_id=slot_id,
    )


def _rewrite_failure_diagnostic(slot_id: str, target: TargetRef, reason: str) -> RenderDiagnostic:
    if reason == "unsafe_paragraph_markup":
        return RenderDiagnostic(
            severity="error",
            code="unsafe_paragraph_markup",
            message=f"DOCX paragraph contains unsupported direct child markup at {_target_label(target)}; rewrite skipped for slot {slot_id}.",
            slot_id=slot_id,
        )
    return _missing_placeholder_diagnostic(slot_id, target)


def _target_label(target: TargetRef) -> str:
    if _is_header_footer_target(target):
        story = "header" if target.target_kind.startswith("header_") else "footer"
        section_index = target.page_index if target.page_index is not None else 0
        if target.table_index is not None and target.row_index is not None and target.column_index is not None:
            return f"{story}[{section_index}].table[{target.table_index}].cell[{target.row_index},{target.column_index}]"
        if target.paragraph_index is not None:
            return f"{story}[{section_index}].paragraph[{target.paragraph_index}]"
        return target.target_kind
    if target.paragraph_index is not None:
        return f"paragraph[{target.paragraph_index}]"
    if target.table_index is not None and target.row_index is not None and target.column_index is not None:
        return f"table[{target.table_index}].cell[{target.row_index},{target.column_index}]"
    return target.target_kind
