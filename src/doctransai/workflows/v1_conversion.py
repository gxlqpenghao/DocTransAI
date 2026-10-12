from __future__ import annotations

import json
import re
from dataclasses import replace
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook

from doctransai.data_layer import read_parsed_document_bundle
from doctransai.mapping.agent_resolver import AgentDecision, _source_overview
from doctransai.mapping.cluster_planner import build_source_cluster_plan
from doctransai.mapping.full_agent_context import FULL_AGENT_CONTEXT_SCHEMA_VERSION
from doctransai.mapping.orchestrator import build_render_plan_with_agent, build_render_plan_with_decisions
from doctransai.mapping.planner import build_render_plan
from doctransai.models.render_plan import PlanDiagnostic, RenderPlan, render_plan_from_dict
from doctransai.renderers.docx_package import inspect_docx_media_usage
from doctransai.renderers.pipeline import renderer_for_format
from doctransai.reports import build_conversion_report, build_full_agent_context_summary, build_provenance_report
from doctransai.source_model import AdaptedSourceModel, adapt_source_corpus
from doctransai.templates import analyze_template, materialize_target_template
from doctransai.templates import normalize_sample_template as normalize_sample_template_file
from doctransai.validation import OutputArtifactSnapshot, OutputCheckResult, check_output, evaluate_bundle_gate
from doctransai.workflows.performance import WorkflowPerformanceMonitor


@dataclass(slots=True)
class ConversionWorkflowResult:
    render_plan: RenderPlan
    render_result: Any
    output_check_result: OutputCheckResult
    conversion_report: dict[str, Any]
    provenance_report: dict[str, Any]
    performance_report: dict[str, Any]
    bundle_gate_result: Any


def run_v1_conversion(
    *,
    bundle_path: str | Path,
    template_path: str | Path,
    output_path: str | Path,
    source_file_path: str | Path | None = None,
    primary_record_id: str | None = None,
    source_instance_id: str | None = None,
    minimum_severity: str = "error",
    on_trigger: str = "warn",
    use_mapping_agent: bool = False,
    decision_client=None,
    decision_payload_path: str | Path | None = None,
    normalize_sample_template: bool = False,
) -> ConversionWorkflowResult:
    resolved_bundle_path = Path(bundle_path).expanduser().resolve()
    resolved_template_path = Path(template_path).expanduser().resolve()
    resolved_output_path = Path(output_path).expanduser().resolve()
    monitor = WorkflowPerformanceMonitor()

    with monitor.step("materialize_template", template_path=str(resolved_template_path), output_path=str(resolved_output_path)) as metadata:
        template_materialization = materialize_target_template(
            resolved_template_path,
            resolved_output_path,
        )
        resolved_template_path = template_materialization.analysis_template_path
        resolved_output_path = template_materialization.output_path
        metadata.update(template_materialization.metadata)

    with monitor.step(
        "normalize_sample_template",
        enabled=normalize_sample_template,
        template_path=str(resolved_template_path),
    ) as metadata:
        template_normalization = normalize_sample_template_file(
            resolved_template_path,
            resolved_output_path,
            enabled=normalize_sample_template,
        )
        resolved_template_path = template_normalization.analysis_template_path
        metadata.update(template_normalization.metadata)

    with monitor.step("read_parsed_document_bundle", bundle_path=str(resolved_bundle_path)) as metadata:
        corpus = read_parsed_document_bundle(resolved_bundle_path)
        metadata["document_count"] = _safe_len(getattr(corpus, "documents", None))
        metadata["parse_error_count"] = _safe_len(getattr(corpus, "parse_errors", None))

    with monitor.step("evaluate_bundle_gate", minimum_severity=minimum_severity, on_trigger=on_trigger) as metadata:
        gate_result = evaluate_bundle_gate(corpus, minimum_severity=minimum_severity, on_trigger=on_trigger)
        metadata["blocked"] = gate_result.blocked
        metadata["issue_count"] = _safe_len(getattr(gate_result, "issues", None))
    if gate_result.blocked:
        raise RuntimeError("Bundle gate blocked conversion due to diagnostics.")

    with monitor.step("adapt_source_corpus") as metadata:
        source_model = adapt_source_corpus(corpus)
        _add_source_model_counts(metadata, source_model)
    if source_file_path is not None:
        with monitor.step("filter_source_model", source_file_path=str(source_file_path)) as metadata:
            source_model = _filter_source_model_by_source_path(source_model, source_file_path)
            _add_source_model_counts(metadata, source_model)
    with monitor.step("analyze_template", template_path=str(resolved_template_path)) as metadata:
        template_model = analyze_template(resolved_template_path)
        metadata["format_key"] = template_model.format_key
        metadata["slot_count"] = len(template_model.slots)
        metadata["repeat_region_count"] = len(template_model.repeat_regions)
    with monitor.step("snapshot_template", format_key=template_model.format_key):
        template_snapshot = _snapshot_for_path(resolved_template_path, format_key=template_model.format_key)
    with monitor.step("build_render_plan", use_mapping_agent=use_mapping_agent or decision_client is not None, decision_payload=decision_payload_path is not None) as metadata:
        if decision_payload_path is not None:
            decisions = _decisions_from_payload(decision_payload_path)
            metadata["decision_count"] = len(decisions)
            plan_payload = build_render_plan_with_decisions(
                template_model=template_model,
                source_model=source_model,
                source_bundle_path=str(resolved_bundle_path),
                output_path=str(resolved_output_path),
                primary_record_id=primary_record_id,
                source_instance_id=source_instance_id,
                decisions=decisions,
            )
        elif use_mapping_agent or decision_client is not None:
            plan_payload = build_render_plan_with_agent(
                template_model=template_model,
                source_model=source_model,
                source_bundle_path=str(resolved_bundle_path),
                output_path=str(resolved_output_path),
                primary_record_id=primary_record_id,
                source_instance_id=source_instance_id,
                decision_client=decision_client,
            )
        else:
            plan_payload = build_render_plan(
                template_model=template_model,
                source_model=source_model,
                source_bundle_path=str(resolved_bundle_path),
                output_path=str(resolved_output_path),
                primary_record_id=primary_record_id,
                source_instance_id=source_instance_id,
            )
        metadata["text_fill_count"] = _safe_len(plan_payload.get("text_fills"))
        metadata["image_placement_count"] = _safe_len(plan_payload.get("image_placements"))
        metadata["unresolved_slot_count"] = _safe_len(plan_payload.get("unresolved_slots"))
    with monitor.step("parse_render_plan"):
        render_plan = render_plan_from_dict(plan_payload)
        render_plan.diagnostics.extend(_materialization_plan_diagnostics(template_materialization.diagnostics))
        render_plan.diagnostics.extend(_materialization_plan_diagnostics(template_normalization.diagnostics))
    with monitor.step("build_agent_context_summary") as metadata:
        full_agent_context_summary = _build_agent_context_report_summary(
            source_model=source_model,
            template_model=template_model,
            render_plan=render_plan,
            source_instance_id=source_instance_id,
        )
        metadata.update(full_agent_context_summary)

    with monitor.step("render_output", output_format=render_plan.output.output_format) as metadata:
        renderer = renderer_for_format(render_plan.output.output_format)
        render_result = renderer.render(render_plan, bundle_dir=resolved_bundle_path.parent)
        metadata["successful"] = render_result.is_successful
        metadata["applied_text_slot_count"] = len(render_result.applied_text_slots)
        metadata["applied_image_slot_count"] = len(render_result.applied_image_slots)
    with monitor.step("snapshot_output", output_format=render_plan.output.output_format):
        snapshot = _snapshot_for_output(render_plan)
        snapshot.applied_text_slots = set(render_result.applied_text_slots)
        snapshot.applied_image_slots = set(render_result.applied_image_slots)
    with monitor.step("check_output") as metadata:
        output_check_result = check_output(
            render_plan,
            [snapshot],
            asset_base_dir=resolved_bundle_path.parent,
            template_artifact=template_snapshot,
        )
        metadata["issue_count"] = len(output_check_result.issues)

    with monitor.step("build_reports"):
        conversion_report, conversion_markdown = build_conversion_report(
            render_plan,
            render_result,
            output_check_result,
            full_agent_context_summary=full_agent_context_summary,
        )
        provenance_report, provenance_markdown = build_provenance_report(render_plan, render_result, output_check_result)
    performance_summary = {
        "source_record_count": len(source_model.records),
        "source_image_asset_count": len(source_model.image_assets),
        "template_slot_count": len(template_model.slots),
        "text_fill_count": len(render_plan.text_fills),
        "image_placement_count": len(render_plan.image_placements),
        "unresolved_slot_count": len(render_plan.unresolved_slots),
        "output_issue_count": len(output_check_result.issues),
        "template_materialization": dict(template_materialization.metadata),
        "template_normalization": dict(template_normalization.metadata),
        "full_agent_context_summary": dict(full_agent_context_summary),
    }
    performance_report = monitor.report(summary=performance_summary)
    with monitor.step("write_reports"):
        _write_reports(
            resolved_output_path,
            conversion_report=conversion_report,
            conversion_markdown=conversion_markdown,
            provenance_report=provenance_report,
            provenance_markdown=provenance_markdown,
            performance_report=performance_report,
        )
    performance_report = monitor.report(summary=performance_summary)
    _write_performance_report(resolved_output_path, performance_report)

    return ConversionWorkflowResult(
        render_plan=render_plan,
        render_result=render_result,
        output_check_result=output_check_result,
        conversion_report=conversion_report,
        provenance_report=provenance_report,
        performance_report=performance_report,
        bundle_gate_result=gate_result,
    )


def _filter_source_model_by_source_path(
    source_model: AdaptedSourceModel,
    source_file_path: str | Path,
) -> AdaptedSourceModel:
    resolved_source_path = str(Path(source_file_path).expanduser().resolve())
    documents = [document for document in source_model.documents if document.source_path == resolved_source_path]
    if not documents:
        raise ValueError(f"Source file path not found in bundle: {resolved_source_path}")

    document_ids = {document.document_id for document in documents}

    def _record_matches(record) -> bool:
        return record.document_id in document_ids and record.source_path == resolved_source_path

    image_assets = [
        asset
        for asset in source_model.image_assets
        if asset.raw.get("source_path") == resolved_source_path
    ]
    diagnostics = [
        diagnostic
        for diagnostic in source_model.diagnostics
        if diagnostic.get("source_path") == resolved_source_path
        or diagnostic.get("document_id") in document_ids
    ]

    return replace(
        source_model,
        documents=documents,
        records=[record for record in source_model.records if _record_matches(record)],
        image_assets=image_assets,
        diagnostics=diagnostics,
        evidence_records=[record for record in source_model.evidence_records if _record_matches(record)],
        table_records=[record for record in source_model.table_records if _record_matches(record)],
        cell_range_records=[record for record in source_model.cell_range_records if _record_matches(record)],
        table_row_records=[record for record in source_model.table_row_records if _record_matches(record)],
        table_cell_records=[record for record in source_model.table_cell_records if _record_matches(record)],
        xlsx_row_records=[record for record in source_model.xlsx_row_records if _record_matches(record)],
        sheet_records=[record for record in source_model.sheet_records if _record_matches(record)],
    )


def _build_agent_context_report_summary(
    *,
    source_model: AdaptedSourceModel,
    template_model,
    render_plan: RenderPlan,
    source_instance_id: str | None,
) -> dict[str, Any]:
    cluster_plan = build_source_cluster_plan(source_model=source_model, template_model=template_model)
    repeat_region_cluster_hints = [
        {
            "region_id": region.region_id,
            "candidate_cluster_ids": list(region.candidate_cluster_ids),
        }
        for region in cluster_plan.repeat_region_plans
        if region.candidate_cluster_ids
    ]
    summary_context = {
        "schema_version": FULL_AGENT_CONTEXT_SCHEMA_VERSION,
        "full_source_context": {
            "record_count": len(source_model.records),
            "image_asset_count": len(source_model.image_assets),
            "diagnostic_count": len(source_model.diagnostics),
        },
        "task_context": {
            "primary_record_id": render_plan.record_plans[0].record_id if render_plan.record_plans else None,
            "source_instance_id": source_instance_id,
            "repeat_region_cluster_hints": repeat_region_cluster_hints,
        },
    }
    return build_full_agent_context_summary(
        summary_context,
        source_overview=_report_source_overview(source_model=source_model, render_plan=render_plan),
        agent_payload_summary=_agent_payload_summary_from_plan(render_plan),
    )


def _report_source_overview(*, source_model: AdaptedSourceModel, render_plan: RenderPlan) -> dict[str, Any]:
    if not source_model.records:
        return {}
    record_id = render_plan.record_plans[0].record_id if render_plan.record_plans else None
    primary_record = next(
        (record for record in source_model.records if record.record_id == record_id),
        source_model.records[0],
    )
    return _source_overview(source_model, primary_record)


def _agent_payload_summary_from_plan(render_plan: RenderPlan) -> dict[str, Any]:
    for diagnostic in render_plan.diagnostics:
        if diagnostic.code not in {"agent_input_payload_preflight", "agent_input_payload_too_large"}:
            continue
        estimated = _message_int(diagnostic.message, "estimated_request_bytes")
        limit = _message_int(diagnostic.message, "max_request_body_bytes")
        estimated_chars = _message_int(diagnostic.message, "estimated_input_chars")
        char_limit = _message_int(diagnostic.message, "max_input_chars")
        request_body_exceeded = _message_bool(diagnostic.message, "request_body_exceeded") or False
        input_length_exceeded = _message_bool(diagnostic.message, "input_length_exceeded") or False
        within_limit = _message_bool(diagnostic.message, "within_limit")
        preflight_blocked = _message_bool(diagnostic.message, "preflight_blocked") or False
        return {
            "estimated_request_bytes": estimated,
            "max_request_body_bytes": limit,
            "estimated_input_chars": estimated_chars,
            "max_input_chars": char_limit,
            "request_body_exceeded": request_body_exceeded,
            "input_length_exceeded": input_length_exceeded,
            "within_limit": within_limit if within_limit is not None else diagnostic.code != "agent_input_payload_too_large",
            "preflight_blocked": preflight_blocked,
            "agent_input_packaging": {
                "mode": _message_string(diagnostic.message, "packaging_mode")
                or ("preflight_blocked" if preflight_blocked else "full"),
                "full_estimated_bytes": _message_optional_int(diagnostic.message, "full_estimated_bytes"),
                "compact_estimated_bytes": _message_optional_int(diagnostic.message, "compact_estimated_bytes"),
                "selected_estimated_bytes": _message_optional_int(diagnostic.message, "selected_estimated_bytes") or estimated,
                "limit_bytes": _message_optional_int(diagnostic.message, "limit_bytes") or limit,
                "full_estimated_chars": _message_optional_int(diagnostic.message, "full_estimated_chars"),
                "compact_estimated_chars": _message_optional_int(diagnostic.message, "compact_estimated_chars"),
                "selected_estimated_chars": _message_optional_int(diagnostic.message, "selected_estimated_chars") or estimated_chars,
                "char_limit": _message_optional_int(diagnostic.message, "char_limit") or char_limit,
                "request_body_exceeded": request_body_exceeded,
                "input_length_exceeded": input_length_exceeded,
                "slot_packets_compacted": _message_bool(diagnostic.message, "slot_packets_compacted") or False,
                "full_context_compacted": _message_bool(diagnostic.message, "full_context_compacted") or False,
                "legacy_fields_removed": _message_list(diagnostic.message, "legacy_fields_removed"),
                "phase": _message_string(diagnostic.message, "phase"),
                "within_limit": within_limit if within_limit is not None else diagnostic.code != "agent_input_payload_too_large",
                "preflight_blocked": preflight_blocked,
                "diagnostic_code": diagnostic.code,
            },
        }
    return {
        "estimated_request_bytes": None,
        "max_request_body_bytes": None,
        "estimated_input_chars": None,
        "max_input_chars": None,
        "request_body_exceeded": False,
        "input_length_exceeded": False,
        "within_limit": None,
        "preflight_blocked": False,
        "agent_input_packaging": {
            "mode": "not_measured",
            "full_estimated_bytes": None,
            "compact_estimated_bytes": None,
            "selected_estimated_bytes": None,
            "limit_bytes": None,
            "full_estimated_chars": None,
            "compact_estimated_chars": None,
            "selected_estimated_chars": None,
            "char_limit": None,
            "request_body_exceeded": False,
            "input_length_exceeded": False,
            "slot_packets_compacted": False,
            "full_context_compacted": False,
            "legacy_fields_removed": [],
            "phase": None,
            "within_limit": None,
            "preflight_blocked": False,
            "diagnostic_code": None,
        },
    }


def _message_int(message: str, key: str) -> int | None:
    match = re.search(rf"{re.escape(key)}=(\d+)", message)
    return int(match.group(1)) if match else None


def _message_optional_int(message: str, key: str) -> int | None:
    match = re.search(rf"{re.escape(key)}=(\d+|None|null)", message)
    if match is None:
        return None
    raw = match.group(1)
    return int(raw) if raw.isdigit() else None


def _message_bool(message: str, key: str) -> bool | None:
    match = re.search(rf"{re.escape(key)}=(true|false)", message)
    if match is None:
        return None
    return match.group(1) == "true"


def _message_string(message: str, key: str) -> str | None:
    match = re.search(rf"{re.escape(key)}=([^;]+)", message)
    if match is None:
        return None
    value = match.group(1).strip()
    return value if value and value not in {"None", "null"} else None


def _message_list(message: str, key: str) -> list[str]:
    value = _message_string(message, key)
    if value is None:
        return []
    return [item for item in (part.strip() for part in value.split(",")) if item]


def _snapshot_for_output(plan: RenderPlan) -> OutputArtifactSnapshot:
    output_path = Path(plan.output.output_path).expanduser().resolve()
    format_key = plan.output.output_format
    return _snapshot_for_path(output_path, format_key=format_key)


def _snapshot_for_path(path: Path, *, format_key: str) -> OutputArtifactSnapshot:
    if format_key == "docx":
        return _docx_snapshot(path)
    if format_key == "xlsx":
        return _xlsx_snapshot(path)
    if format_key == "pdf":
        return _pdf_snapshot(path)
    return OutputArtifactSnapshot(name=path.name, metadata={"format_key": format_key})


def _docx_snapshot(path: Path) -> OutputArtifactSnapshot:
    document = Document(path)
    media_usage = inspect_docx_media_usage(path)
    text_fragments = [paragraph.text for paragraph in document.paragraphs]
    target_styles: dict[str, object] = {}
    image_targets: list[str] = []
    for index, paragraph in enumerate(document.paragraphs):
        target_styles[f"paragraph[{index}]"] = paragraph.style.name if paragraph.style is not None else None
        if _paragraph_has_image(paragraph):
            image_targets.append(f"paragraph[{index}]")
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            for column_index, cell in enumerate(row.cells):
                key = f"table[{table_index}].cell[{row_index},{column_index}]"
                text_fragments.extend(paragraph.text for paragraph in cell.paragraphs)
                target_styles[key] = cell.paragraphs[0].style.name if cell.paragraphs and cell.paragraphs[0].style is not None else None
                for paragraph_index, paragraph in enumerate(cell.paragraphs):
                    paragraph_key = f"{key}.paragraph[{paragraph_index}]"
                    target_styles[paragraph_key] = paragraph.style.name if paragraph.style is not None else None
                    if _paragraph_has_image(paragraph):
                        image_targets.append(paragraph_key)
                if any(_paragraph_has_image(paragraph) for paragraph in cell.paragraphs):
                    image_targets.append(key)
    return OutputArtifactSnapshot(
        name=path.name,
        text_fragments=text_fragments,
        metadata={
            "format_key": "docx",
            "paragraph_count": len(document.paragraphs),
            "table_count": len(document.tables),
            "target_styles": target_styles,
            "image_targets": image_targets,
            "table_cell_tokens": _docx_table_cell_tokens(document),
            "docx_drawing_count": media_usage.drawing_count,
            "docx_media_part_count": len(media_usage.media_parts),
            "docx_referenced_media_part_count": len(media_usage.referenced_media_parts),
            "docx_orphan_media_part_count": len(media_usage.orphan_media_parts),
        },
    )


def _xlsx_snapshot(path: Path) -> OutputArtifactSnapshot:
    workbook = load_workbook(path)
    text_fragments: list[str] = []
    sheet_names = list(workbook.sheetnames)
    merged_ranges: dict[str, list[str]] = {}
    freeze_panes: dict[str, str | None] = {}
    cell_styles: dict[str, int] = {}
    image_targets: list[str] = []
    for worksheet in workbook.worksheets:
        merged_ranges[worksheet.title] = [str(item) for item in worksheet.merged_cells.ranges]
        freeze_panes[worksheet.title] = worksheet.freeze_panes if isinstance(worksheet.freeze_panes, str) else getattr(worksheet.freeze_panes, "coordinate", None)
        for image in getattr(worksheet, "_images", []):
            anchor = getattr(image, "anchor", None)
            origin = getattr(anchor, "_from", None)
            if origin is not None:
                image_targets.append(f"{worksheet.title}!{_column_letter(origin.col + 1)}{origin.row + 1}")
        for row in worksheet.iter_rows():
            for cell in row:
                cell_styles[f"{worksheet.title}!{cell.coordinate}"] = cell.style_id
                if cell.value is None:
                    continue
                text = str(cell.value).strip()
                if text:
                    text_fragments.append(text)
    return OutputArtifactSnapshot(
        name=path.name,
        text_fragments=text_fragments,
        metadata={
            "format_key": "xlsx",
            "sheet_names": sheet_names,
            "merged_ranges": merged_ranges,
            "freeze_panes": freeze_panes,
            "cell_styles": cell_styles,
            "image_targets": image_targets,
        },
    )


def _pdf_snapshot(path: Path) -> OutputArtifactSnapshot:
    return OutputArtifactSnapshot(
        name=path.name,
        text_fragments=[path.read_bytes().decode("latin-1", errors="ignore")],
        metadata={"format_key": "pdf"},
    )


def _paragraph_has_image(paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing"))


def _docx_table_cell_tokens(document) -> list[dict[str, object]]:
    tokens: list[dict[str, object]] = []
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table._tbl.tr_lst):
            for cell_index, cell in enumerate(row.tc_lst):
                tokens.append(
                    {
                        "table": table_index,
                        "row": row_index,
                        "cell": cell_index,
                        "grid_span": _docx_grid_span(cell),
                        "v_merge": _docx_vertical_merge(cell),
                    }
                )
    return tokens


def _docx_grid_span(cell) -> int:
    values = cell.xpath("./w:tcPr/w:gridSpan/@w:val")
    if not values:
        return 1
    try:
        return int(values[0])
    except (TypeError, ValueError):
        return 1


def _docx_vertical_merge(cell) -> str | None:
    values = cell.xpath("./w:tcPr/w:vMerge/@w:val")
    if values:
        return str(values[0])
    if cell.xpath("./w:tcPr/w:vMerge"):
        return "continue"
    return None


def _column_letter(index: int) -> str:
    letters = ""
    current = index
    while current > 0:
        current, remainder = divmod(current - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def _write_reports(
    output_path: Path,
    *,
    conversion_report: dict[str, Any],
    conversion_markdown: str,
    provenance_report: dict[str, Any],
    provenance_markdown: str,
    performance_report: dict[str, Any],
) -> None:
    conversion_json_path = output_path.with_suffix(".conversion_report.json")
    conversion_md_path = output_path.with_suffix(".conversion_report.md")
    provenance_json_path = output_path.with_suffix(".provenance_report.json")
    provenance_md_path = output_path.with_suffix(".provenance_report.md")
    performance_json_path = output_path.with_suffix(".performance_report.json")

    conversion_json_path.write_text(json.dumps(conversion_report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    conversion_md_path.write_text(conversion_markdown + "\n", encoding="utf-8")
    provenance_json_path.write_text(json.dumps(provenance_report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    provenance_md_path.write_text(provenance_markdown + "\n", encoding="utf-8")
    performance_json_path.write_text(json.dumps(performance_report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _write_performance_report(output_path: Path, performance_report: dict[str, Any]) -> None:
    output_path.with_suffix(".performance_report.json").write_text(
        json.dumps(performance_report, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def _materialization_plan_diagnostics(diagnostics: list[dict[str, Any]]) -> list[PlanDiagnostic]:
    return [
        PlanDiagnostic(
            severity=str(item.get("severity", "info")),
            code=str(item.get("code", "template_materialization")),
            message=str(item.get("message", "")),
            field=item.get("field"),
        )
        for item in diagnostics
        if isinstance(item, dict)
    ]


def _safe_len(value: Any) -> int | None:
    try:
        return len(value)
    except TypeError:
        return None


def _add_source_model_counts(metadata: dict[str, Any], source_model: AdaptedSourceModel) -> None:
    metadata["document_count"] = len(source_model.documents)
    metadata["record_count"] = len(source_model.records)
    metadata["image_asset_count"] = len(source_model.image_assets)
    metadata["diagnostic_count"] = len(source_model.diagnostics)


def _decisions_from_payload(payload_path: str | Path) -> list[AgentDecision]:
    payload = json.loads(Path(payload_path).expanduser().resolve().read_text(encoding="utf-8"))
    return [
        AgentDecision(
            slot_id=str(item["slot_id"]),
            decision=str(item["decision"]),
            candidate_id=item.get("candidate_id"),
            generated_value=item.get("generated_value"),
            confidence=float(item.get("confidence", 0.0)),
            reason=str(item.get("reason", "")),
            candidate_ids=[str(candidate_id) for candidate_id in item.get("candidate_ids", [])],
        )
        for item in payload.get("decisions", [])
        if isinstance(item, dict)
    ]
