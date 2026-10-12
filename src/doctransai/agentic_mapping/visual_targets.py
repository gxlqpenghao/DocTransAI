from __future__ import annotations

from typing import Any

from docx import Document
from openpyxl import load_workbook

from doctransai.models.template import TemplateModel


def resolve_visual_target_mapping(
    payload: dict[str, Any],
    template_model: TemplateModel,
) -> dict[str, Any]:
    if payload.get("target_kind") != "visual_text_anchor":
        return payload
    if _has_structured_target_address(payload):
        return _structured_visual_target(payload, template_model)

    label = _optional_string(payload.get("label_or_nearby_text") or payload.get("nearby_text") or payload.get("notes"))
    if label is None:
        return {**payload, "format_key": payload.get("format_key") or template_model.format_key}

    if template_model.format_key == "docx":
        return _resolve_docx_visual_text_anchor(payload, template_model, label)
    if template_model.format_key == "xlsx":
        return _resolve_xlsx_visual_text_anchor(payload, template_model, label)
    return {**payload, "format_key": payload.get("format_key") or template_model.format_key, "notes": label}


def _structured_visual_target(
    payload: dict[str, Any],
    template_model: TemplateModel,
) -> dict[str, Any]:
    resolved = {**payload, "format_key": payload.get("format_key") or template_model.format_key}
    if template_model.format_key == "xlsx" and payload.get("cell_ref"):
        resolved["target_kind"] = "cell"
    return resolved


def _resolve_docx_visual_text_anchor(
    payload: dict[str, Any],
    template_model: TemplateModel,
    label: str,
) -> dict[str, Any]:
    try:
        document = Document(template_model.template_path)
    except Exception:  # noqa: BLE001
        return {**payload, "format_key": payload.get("format_key") or template_model.format_key, "notes": label}

    matches: list[dict[str, Any]] = []
    for paragraph_index, paragraph in enumerate(document.paragraphs):
        if _visual_label_matches(paragraph.text, label):
            matches.append({
                **payload,
                "format_key": payload.get("format_key") or template_model.format_key,
                "paragraph_index": paragraph_index,
                "notes": label,
            })

    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            for column_index, cell in enumerate(row.cells):
                for cell_paragraph_index, paragraph in enumerate(cell.paragraphs):
                    if _visual_label_matches(paragraph.text, label):
                        matches.append({
                            **payload,
                            "format_key": payload.get("format_key") or template_model.format_key,
                            "table_index": table_index,
                            "row_index": row_index,
                            "column_index": column_index,
                            "cell_paragraph_index": cell_paragraph_index,
                            "notes": label,
                        })
    if matches:
        return _choose_docx_visual_match(matches, payload)
    return {**payload, "format_key": payload.get("format_key") or template_model.format_key, "notes": label}


def _choose_docx_visual_match(matches: list[dict[str, Any]], payload: dict[str, Any]) -> dict[str, Any]:
    if len(matches) == 1:
        return matches[0]
    page_index = payload.get("page_index")
    approx_location = str(payload.get("approx_location") or "").casefold()
    if (isinstance(page_index, int) and page_index > 0) or any(
        marker in approx_location
        for marker in ("bottom", "center", "middle", "下", "中")
    ):
        return matches[-1]
    return matches[0]


def _resolve_xlsx_visual_text_anchor(
    payload: dict[str, Any],
    template_model: TemplateModel,
    label: str,
) -> dict[str, Any]:
    try:
        workbook = load_workbook(template_model.template_path, data_only=False)
    except Exception:  # noqa: BLE001
        return {**payload, "format_key": payload.get("format_key") or template_model.format_key, "notes": label}

    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                value = "" if cell.value is None else str(cell.value)
                if not _visual_label_matches(value, label):
                    continue
                target_cell = _first_blank_cell_to_right(worksheet, cell.row, cell.column)
                if target_cell is None:
                    return {**payload, "format_key": "xlsx", "notes": label}
                return {
                    **payload,
                    "format_key": "xlsx",
                    "target_kind": "cell",
                    "sheet_name": worksheet.title,
                    "cell_ref": target_cell.coordinate,
                    "row_index": target_cell.row - 1,
                    "column_index": target_cell.column - 1,
                    "notes": label,
                }
    return {**payload, "format_key": "xlsx", "notes": label}


def _first_blank_cell_to_right(worksheet, row_index: int, column_index: int):
    for candidate_column in range(column_index + 1, max(worksheet.max_column + 2, column_index + 2)):
        candidate = worksheet.cell(row=row_index, column=candidate_column)
        if candidate.value in (None, ""):
            return candidate
    return None


def _has_structured_target_address(payload: dict[str, Any]) -> bool:
    return any(
        payload.get(key) is not None
        for key in (
            "paragraph_index",
            "table_index",
            "row_index",
            "column_index",
            "cell_ref",
            "form_field_name",
        )
    )


def _visual_label_matches(text: str, label: str) -> bool:
    if label in text:
        return True
    return _compact_visual_text(label) in _compact_visual_text(text)


def _compact_visual_text(text: str) -> str:
    return "".join(str(text).split())


def _optional_string(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value)
    return text if text else None


__all__ = ["resolve_visual_target_mapping"]
