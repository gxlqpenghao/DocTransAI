"""Versioned, atomic edits to a DOCX copy. No semantic mapping or LLM calls."""
from __future__ import annotations

import copy
import hashlib
from io import BytesIO
import json
import os
from pathlib import Path
import shutil
import tempfile
from typing import Callable
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import _Row
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from PIL import Image


class DocxEditor:
    def __init__(self, template: Path, output: Path, asset_path: Callable[[str], Path]):
        self.template, self.output = template.resolve(), output.resolve()
        if self.template == self.output:
            raise ValueError("Output must not overwrite the template")
        self.original_hash = hashlib.sha256(self.template.read_bytes()).hexdigest()
        self.doc = Document(self.template)
        self.asset_path = asset_path
        self.nodes, self.next_id = {}, 0
        self.version = 0
        self.applied_batches: dict[str, str] = {}
        self.touched_parts = {"word/document.xml"}
        self.inspect()
        shutil.copy2(self.template, self.output)

    @staticmethod
    def _address(node):
        element = node._element
        indexes = []
        while element.getparent() is not None:
            parent = element.getparent()
            indexes.append(parent.index(element))
            element = parent
        return str(node.part.partname), tuple(reversed(indexes))

    def inspect(self, *, _restore_ids=None, detail=False) -> dict:
        known = {id(node._element): key for key, node in self.nodes.items()}
        entries = []
        live = {}

        def add(node, where, kind):
            key = known.get(id(node._element))
            if key is None and _restore_ids is not None:
                key = _restore_ids.get(self._address(node))
            if key is None:
                key = f"n{self.next_id}"
                self.next_id += 1
            live[key] = node
            known[id(node._element)] = key
            text = node.text if hasattr(node, "text") else "\t".join(c.text for c in node.cells)
            entries.append(dict(id=key, kind=kind, where=where, text=text,
                                char_count=len(text), nested_tables=bool(getattr(node, "tables", [])),
                                body_index=node._element.getparent().index(node._element)
                                if node._element.getparent() is self.doc._element.body else None))
            if isinstance(node, Paragraph):
                try:
                    self.guard_text(node)
                    entries[-1].update(whole_text_replace_allowed=True, protection_reason=None)
                except ValueError as exc:
                    entries[-1].update(whole_text_replace_allowed=False, protection_reason=str(exc))
                entries[-1].update(self.paragraph_details(node) if detail else dict(
                    paragraph_style=node.style.name if node.style else None, run_count=len(node.runs)))
            return key

        for i, p in enumerate(self.doc.paragraphs):
            add(p, f"body paragraph {i}", "paragraph")
        def add_tables(tables, prefix):
            for ti, table in enumerate(tables):
                seen = set()
                for ri, row in enumerate(table.rows):
                    where = f"{prefix}table {ti} row {ri}"
                    row_id = add(row, where, "row")
                    row_entry = entries[-1]
                    row_entry.update(height_pt=row.height.pt if row.height else None,
                                     height_rule=str(row.height_rule) if row.height_rule is not None else "inherited",
                                     cant_split=any(x.get(qn("w:val"), "1") not in ("0", "false", "off")
                                                    for x in row._tr.xpath("./w:trPr/w:cantSplit")))
                    row_entry["physical_cells"] = [dict(grid_span=tc.grid_span, vertical_merge=tc.vMerge)
                                                   for tc in row._tr.tc_lst]
                    row_entry["content_cell_ids"] = []
                    for ci, cell in enumerate(row.cells):
                        if cell._tc in seen:
                            continue
                        seen.add(cell._tc)
                        cell_id = add(cell, f"{where} column {ci}", "cell")
                        row_entry["content_cell_ids"].append(cell_id)
                        entries[-1].update(row_id=row_id, grid_span=cell._tc.grid_span)
                        for pi, paragraph in enumerate(cell.paragraphs):
                            add(paragraph, f"cell {cell_id} paragraph {pi}", "paragraph")
                            entries[-1]["cell_id"] = cell_id
                        add_tables(cell.tables, f"cell {cell_id} ")
        add_tables(self.doc.tables, "")
        for si, section in enumerate(self.doc.sections):
            for name in ("header", "footer", "first_page_header", "first_page_footer", "even_page_header", "even_page_footer"):
                part = getattr(section, name)
                if part.is_linked_to_previous:
                    continue
                for pi, p in enumerate(part.paragraphs):
                    add(p, f"section {si} {name} paragraph {pi}", "paragraph")
                add_tables(part.tables, f"section {si} {name} ")
        self.nodes = live
        return dict(version=self.version, nodes=entries, image_count=len(self.doc.inline_shapes),
                    supported_operations=["set_text", "replace_text", "replace_range", "replace_ranges", "set_images", "repeat_paragraphs", "repeat_rows", "set_list", "format_row", "format_paragraph"],
                    limitations=["No nested-table replacement or floating textbox editing; repeat_rows requires a complete vertical-merge group"])

    @staticmethod
    def clone_structure(element):
        if element.xpath(".//w:sectPr | .//w:bookmarkStart | .//w:bookmarkEnd | .//w:commentRangeStart | .//w:commentRangeEnd"):
            raise ValueError("Cannot clone section boundaries, bookmarks or comments")
        clone = copy.deepcopy(element)
        # These editing-session identities must not be duplicated across clones.
        for child in clone.iter():
            for name in ("{http://schemas.microsoft.com/office/word/2010/wordml}paraId",
                         "{http://schemas.microsoft.com/office/word/2010/wordml}textId"):
                child.attrib.pop(name, None)
        return clone

    @staticmethod
    def paragraph_details(p):
        runs, offset = [], 0
        for index, run in enumerate(p.runs):
            runs.append(dict(index=index, start=offset, end=offset + len(run.text), text=run.text,
                             bold=run.bold, italic=run.italic, underline=str(run.underline) if run.underline is not None else None,
                             hidden=run.font.hidden, color=str(run.font.color.rgb) if run.font.color.rgb else None,
                             size_pt=run.font.size.pt if run.font.size else None))
            offset += len(run.text)
        num = p._p.xpath("./w:pPr/w:numPr")
        return dict(runs=runs, paragraph_style=p.style.name if p.style else None,
                    numbering_xml=num[0].xml if num else None,
                    range_unit="Python Unicode code points; end exclusive; offsets require current version",
                    protected=bool(p._p.xpath(".//w:drawing | .//w:pict | .//w:fldChar | .//w:fldSimple | .//w:hyperlink | .//w:sdt | .//w:ins | .//w:del | .//w:bookmarkStart | .//w:bookmarkEnd | .//w:commentRangeStart | .//w:commentRangeEnd")))

    def set_list(self, node, items, style, start):
        if style not in ("decimal", "bullet") or not items:
            raise ValueError("set_list needs items and decimal/bullet style")
        paragraphs = self.paragraphs(node)
        for p in paragraphs:
            self.guard_text(p)
        prototype = paragraphs[0]
        # A new numbering instance isolates this list from every other list.
        numbering = self.doc.part.numbering_part.element
        abstract_id = max([int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))] + [-1]) + 1
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        lvl = OxmlElement("w:lvl"); lvl.set(qn("w:ilvl"), "0")
        for tag, value in (("start", str(start)), ("numFmt", style), ("lvlText", "%1." if style == "decimal" else "•"), ("suff", "tab")):
            child = OxmlElement("w:" + tag); child.set(qn("w:val"), value); lvl.append(child)
        props = OxmlElement("w:pPr")
        indent = OxmlElement("w:ind"); indent.set(qn("w:left"), "360"); indent.set(qn("w:hanging"), "360")
        props.append(indent); lvl.append(props); abstract.append(lvl)
        first_num = numbering.find(qn("w:num"))
        if first_num is None:
            numbering.append(abstract)
        else:
            first_num.addprevious(abstract)
        num = numbering.add_num(abstract_id)
        anchor = prototype._p
        for text in items:
            element = self.clone_structure(prototype._p)
            anchor.addnext(element); anchor = element
            p = Paragraph(element, prototype._parent)
            self.put_text(p, text)
            pr = p._p.get_or_add_pPr().get_or_add_numPr()
            pr.get_or_add_ilvl().val = 0
            pr.get_or_add_numId().val = num.numId
            p.paragraph_format.keep_with_next = False
        for p in paragraphs:
            p._p.getparent().remove(p._p)
        self.touched_parts.add("word/numbering.xml")

    def repeat_rows(self, edit):
        ids = edit["targets"]
        rows = [self.nodes[x] for x in ids]
        if not rows or len(set(ids)) != len(ids) or any(not isinstance(r, _Row) for r in rows):
            raise ValueError("repeat_rows requires distinct row IDs")
        if any(a._tr.getnext() is not b._tr for a, b in zip(rows, rows[1:])):
            raise ValueError("Repeated rows must be consecutive")
        # Vertical merges may only be copied when their entire group is selected.
        if any(tc.vMerge == "continue" for tc in rows[0]._tr.tc_lst):
            raise ValueError("Selection starts inside a vertical merge")
        following = rows[-1]._tr.getnext()
        if following is not None and following.tag == qn("w:tr") and any(tc.vMerge == "continue" for tc in following.tc_lst):
            raise ValueError("Selection ends inside a vertical merge")
        prototypes = []
        for row in rows:
            for tc in row._tr.tc_lst:
                if tc.vMerge == "continue":
                    continue
                if tc.xpath(".//w:tbl | .//w:drawing | .//w:pict | .//w:bookmarkStart | .//w:fldChar | .//w:sdt"):
                    raise ValueError("Row prototypes with nested tables, drawings, bookmarks, fields or controls are unsupported")
                prototypes.append(tc)
        for item in edit["items"]:
            if len(item["contents"]) != len(prototypes):
                raise ValueError(f"Each item needs {len(prototypes)} contents entries for physical cells, excluding vertical continuations")
        from docx.table import _Cell
        anchor = rows[-1]._tr
        for item in edit["items"]:
            cells = []
            cloned_rows = []
            for original in rows:
                element = self.clone_structure(original._tr)
                anchor.addnext(element); anchor = element
                cloned_rows.append(element)
                cells.extend(_Cell(tc, original._parent) for tc in element.tc_lst if tc.vMerge != "continue")
            for cell, content in zip(cells, item["contents"]):
                if "asset_ids" in content:
                    self.add_images(cell, content["asset_ids"], content.get("max_height_inches", 2.5))
                else:
                    self.put_text(cell, content["text"], content.get("alignment"))
            if edit.get("keep_together", True):
                for index, tr in enumerate(cloned_rows):
                    pr = tr.get_or_add_trPr()
                    if pr.find(qn("w:cantSplit")) is None:
                        pr.append(OxmlElement("w:cantSplit"))
                    for tc in tr.tc_lst:
                        cell = _Cell(tc, rows[0]._parent)
                        for pi, paragraph in enumerate(cell.paragraphs):
                            paragraph.paragraph_format.keep_with_next = index < len(cloned_rows) - 1 or pi < len(cell.paragraphs) - 1
        self.touched_parts.add(str(rows[0].part.partname).lstrip("/"))
        for row in rows:
            row._tr.getparent().remove(row._tr)

    @staticmethod
    def paragraphs(node):
        if getattr(node, "tables", []):
            raise ValueError("Nested-table replacement is not supported")
        return node.paragraphs if hasattr(node, "paragraphs") else [node]

    @staticmethod
    def guard_text(p, *, preserve_anchors=False):
        if p._p.xpath(".//w:drawing | .//w:pict | .//w:fldChar | .//w:fldSimple | .//w:hyperlink | .//w:sdt | .//w:ins | .//w:del"):
            raise ValueError("Text operation would damage an image, field, hyperlink, tracked change or content control")
        if not preserve_anchors and p._p.xpath(".//w:bookmarkStart | .//w:bookmarkEnd | .//w:commentRangeStart | .//w:commentRangeEnd"):
            raise ValueError("Whole-text replacement would remove bookmarks/comments; use a local replace_text/replace_range that does not cross their boundaries")

    def put_text(self, node, text, alignment=None):
        paragraphs = self.paragraphs(node)
        for p in paragraphs:
            self.guard_text(p)
        p = paragraphs[0]
        rpr = copy.deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
        p.clear()
        run = p.add_run(text)
        if rpr is not None:
            run._r.insert(0, rpr)
        for extra in paragraphs[1:]:
            extra._p.getparent().remove(extra._p)
        if alignment:
            p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                           "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}[alignment]

    def replace_text(self, node, old, new):
        paragraphs = self.paragraphs(node)
        if not old or sum(p.text.count(old) for p in paragraphs) != 1:
            raise ValueError("old must occur exactly once within one paragraph; inspect the target again")
        p = next(p for p in paragraphs if old in p.text)
        self.replace_range(p, p.text.index(old), p.text.index(old) + len(old), old, new)

    def replace_range(self, p, start, end, expected, new, style_run_index=None, _style_override=None):
        if not isinstance(p, Paragraph):
            raise ValueError("replace_range requires a paragraph ID")
        if not 0 <= start < end <= len(p.text) or p.text[start:end] != expected:
            raise ValueError(f"Range/expected text mismatch at [{start}, {end}); actual={p.text[start:end]!r}, paragraph length={len(p.text)}; inspect the current paragraph")
        self.guard_text(p, preserve_anchors=True)
        if "".join(r.text for r in p.runs) != p.text:
            raise ValueError("This paragraph contains unsupported non-run text")
        position, matches = 0, []
        for run in p.runs:
            text = run.text
            right = position + len(text)
            if right > start and position < end:
                matches.append((run, position, right))
            position = right
        first, left, _ = matches[0]
        last, last_left, _ = matches[-1]
        children = list(p._p)
        span = children[children.index(first._r):children.index(last._r) + 1]
        point_anchors = set()
        for a, b in zip(children, children[1:]):
            if (a.tag == qn("w:bookmarkStart") and b.tag == qn("w:bookmarkEnd") and
                    a.get(qn("w:id")) is not None and a.get(qn("w:id")) == b.get(qn("w:id"))):
                point_anchors.update((a, b))
        if any(child not in point_anchors and child.tag in {qn("w:bookmarkStart"), qn("w:bookmarkEnd"), qn("w:commentRangeStart"), qn("w:commentRangeEnd")} for child in span):
            raise ValueError("Replacement crosses a bookmark/comment boundary; select a smaller range")
        prefix, suffix = first.text[:start - left], last.text[end - last_left:]
        fills = (p._p.xpath("./w:pPr/w:shd/@w:fill") or
                 p._p.xpath("ancestor::w:tc[1]/w:tcPr/w:shd/@w:fill") or
                 self.doc._element.xpath("./w:background/@w:color"))
        background = str(fills[0]).upper() if fills else "FFFFFF"
        def visible(run):
            rgb = run.font.color.rgb
            return run.font.hidden is not True and (rgb is None or str(rgb).upper() != background)
        # A match may start in an invisible X marker and continue into a visible
        # placeholder. Use the visible matched text's style, without recoloring
        # any prefix/suffix outside the replacement range.
        style_run = next((r for r, _, _ in matches if visible(r)), first)
        if style_run_index is not None:
            if not 0 <= style_run_index < len(p.runs):
                raise ValueError("style_run_index outside paragraph")
            style_run = p.runs[style_run_index]
        style = copy.deepcopy(style_run._r.rPr)
        if _style_override and _style_override[0]:
            style = copy.deepcopy(_style_override[1])
        suffix_style = copy.deepcopy(last._r.rPr)
        for run, _, _ in matches:
            run.text = ""
        first.text = prefix
        element = OxmlElement("w:r")
        first._r.addnext(element)
        inserted = Run(element, p)
        if style is not None:
            element.insert(0, style)
        inserted.text = new
        if last is first and suffix:
            element_after = OxmlElement("w:r")
            element.addnext(element_after)
            if suffix_style is not None:
                element_after.insert(0, suffix_style)
            Run(element_after, p).text = suffix
        elif last is not first:
            last.text = suffix

    def replace_ranges(self, p, replacements):
        """All coordinates and style indexes refer to one paragraph snapshot."""
        if not isinstance(p, Paragraph) or not replacements:
            raise ValueError("replace_ranges requires a paragraph and nonempty replacements")
        text = p.text
        ordered = sorted(replacements, key=lambda item: item["start"])
        for index, item in enumerate(ordered):
            if not 0 <= item["start"] < item["end"] <= len(text) or text[item["start"]:item["end"]] != item["expected"]:
                raise ValueError("replace_ranges expected text mismatch in base snapshot")
            if index and ordered[index-1]["end"] > item["start"]:
                raise ValueError("replace_ranges intervals overlap")
            style_index = item.get("style_run_index")
            if style_index is not None and not 0 <= style_index < len(p.runs):
                raise ValueError("style_run_index outside base paragraph")
        # Explicit styles must not drift as preceding edits add runs.
        styles = [copy.deepcopy(p.runs[item["style_run_index"]]._r.rPr)
                  if item.get("style_run_index") is not None else None for item in ordered]
        for item, style in reversed(list(zip(ordered, styles))):
            self.replace_range(p, item["start"], item["end"], item["expected"], item["text"],
                               _style_override=(item.get("style_run_index") is not None, style))

    def add_images(self, node, ids, max_height=2.5):
        paragraphs = self.paragraphs(node)
        p = paragraphs[0]
        # This operation explicitly replaces the image region, unlike text edits.
        p.clear()
        for extra in paragraphs[1:]:
            extra._p.getparent().remove(extra._p)
        body_width = min(s.page_width.inches - s.left_margin.inches - s.right_margin.inches for s in self.doc.sections)
        container = node._parent if isinstance(node, Paragraph) else node
        width = min(body_width, container.width.inches - .18) if getattr(container, "width", None) else body_width
        for index, ref in enumerate(ids):
            path = self.asset_path(ref)
            with Image.open(path) as image:
                w, h = image.size
            if index:
                p.add_run().add_break()
            p.add_run().add_picture(str(path), width=Inches(max(.1, min(width, max_height * w / h))))

    def apply(self, *, batch_id: str, version: int, edits: list[dict]) -> dict:
        if not batch_id or not edits:
            raise ValueError("A nonempty batch_id and edits are required")
        digest = hashlib.sha256(json.dumps(dict(version=version, edits=edits), sort_keys=True, ensure_ascii=False).encode()).hexdigest()
        if batch_id in self.applied_batches:
            if self.applied_batches[batch_id] != digest:
                raise ValueError("batch_id already used for different edits")
            return dict(already_applied=True, **self.inspect())
        if version != self.version:
            raise ValueError(f"Stale document version {version}; current version is {self.version}")
        prior_ids = {self._address(node): key for key, node in self.nodes.items()}
        prior_next_id, prior_parts = self.next_id, set(self.touched_parts)
        for edit in edits:
            if edit["op"] == "repeat_paragraphs":
                expected = len(edit["targets"])
                actual = [len(item["contents"]) for item in edit["items"]]
                if any(count != expected for count in actual):
                    raise ValueError(f"Each repeated item needs {expected} contents entries in target order; got {actual}. Include the first AND subsequent complete groups; do not fill the prototypes first.")
        try:
            for operation_index, edit in enumerate(edits):
                op = edit["op"]
                if op == "repeat_rows":
                    self.repeat_rows(edit)
                elif op == "repeat_paragraphs":
                    ids = edit["targets"]
                    if len(ids) != len(set(ids)):
                        raise ValueError("Repeated paragraph IDs must be distinct")
                    originals = [self.nodes[x] for x in ids]
                    if any(not isinstance(p, Paragraph) for p in originals):
                        raise ValueError("repeat_paragraphs requires paragraph IDs")
                    for p in originals:
                        self.guard_text(p)
                    if any(a._p.getnext() is not b._p for a, b in zip(originals, originals[1:])):
                        raise ValueError("Repeated paragraphs must be consecutive")
                    self.touched_parts.add(str(originals[0].part.partname).lstrip("/"))
                    anchor = originals[-1]._p
                    for item in edit["items"]:
                        if len(item["contents"]) != len(originals):
                            raise ValueError("Each item needs one content per original paragraph")
                        for i, (original, content) in enumerate(zip(originals, item["contents"])):
                            element = self.clone_structure(original._p)
                            anchor.addnext(element)
                            anchor = element
                            p = Paragraph(element, original._parent)
                            if "asset_ids" in content:
                                self.add_images(p, content["asset_ids"], content.get("max_height_inches", 2.5))
                            else:
                                self.put_text(p, content["text"], content.get("alignment"))
                            p.paragraph_format.keep_together = edit.get("keep_together", True)
                            p.paragraph_format.keep_with_next = edit.get("keep_together", True) and i < len(originals) - 1
                    for ref, p in zip(ids, originals):
                        p._p.getparent().remove(p._p)
                        del self.nodes[ref]
                else:
                    node = self.nodes[edit["target"]]
                    self.touched_parts.add(str(node.part.partname).lstrip("/"))
                    if op == "replace_ranges":
                        self.replace_ranges(node, edit["replacements"])
                    elif op == "replace_range":
                        self.replace_range(node, edit["start"], edit["end"], edit["expected"], edit["text"], edit.get("style_run_index"))
                    elif op == "set_list":
                        self.set_list(node, edit["items"], edit["style"], edit.get("start", 1))
                    elif op == "set_text":
                        self.put_text(node, edit["text"], edit.get("alignment"))
                    elif op == "replace_text":
                        self.replace_text(node, edit["old"], edit["text"])
                    elif op == "set_images":
                        self.add_images(node, edit["asset_ids"], edit.get("max_height_inches", 2.5))
                    elif op == "format_row":
                        if not isinstance(node, _Row):
                            raise ValueError("format_row requires a row ID")
                        options = edit["format"]
                        if "height_pt" in options and "height_rule" not in options:
                            raise ValueError("height_pt requires height_rule")
                        if "height_rule" in options:
                            rule = options["height_rule"]
                            if rule != "auto" and "height_pt" not in options:
                                raise ValueError("Non-auto row height requires height_pt")
                            node.height_rule = {"auto": WD_ROW_HEIGHT_RULE.AUTO, "at_least": WD_ROW_HEIGHT_RULE.AT_LEAST,
                                                "exact": WD_ROW_HEIGHT_RULE.EXACTLY}[rule]
                            node.height = None if rule == "auto" else Pt(options["height_pt"])
                        if "allow_split" in options:
                            pr = node._tr.get_or_add_trPr()
                            for element in list(pr.findall(qn("w:cantSplit"))):
                                pr.remove(element)
                            if not options["allow_split"]:
                                pr.append(OxmlElement("w:cantSplit"))
                    elif op == "format_paragraph":
                        for p in self.paragraphs(node):
                            options = edit["format"]
                            for key in ("keep_with_next", "keep_together", "page_break_before"):
                                if key in options:
                                    setattr(p.paragraph_format, key, options[key])
                            for key in ("space_before", "space_after", "line_spacing"):
                                if key + "_pt" in options:
                                    setattr(p.paragraph_format, key, Pt(options[key + "_pt"]))
                            if "font_size_pt" in options:
                                for run in p.runs or [p.add_run()]:
                                    run.font.size = Pt(options["font_size_pt"])
                    else:
                        raise ValueError(f"Unsupported edit operation: {op}")
                self.inspect()  # Retire removed descendants before the next edit in this batch.
            self.save()
        except Exception as exc:
            # Deepcopy of python-docx proxies may produce a second, disconnected
            # part tree. Reload the last atomic file and bind the old IDs to it.
            self.doc = Document(self.output)
            self.nodes, self.next_id, self.touched_parts = {}, prior_next_id, prior_parts
            self.inspect(_restore_ids=prior_ids)
            exc.edit_failure = dict(failed_operation_index=operation_index, rolled_back=True, input_version=version, output_version=self.version, message_state="Uncommitted intermediate state; whole batch rolled back")
            raise
        self.version += 1
        self.applied_batches[batch_id] = digest
        return dict(already_applied=False, **self.inspect())

    def save(self):
        buffer = BytesIO()
        self.doc.save(buffer)
        allowed = {"[Content_Types].xml", *self.touched_parts}
        for name in self.touched_parts:
            part = Path(name)
            allowed.add(str(part.parent / "_rels" / (part.name + ".rels")))
        fd, tmp = tempfile.mkstemp(dir=self.output.parent, suffix=".docx")
        os.close(fd)
        try:
            with zipfile.ZipFile(self.template) as original, zipfile.ZipFile(buffer) as edited, zipfile.ZipFile(tmp, "w") as out:
                names = set(original.namelist())
                for info in original.infolist():
                    out.writestr(info, edited.read(info.filename) if info.filename in allowed else original.read(info.filename))
                for info in edited.infolist():
                    if info.filename not in names:
                        out.writestr(info, edited.read(info.filename))
            os.replace(tmp, self.output)
        finally:
            if os.path.exists(tmp):
                os.unlink(tmp)
