"""Atomic workbook edits with explicit addresses; no business-field inference."""
from __future__ import annotations

from copy import copy
from datetime import datetime
import hashlib
import json
import math
import os
import posixpath
from pathlib import Path
import re
import shutil
import tempfile
import warnings
import zipfile

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter, range_boundaries, coordinate_to_tuple
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.workbook.properties import CalcProperties


class XlsxEditor:
    format_key = 'xlsx'
    operations = ['set_cells', 'set_formula', 'copy_rows', 'set_cell_images', 'format_cells', 'set_print_area']

    def __init__(self, template: Path, output: Path, asset_path):
        self.template, self.output = template.resolve(), output.resolve()
        if self.template == self.output:
            raise ValueError('Output must not overwrite template')
        self.asset_path = asset_path
        self.original_hash = hashlib.sha256(self.template.read_bytes()).hexdigest()
        with zipfile.ZipFile(template) as z:
            unsupported = [n for n in z.namelist() if any(k in n for k in
                ('vbaProject', 'slicer', 'pivotTable', 'pivotCache', 'externalLink', 'activeX', 'vmlDrawing', 'ctrlProps', 'threadedComments', 'xl/metadata.xml'))]
            if unsupported:
                raise ValueError('Unsupported workbook objects require a preserving adapter: ' + ', '.join(unsupported[:5]))
            # openpyxl drops arbitrary drawing shapes. Refuse rather than silently lose them.
            from lxml import etree
            for name in z.namelist():
                if name.startswith('xl/drawings/drawing') and name.endswith('.xml'):
                    tree = etree.fromstring(z.read(name))
                    if tree.xpath('//*[local-name()="sp" or local-name()="grpSp" or local-name()="cxnSp"]'):
                        raise ValueError('Drawing shapes are not supported by this workbook editor')
        self.version, self.applied_batches, self.nodes = 0, {}, {}
        self.book = self._load(template)
        self.original_formulas = self._formulas()
        self.formula_changes = set()
        shutil.copy2(template, output)
        self.inspect()

    def _preserve_custom_xml(self, saved):
        """Retain opaque custom-XML metadata and its package declarations."""
        from lxml import etree
        with zipfile.ZipFile(self.template) as original:
            opaque={n:original.read(n) for n in original.namelist() if n.startswith('customXml/')}
            if not opaque:return
            with zipfile.ZipFile(saved) as edited:
                parts={n:edited.read(n) for n in edited.namelist()}
            parts.update(opaque)
            for name in original.namelist():
                if not name.endswith('.rels') or name.startswith('customXml/'):continue
                old=etree.fromstring(original.read(name))
                extras=[r for r in old if r.get('Type','').endswith('/customXml')]
                if not extras:continue
                if name not in parts:raise ValueError('Custom XML owner relationship part was not retained')
                new=etree.fromstring(parts[name]);ids={r.get('Id') for r in new}
                for relation in extras:
                    relation=copy(relation)
                    index=1
                    while f'rId{index}' in ids:index+=1
                    relation.set('Id',f'rId{index}');ids.add(f'rId{index}');new.append(relation)
                parts[name]=etree.tostring(new,xml_declaration=True,encoding='UTF-8')
            old=etree.fromstring(original.read('[Content_Types].xml'))
            new=etree.fromstring(parts['[Content_Types].xml'])
            known={e.get('PartName') for e in new}
            extensions={e.get('Extension') for e in new}
            opaque_extensions={n.rsplit('.',1)[-1] for n in opaque}
            for entry in old:
                if entry.get('PartName','').startswith('/customXml/') and entry.get('PartName') not in known:
                    new.append(copy(entry))
                elif entry.get('Extension') in opaque_extensions and entry.get('Extension') not in extensions:
                    new.append(copy(entry));extensions.add(entry.get('Extension'))
            parts['[Content_Types].xml']=etree.tostring(new,xml_declaration=True,encoding='UTF-8')
        with zipfile.ZipFile(saved,'w',compression=zipfile.ZIP_DEFLATED) as target:
            for name,data in parts.items():target.writestr(name,data)

    @staticmethod
    def _validate_relationships(path):
        from lxml import etree
        with zipfile.ZipFile(path) as package:
            names=set(package.namelist())
            for name in names:
                if not name.endswith('.rels'):continue
                base=name.rsplit('/_rels/',1)[0] if '/_rels/' in name else ''
                for rel in etree.fromstring(package.read(name)):
                    if rel.get('TargetMode')=='External':continue
                    target=rel.get('Target','')
                    resolved=target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(base,target))
                    if resolved not in names:
                        raise ValueError(f'Workbook relationship target missing: {resolved}')

    @staticmethod
    def _load(path):
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter('always')
            book = load_workbook(path, data_only=False, rich_text=True)
        if caught:
            raise ValueError('Workbook reader cannot guarantee preservation: ' + '; '.join(str(w.message) for w in caught))
        return book

    def _formulas(self):
        return {(s.title, c.coordinate): self._formula_value(c.value) for s in self.book for c in s._cells.values() if c.data_type == 'f'}

    @staticmethod
    def _formula_value(value):
        return value if isinstance(value, str) else json.dumps(vars(value), sort_keys=True)

    def resolve(self, target):
        match = re.fullmatch(r's(\d+)!([A-Z]+[1-9]\d*(?::[A-Z]+[1-9]\d*)?)', target)
        if not match or int(match[1]) >= len(self.book.worksheets):
            raise ValueError('Use an exact address such as s0!A3 or s0!A3:D5 from inspect_document')
        bounds = range_boundaries(match[2])
        a, b, c, d = bounds
        if not (a <= c <= 16384 and b <= d <= 1048576) or (c-a+1)*(d-b+1) > 2000:
            raise ValueError('Range is invalid or exceeds 2000 cells; request a smaller range')
        return self.book.worksheets[int(match[1])], bounds

    @staticmethod
    def _value(value):
        return value.isoformat() if isinstance(value, datetime) else value if isinstance(value, (str,int,float,bool,type(None))) else str(value)

    @staticmethod
    def _is_material_cell(cell):
        return cell.value is not None or cell.data_type == 'f'

    @classmethod
    def _used_bounds(cls, ws):
        """Find a non-destructive structural range from existing worksheet objects."""
        coordinates = [(cell.row, cell.column) for cell in ws._cells.values() if cell.has_style or cls._is_material_cell(cell)]
        for merged in ws.merged_cells.ranges:
            coordinates.extend(((merged.min_row, merged.min_col), (merged.max_row, merged.max_col)))
        for image in ws._images:
            anchor = cls._image_cell(image)
            if anchor:
                coordinates.append(coordinate_to_tuple(anchor))
        if not coordinates:
            return (1, 1, 1, 1)
        rows = [row for row, _ in coordinates]
        columns = [column for _, column in coordinates]
        return min(columns), min(rows), max(columns), max(rows)

    def inspect(self, *, detail=False, node_ids=None):
        sheets, nodes = [], []
        requested = {}
        for target in node_ids or []:
            ws, bounds = self.resolve(target)
            requested.setdefault(ws.title, []).append(bounds)
        for si, ws in enumerate(self.book):
            used_a, used_b, used_c, used_d = self._used_bounds(ws)
            material_count = sum(self._is_material_cell(cell) for cell in ws._cells.values())
            sheets.append(dict(id=f's{si}', name=ws.title, state=ws.sheet_state, rows=ws.max_row, columns=ws.max_column,
                merged_ranges=[str(r) for r in ws.merged_cells.ranges], formulas=sum(c.data_type=='f' for c in ws._cells.values()),
                images=[dict(anchor=self._image_cell(im), width=im.width, height=im.height) for im in ws._images],
                print_area=str(ws.print_area), print_title_rows=ws.print_title_rows, print_title_columns=ws.print_title_cols, freeze_panes=str(ws.freeze_panes),
                tables=list(ws.tables), validation_count=len(ws.data_validations.dataValidation),
                conditional_formats=len(ws.conditional_formatting),
                used_range=f'{get_column_letter(used_a)}{used_b}:{get_column_letter(used_c)}{used_d}',
                non_empty_cells=material_count))
            if not node_ids:
                view_c, view_d = min(used_c, 40), min(used_d, 20)
                sheets[-1]['default_view'] = f'A1:{get_column_letter(view_c)}{view_d}'
                sheets[-1]['view_truncated'] = used_c > 40 or used_d > 20
                sheets[-1]['default_view_note'] = 'Only non-empty cells and formulas are listed; use inspect_document with an exact range for blank cells and formatting.'
                sheets[-1]['omitted_non_empty_cells'] = sum(self._is_material_cell(cell) and
                    (cell.row > view_d or cell.column > view_c) for cell in ws._cells.values())
                cell_iter = (cell for (row, column), cell in sorted(ws._cells.items())
                             if row <= view_d and column <= view_c and self._is_material_cell(cell))
                ranges = None
            else:
                ranges = requested.get(ws.title, [])
                cell_iter = None
            seen=set()
            if cell_iter is None:
                cell_iter = (ws._cells.get((ri, ci)) or Cell(ws, row=ri, column=ci)
                             for a, b, c, d in ranges
                             for ri in range(b, d + 1)
                             for ci in range(a, c + 1))
            for cell in cell_iter:
                if cell.coordinate in seen:
                    continue
                seen.add(cell.coordinate)
                key=f's{si}!{cell.coordinate}'
                merged=next((str(m) for m in ws.merged_cells.ranges if cell.coordinate in m), None)
                n=dict(id=key,kind='cell',data_type=cell.data_type,value_type=type(cell.value).__name__,text=str(self._value(cell.value) or ''),value=self._value(cell.value),
                    formula=cell.value if cell.data_type=='f' else None, merged_range=merged,
                    writable=not isinstance(cell,MergedCell) and cell.data_type!='f' and not ws.protection.sheet,
                    number_format=cell.number_format)
                if detail:
                    n.update(row_height=ws.row_dimensions.get(cell.row).height if ws.row_dimensions.get(cell.row) else None,
                        column_width=self._column_width(ws, cell.column),
                        wrap_text=cell.alignment.wrap_text,horizontal=cell.alignment.horizontal,
                        font_size=cell.font.sz)
                nodes.append(n)
        self.nodes={n['id']:n for n in nodes}
        return dict(format='xlsx',version=self.version,sheets=sheets,nodes=nodes,
            image_count=sum(len(s._images) for s in self.book),supported_operations=self.operations,
            limitations=['copy_rows writes empty destinations; it does not insert/shift existing rows',
                         'Formula caches recalculate when opened by a spreadsheet engine',
                         'Complex drawing shapes, pivots, macros and external links are rejected'])

    @staticmethod
    def _column_width(ws, column):
        explicit = ws.column_dimensions.get(get_column_letter(column))
        if explicit is not None and explicit.width is not None:return explicit.width
        for dimension in ws.column_dimensions.values():
            if dimension.min and dimension.max and dimension.min <= column <= dimension.max:
                if dimension.width is not None:return dimension.width
        return ws.sheet_format.defaultColWidth or 8.43

    @staticmethod
    def _image_cell(image):
        a=image.anchor
        return a if isinstance(a,str) else f'{get_column_letter(a._from.col+1)}{a._from.row+1}' if hasattr(a,'_from') else None

    @staticmethod
    def _editable(ws, cell, *, formula=False):
        if ws.protection.sheet: raise ValueError('Sheet is protected')
        if isinstance(cell,MergedCell): raise ValueError('Write only the top-left cell of a merged range')
        if cell.data_type=='f' and not formula: raise ValueError('Formula cell is protected; use set_formula with expected_formula')

    def _set_cells(self, ws, bounds, edit):
        a,b,c,d=bounds; values=edit['values']
        if len(values)!=d-b+1 or any(len(row)!=c-a+1 for row in values):
            raise ValueError('values shape must exactly match target range')
        for ri,row in enumerate(values,b):
            for ci,value in enumerate(row,a):
                cell=ws.cell(ri,ci);self._editable(ws,cell)
                if isinstance(value,dict):
                    date_text = value['date']
                    value=datetime.fromisoformat(date_text)
                    if value.tzinfo:raise ValueError('Excel dates must have no timezone')
                    default_format = 'yyyy-mm-dd hh:mm' if 'T' in date_text or ' ' in date_text else 'yyyy-mm-dd'
                    cell.number_format=edit.get('number_format', default_format)
                if isinstance(value, str) and len(value) > 32767:
                    raise ValueError('Cell text exceeds Excel limit; summarize explicitly instead of truncating')
                if isinstance(value, float) and not math.isfinite(value):
                    raise ValueError('Cell number must be finite')
                cell.value=value
                if isinstance(value,str) and value.startswith('='):
                    cell.data_type='s'  # Literal text is never an implicit executable formula.
                if 'number_format' in edit:cell.number_format=edit['number_format']

    def _copy_rows(self, ws, bounds, destinations):
        a,b,c,d=bounds
        source=list(ws.iter_rows(min_row=b,max_row=d,min_col=a,max_col=c))
        for row in source:
            for cell in row:
                if cell.hyperlink or cell.comment:raise ValueError('Row copying with links/comments is unsupported')
                if cell.data_type=='f' and not isinstance(cell.value,str):raise ValueError('Array/shared formula row copying is unsupported')
        merges=[]
        for m in ws.merged_cells.ranges:
            if m.max_row>=b and m.min_row<=d and m.max_col>=a and m.min_col<=c:
                if not (b<=m.min_row<=m.max_row<=d and a<=m.min_col<=m.max_col<=c):
                    raise ValueError('Source must include entire merged ranges')
                merges.append(copy(m))
        if ws.tables or len(ws.conditional_formatting) or ws.data_validations.dataValidation:
            raise ValueError('Row copying with tables/conditional formats/validation needs an explicit range-extension adapter')
        if any(self._image_cell(im) and b<=ws[self._image_cell(im)].row<=d for im in ws._images):
            raise ValueError('Copy rows before adding images; attach each record image explicitly')
        occupied=set()
        for start in destinations:
            if start<1 or start+d-b>1048576:raise ValueError('Destination row outside worksheet')
            rows=set(range(start,start+d-b+1))
            if rows & (occupied|set(range(b,d+1))):raise ValueError('Copy destinations overlap')
            occupied|=rows
            if any(m.max_row>=start and m.min_row<=start+d-b and m.max_col>=a and m.min_col<=c for m in ws.merged_cells.ranges):
                raise ValueError('Copy destination intersects a merged range')
            for ri in rows:
                for ci in range(a,c+1):
                    cell=ws.cell(ri,ci)
                    if cell.value is not None or isinstance(cell,MergedCell) or cell.hyperlink or cell.comment:
                        raise ValueError('Copy destination must be empty and unmerged')
            if any(self._image_cell(im) and ws[self._image_cell(im)].row in rows for im in ws._images):
                raise ValueError('Copy destination contains images')
            for ri,row in enumerate(source,start):
                dim=copy(ws.row_dimensions[b+ri-start]);dim.index=ri;ws.row_dimensions[ri]=dim
                for original in row:
                    target=ws.cell(ri,original.column)
                    target._style=copy(original._style)
                    if isinstance(original,MergedCell):continue
                    target.value=Translator(original.value,origin=original.coordinate).translate_formula(target.coordinate) if original.data_type=='f' else original.value
                    target.data_type=original.data_type
                    if target.data_type=='f':self.formula_changes.add((ws.title,target.coordinate))
            for m in merges:
                m.shift(row_shift=start-b);ws.merge_cells(str(m));m.shift(row_shift=b-start)

    def _images(self, ws, bounds, refs):
        a,b,c,d=bounds
        if a!=c or b!=d:raise ValueError('Image target must be a single cell or merged-range top-left cell')
        cell=ws.cell(b,a);self._editable(ws,cell)
        merge=next((m for m in ws.merged_cells.ranges if cell.coordinate in m),None)
        if merge:c,d=merge.max_col,merge.max_row
        width=sum(self._column_width(ws,i)*7+5 for i in range(a,c+1))
        height=sum(((ws.row_dimensions.get(i).height if ws.row_dimensions.get(i) else None) or ws.sheet_format.defaultRowHeight or 15)*96/72 for i in range(b,d+1))
        ws._images=[im for im in ws._images if self._image_cell(im)!=cell.coordinate]
        for index,ref in enumerate(refs):
            im=XLImage(str(self.asset_path(ref)));available=max(1,(height-4)/len(refs))
            scale=min(max(1,width-4)/im.width,available/im.height)
            im.width*=scale;im.height*=scale
            marker=AnchorMarker(col=a-1,row=b-1,colOff=pixels_to_EMU(2),rowOff=pixels_to_EMU(2+index*available))
            im.anchor=OneCellAnchor(_from=marker,ext=XDRPositiveSize2D(pixels_to_EMU(im.width),pixels_to_EMU(im.height)))
            ws.add_image(im)

    def apply(self, *, batch_id, version, edits):
        digest=hashlib.sha256(json.dumps(dict(version=version,edits=edits),sort_keys=True,ensure_ascii=False).encode()).hexdigest()
        if batch_id in self.applied_batches:
            if self.applied_batches[batch_id]!=digest:raise ValueError('batch_id already used for different edits')
            return dict(already_applied=True,**self.inspect())
        if version!=self.version:raise ValueError('Stale workbook version')
        if not batch_id or not edits:raise ValueError('Nonempty batch required')
        prior_changes=set(self.formula_changes)
        try:
            for edit in edits:
                ws,bounds=self.resolve(edit['target']);a,b,c,d=bounds
                if ws.protection.sheet:raise ValueError('Sheet is protected')
                op=edit['op']
                if op=='set_cells':self._set_cells(ws,bounds,edit)
                elif op=='set_formula':
                    if a!=c or b!=d:raise ValueError('Formula requires one cell')
                    cell=ws.cell(b,a);self._editable(ws,cell,formula=True)
                    current=cell.value if cell.data_type=='f' else None
                    if current!=edit['expected_formula']:raise ValueError('expected_formula mismatch')
                    if not edit['formula'].startswith('=') or '#REF!' in edit['formula']:raise ValueError('Invalid formula')
                    cell.value=edit['formula'];self.formula_changes.add((ws.title,cell.coordinate))
                elif op=='copy_rows':self._copy_rows(ws,bounds,edit['destination_rows'])
                elif op=='set_cell_images':self._images(ws,bounds,edit['asset_ids'])
                elif op=='format_cells':
                    f=edit['format']
                    for row in ws.iter_rows(min_row=b,max_row=d,min_col=a,max_col=c):
                        for cell in row:
                            if isinstance(cell,MergedCell):continue
                            alignment=copy(cell.alignment)
                            for key in ('wrap_text','horizontal','vertical'):
                                if key in f:setattr(alignment,key,f[key])
                            cell.alignment=alignment
                            if 'number_format' in f:cell.number_format=f['number_format']
                            if 'font_size' in f:font=copy(cell.font);font.sz=f['font_size'];cell.font=font
                    if 'row_height' in f:
                        for ri in range(b,d+1):ws.row_dimensions[ri].height=f['row_height']
                    if 'column_width' in f:
                        for ci in range(a,c+1):ws.column_dimensions[get_column_letter(ci)].width=f['column_width']
                elif op=='set_print_area':
                    outside = [cell.coordinate for cell in ws._cells.values()
                               if cell.value is not None and not (a <= cell.column <= c and b <= cell.row <= d)]
                    outside_images = [self._image_cell(image) for image in ws._images
                                      if self._image_cell(image) and not (a <= ws[self._image_cell(image)].column <= c and
                                                                           b <= ws[self._image_cell(image)].row <= d)]
                    if outside or outside_images:
                        raise ValueError('print area cannot hide existing content or images')
                    ws.print_area=f'{get_column_letter(a)}{b}:{get_column_letter(c)}{d}'
                    if 'orientation' in edit:ws.page_setup.orientation=edit['orientation']
                    if 'fit_to_width' in edit:
                        ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.fitToWidth=edit['fit_to_width'];ws.page_setup.fitToHeight=0
                    if 'repeat_rows' in edit:ws.print_title_rows=edit['repeat_rows']
                    if 'repeat_columns' in edit:ws.print_title_cols=edit['repeat_columns']
                else:raise ValueError(f'Unsupported XLSX operation: {op}')
            for key,formula in self.original_formulas.items():
                if key not in self.formula_changes and self._formula_value(self.book[key[0]][key[1]].value)!=formula:
                    raise ValueError(f'Unintended formula change: {key}')
            self.book.calculation=CalcProperties(fullCalcOnLoad=True,forceFullCalc=True,calcMode='auto')
            fd,tmp=tempfile.mkstemp(dir=self.output.parent,suffix='.xlsx');os.close(fd)
            try:
                self.book.save(tmp)
                self._preserve_custom_xml(Path(tmp))
                self._validate_relationships(Path(tmp))
                restored=self._load(tmp)  # Reopen images after save; image streams are single-use.
                os.replace(tmp,self.output);self.book=restored
            finally:
                if os.path.exists(tmp):os.unlink(tmp)
        except Exception:
            self.book=self._load(self.output);self.formula_changes=prior_changes;self.inspect();raise
        self.version+=1;self.applied_batches[batch_id]=digest
        return dict(already_applied=False,**self.inspect())
