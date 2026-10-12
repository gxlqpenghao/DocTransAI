from __future__ import annotations

import base64
import inspect
import json
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Callable

from docx import Document
from jsonschema import Draft202012Validator
from openpyxl import load_workbook

from doctransai.mapping.cluster_planner import (
    RepeatRegionClusterPlan,
    SourceClusterPlan,
    SourceClusterPlanningResult,
    build_source_cluster_plan,
)
from doctransai.mapping.full_agent_context import build_full_agent_context
from doctransai.mapping.prompts import (
    build_mapping_repair_prompt,
    build_mapping_system_prompt,
    build_mapping_user_prompt,
    build_repair_system_prompt,
)
from doctransai.mapping.type_inference import infer_slot_value_type, infer_source_value_type
from doctransai.llm import (
    estimate_chat_input_chars,
    estimate_chat_request_bytes,
    generate_text_with_doctransai_llm,
    resolve_llm_input_char_limit,
    resolve_llm_request_body_limit_bytes,
)
from doctransai.models.mapping import ImageCandidate, UnresolvedSlot
from doctransai.models.source_instance import SourceCluster, SourceInstance, SourceInstanceIndex
from doctransai.models.render_plan import TextFill
from doctransai.models.source import ProvenanceRef, SourceModel, SourceRecord
from doctransai.source_model import build_source_instance_index
from doctransai.models.template import TemplateModel, TemplateSlot
from doctransai.templates.visual_preview import render_template_visual_pages


AGENT_MAPPING_SCHEMA_VERSION = "doctransai_agent_mapping_decision_v1"
_GENERIC_TEXT_KEYS = {"text", "table_text", "range_text", "sheet_summary", "body_text", "row_text"}
_DEFAULT_MAX_TEXT_CANDIDATES_PER_SLOT: int | None = None
_DEFAULT_MAX_IMAGE_CANDIDATES_PER_SLOT: int | None = None
_SOURCE_OVERVIEW_RECORD_LIMIT = 256
_SOURCE_OVERVIEW_VALUE_MAX_CHARS = 1200
_SOURCE_OVERVIEW_DIAGNOSTIC_LIMIT = 30
_IMAGE_CONTEXT_VALUE_MAX_CHARS = 180
_SOURCE_INSTANCE_CLUSTER_INSTANCE_LIMIT = 32
_PACKAGING_MODE_FULL = "full"
_PACKAGING_MODE_COMPACT_SLOT_PACKETS = "compact_slot_packets"
_PACKAGING_MODE_PREFLIGHT_BLOCKED = "preflight_blocked"
_LEGACY_SHARED_CONTEXT_KEYS = [
    "source_overview",
    "source_document_view",
    "template_overview",
    "source_instance_overview",
    "source_cluster_plan",
]
_COMPACT_VALUE_PREVIEW_MAX_CHARS = 240


@dataclass(slots=True)
class AgentCandidatePacket:
    candidate_id: str
    candidate_kind: str
    value_label: str
    value_text: str
    value_type: str
    relation: str
    confidence_hint: float
    locator: dict[str, Any] = field(default_factory=dict)
    source_path: str | None = None
    provenance_refs: list[dict[str, Any]] = field(default_factory=list)
    source_object_refs: list[dict[str, Any]] = field(default_factory=list)


@dataclass(slots=True)
class AgentSlotPacket:
    slot_id: str
    raw_name: str
    slot_kind: str
    expected_value_type: str
    context: dict[str, Any] = field(default_factory=dict)
    candidates: list[AgentCandidatePacket] = field(default_factory=list)


@dataclass(slots=True)
class AgentDecision:
    slot_id: str
    decision: str
    candidate_id: str | None
    generated_value: str | None
    confidence: float
    reason: str
    candidate_ids: list[str] = field(default_factory=list)


@dataclass(slots=True)
class AgentResolutionTrace:
    decisions: list[AgentDecision] = field(default_factory=list)
    raw_output_text: str = ""
    normalized_payload: dict[str, Any] = field(default_factory=dict)
    repaired: bool = False
    repair_raw_output_text: str | None = None
    validation_errors: list[str] = field(default_factory=list)
    payload_preflight: dict[str, Any] = field(default_factory=dict)


class AgentResolutionError(ValueError):
    pass


def prepare_agent_slot_packets(
    *,
    source_model: SourceModel,
    primary_record: SourceRecord,
    unresolved_slots: list[UnresolvedSlot],
    template_slots_by_id: dict[str, TemplateSlot],
    max_candidates_per_slot: int | None = None,
    include_text_candidates: bool = True,
) -> list[AgentSlotPacket]:
    packets: list[AgentSlotPacket] = []
    for unresolved_slot in unresolved_slots:
        slot = template_slots_by_id.get(unresolved_slot.slot_id)
        if slot is None:
            continue
        expected_value_type = infer_slot_value_type(slot)

        if slot.slot_kind == "image":
            candidates = _image_candidate_packets(source_model=source_model, primary_record=primary_record, slot=slot)
        elif not include_text_candidates:
            candidates = []
        else:
            candidates = _text_candidate_packets(
                source_model=source_model,
                primary_record=primary_record,
                expected_value_type=expected_value_type,
            )

        if not candidates and slot.slot_kind == "image":
            continue
        ranked_candidates = _rank_candidates_for_slot(candidates, slot)
        candidate_limit = _candidate_limit(slot, max_candidates_per_slot, candidate_count=len(ranked_candidates))

        context = _slot_context(slot, resolution_mode="unresolved")
        if slot.slot_kind != "image" and not include_text_candidates:
            context["text_candidate_material"] = "omitted_per_slot_to_avoid_duplicate_payload_use_shared_source_overview"

        packets.append(
            AgentSlotPacket(
                slot_id=slot.slot_id,
                raw_name=slot.raw_name,
                slot_kind=slot.slot_kind,
                expected_value_type=expected_value_type,
                context=context,
                candidates=ranked_candidates[:candidate_limit],
            )
        )
    return packets


def prepare_agent_review_slot_packets(
    *,
    source_model: SourceModel,
    primary_record: SourceRecord,
    review_text_fills: list[TextFill],
    template_slots_by_id: dict[str, TemplateSlot],
    max_candidates_per_slot: int | None = None,
    include_text_candidates: bool = True,
) -> list[AgentSlotPacket]:
    packets: list[AgentSlotPacket] = []
    for text_fill in review_text_fills:
        slot = template_slots_by_id.get(text_fill.slot_id)
        if slot is None or slot.slot_kind == "image":
            continue
        expected_value_type = infer_slot_value_type(slot)
        candidates = (
            _text_candidate_packets(
                source_model=source_model,
                primary_record=primary_record,
                expected_value_type=expected_value_type,
            )
            if include_text_candidates
            else []
        )
        context = _slot_context(slot, resolution_mode="review")
        if not include_text_candidates:
            context["text_candidate_material"] = "omitted_per_slot_to_avoid_duplicate_payload_use_shared_source_overview"
        context.update(
            {
                "current_value": text_fill.value,
                "current_confidence": text_fill.confidence,
                "current_transform": text_fill.transform,
                "review_reason": _review_reason(text_fill),
            }
        )
        ranked_candidates = _rank_candidates_for_slot(candidates, slot)
        candidate_limit = _candidate_limit(slot, max_candidates_per_slot, candidate_count=len(ranked_candidates))
        packets.append(
            AgentSlotPacket(
                slot_id=slot.slot_id,
                raw_name=slot.raw_name,
                slot_kind=slot.slot_kind,
                expected_value_type=expected_value_type,
                context=context,
                candidates=ranked_candidates[:candidate_limit],
            )
        )
    return packets


def prepare_agent_shared_context(
    *,
    source_model: SourceModel,
    primary_record: SourceRecord,
    template_model: TemplateModel | None,
    template_slots_by_id: dict[str, TemplateSlot],
    source_instance_id: str | None = None,
) -> dict[str, Any]:
    source_overview = _source_overview(source_model, primary_record)
    source_document_view = _full_source_document_view(source_model.records)
    template_overview = _template_overview(template_model, template_slots_by_id)
    source_instance_index = build_source_instance_index(source_model)
    source_cluster_plan = build_source_cluster_plan(source_model=source_model, template_model=template_model) if template_model else None
    source_instance_overview = _source_instance_overview(
        source_instance_index=source_instance_index,
        primary_record=primary_record,
    )
    source_cluster_plan_overview = _source_cluster_plan_overview(source_cluster_plan)
    full_agent_context = build_full_agent_context(
        source_model=source_model,
        primary_record=primary_record,
        template_model=template_model,
        template_slots_by_id=template_slots_by_id,
        source_document_view=source_document_view,
        source_instance_overview=source_instance_overview,
        source_cluster_plan=source_cluster_plan_overview,
        task_context={
            "primary_record_id": primary_record.record_id,
            "source_instance_id": source_instance_id,
            "context_phase": "shared_agent_context",
        },
        template_document_view=template_overview.get("document_view") if isinstance(template_overview, dict) else None,
        template_visual_pages=template_overview.get("visual_pages") if isinstance(template_overview, dict) else None,
    )
    full_source_context = full_agent_context["full_source_context"]
    navigation_context = full_agent_context["navigation_context"]
    return {
        "template_overview": template_overview,
        "source_overview": source_overview,
        "full_agent_context": full_agent_context,
        "source_document_view": full_source_context["document_view"],
        "source_instance_overview": navigation_context["source_instance_overview"],
        "source_cluster_plan": navigation_context["source_cluster_plan"],
    }


def resolve_mapping_decisions(
    slot_packets: list[AgentSlotPacket],
    *,
    decision_client: Callable[..., str] | None = None,
    shared_context: dict[str, Any] | None = None,
) -> list[AgentDecision]:
    return resolve_mapping_decisions_with_trace(
        slot_packets,
        decision_client=decision_client,
        shared_context=shared_context,
    ).decisions


def resolve_mapping_decisions_with_trace(
    slot_packets: list[AgentSlotPacket],
    *,
    decision_client: Callable[..., str] | None = None,
    shared_context: dict[str, Any] | None = None,
) -> AgentResolutionTrace:
    if not slot_packets:
        return AgentResolutionTrace()

    client = decision_client or generate_text_with_doctransai_llm
    system_prompt = build_mapping_system_prompt()
    user_payload, user_prompt, user_parts, payload_preflight = _build_agent_request(
        slot_packets=slot_packets,
        shared_context=shared_context,
        system_prompt=system_prompt,
        phase="initial",
    )
    if payload_preflight["preflight_blocked"]:
        return _payload_too_large_trace(slot_packets, payload_preflight=payload_preflight)

    raw_output = _invoke_decision_client(
        client,
        system_prompt=system_prompt,
        user_prompt=user_prompt,
        user_parts=user_parts,
    )
    payload, validation_errors = _parse_and_validate_payload(raw_output, slot_packets)
    repaired = False
    repair_raw_output_text: str | None = None

    if validation_errors:
        repaired = True
        repair_prompt = build_mapping_repair_prompt(
            slot_payload=user_payload,
            previous_output=raw_output,
            validation_errors=validation_errors,
        )
        repair_system_prompt = build_repair_system_prompt()
        repair_payload_preflight = _agent_payload_preflight(
            system_prompt=repair_system_prompt,
            user_prompt=repair_prompt,
            user_parts=None,
            enforce=True,
            phase="repair",
            agent_input_packaging=_repair_packaging_metadata(user_payload.get("agent_input_packaging")),
        )
        if repair_payload_preflight["preflight_blocked"]:
            return _payload_too_large_trace(
                slot_packets,
                payload_preflight=repair_payload_preflight,
                raw_output_text=raw_output,
                repaired=True,
                validation_errors=validation_errors,
            )
        repair_raw_output_text = _invoke_decision_client(
            client,
            system_prompt=repair_system_prompt,
            user_prompt=repair_prompt,
            user_parts=None,
        )
        payload, validation_errors = _parse_and_validate_payload(repair_raw_output_text, slot_packets)
        if validation_errors:
            details = "\n".join(f"- {item}" for item in validation_errors)
            raise AgentResolutionError(
                "Agent output is still invalid after one structured repair retry:\n" + details
            )

    decisions: list[AgentDecision] = []
    for item in payload["decisions"]:
        slot_id = str(item["slot_id"])
        candidate_id = item.get("candidate_id")
        decision = str(item["decision"])
        selected_candidate_ids = _candidate_ids_from_payload_item(item)
        decisions.append(
            AgentDecision(
                slot_id=slot_id,
                decision=decision,
                candidate_id=str(candidate_id) if candidate_id is not None else None,
                generated_value=item.get("generated_value"),
                confidence=_coerce_confidence(item.get("confidence"), default=0.0 if decision == "unresolved" else 0.75),
                reason=str(item["reason"]),
                candidate_ids=selected_candidate_ids if decision == "candidate" else [],
            )
        )
    return AgentResolutionTrace(
        decisions=decisions,
        raw_output_text=raw_output,
        normalized_payload=payload,
        repaired=repaired,
        repair_raw_output_text=repair_raw_output_text,
        validation_errors=validation_errors,
        payload_preflight=repair_payload_preflight if repaired else payload_preflight,
    )


def _slot_payload(
    slot_packets: list[AgentSlotPacket],
    *,
    shared_context: dict[str, Any] | None = None,
    packaging_mode: str = _PACKAGING_MODE_FULL,
    agent_input_packaging: dict[str, Any] | None = None,
) -> dict[str, Any]:
    payload = {
        "schema_version": AGENT_MAPPING_SCHEMA_VERSION,
        "slots": [
            _compact_slot_packet(packet) if packaging_mode == _PACKAGING_MODE_COMPACT_SLOT_PACKETS else asdict(packet)
            for packet in slot_packets
        ],
    }
    if shared_context:
        payload.update(_shared_context_for_packaging(shared_context, packaging_mode=packaging_mode))
    if agent_input_packaging is not None:
        payload["agent_input_packaging"] = dict(agent_input_packaging)
    return payload


def _build_agent_request(
    *,
    slot_packets: list[AgentSlotPacket],
    shared_context: dict[str, Any] | None,
    system_prompt: str,
    phase: str,
) -> tuple[dict[str, Any], str, list[dict[str, Any]] | None, dict[str, Any]]:
    full_payload = _slot_payload(slot_packets, shared_context=shared_context, packaging_mode=_PACKAGING_MODE_FULL)
    full_user_prompt = build_mapping_user_prompt(full_payload)
    full_user_parts = _mapping_user_parts(shared_context=full_payload, user_prompt=full_user_prompt)
    full_preflight = _agent_payload_preflight(
        system_prompt=system_prompt,
        user_prompt=full_user_prompt,
        user_parts=full_user_parts,
        enforce=True,
        phase=phase,
    )
    if not full_preflight["exceeded"]:
        metadata = _agent_input_packaging_metadata(
            mode=_PACKAGING_MODE_FULL,
            full_estimated_bytes=full_preflight["estimated_request_bytes"],
            compact_estimated_bytes=None,
            selected_estimated_bytes=full_preflight["estimated_request_bytes"],
            limit_bytes=full_preflight["max_request_body_bytes"],
            full_estimated_chars=full_preflight["estimated_input_chars"],
            compact_estimated_chars=None,
            selected_estimated_chars=full_preflight["estimated_input_chars"],
            char_limit=full_preflight["max_input_chars"],
            request_body_exceeded=full_preflight["request_body_exceeded"],
            input_length_exceeded=full_preflight["input_length_exceeded"],
            slot_packets_compacted=False,
            legacy_fields_removed=[],
            phase=phase,
        )
        return _finalized_agent_request(
            slot_packets=slot_packets,
            shared_context=shared_context,
            system_prompt=system_prompt,
            phase=phase,
            packaging_mode=_PACKAGING_MODE_FULL,
            metadata=metadata,
        )

    compact_payload = _slot_payload(
        slot_packets,
        shared_context=shared_context,
        packaging_mode=_PACKAGING_MODE_COMPACT_SLOT_PACKETS,
    )
    compact_user_prompt = build_mapping_user_prompt(compact_payload)
    compact_user_parts = _mapping_user_parts(shared_context=compact_payload, user_prompt=compact_user_prompt)
    compact_preflight = _agent_payload_preflight(
        system_prompt=system_prompt,
        user_prompt=compact_user_prompt,
        user_parts=compact_user_parts,
        enforce=True,
        phase=phase,
    )
    selected_mode = (
        _PACKAGING_MODE_PREFLIGHT_BLOCKED
        if compact_preflight["exceeded"]
        else _PACKAGING_MODE_COMPACT_SLOT_PACKETS
    )
    metadata = _agent_input_packaging_metadata(
        mode=selected_mode,
        full_estimated_bytes=full_preflight["estimated_request_bytes"],
        compact_estimated_bytes=compact_preflight["estimated_request_bytes"],
        selected_estimated_bytes=compact_preflight["estimated_request_bytes"],
        limit_bytes=compact_preflight["max_request_body_bytes"],
        full_estimated_chars=full_preflight["estimated_input_chars"],
        compact_estimated_chars=compact_preflight["estimated_input_chars"],
        selected_estimated_chars=compact_preflight["estimated_input_chars"],
        char_limit=compact_preflight["max_input_chars"],
        request_body_exceeded=compact_preflight["request_body_exceeded"],
        input_length_exceeded=compact_preflight["input_length_exceeded"],
        slot_packets_compacted=True,
        legacy_fields_removed=_legacy_fields_present(shared_context),
        phase=phase,
    )
    return _finalized_agent_request(
        slot_packets=slot_packets,
        shared_context=shared_context,
        system_prompt=system_prompt,
        phase=phase,
        packaging_mode=_PACKAGING_MODE_COMPACT_SLOT_PACKETS,
        metadata=metadata,
    )


def _finalized_agent_request(
    *,
    slot_packets: list[AgentSlotPacket],
    shared_context: dict[str, Any] | None,
    system_prompt: str,
    phase: str,
    packaging_mode: str,
    metadata: dict[str, Any],
) -> tuple[dict[str, Any], str, list[dict[str, Any]] | None, dict[str, Any]]:
    payload = _slot_payload(
        slot_packets,
        shared_context=shared_context,
        packaging_mode=packaging_mode,
        agent_input_packaging=metadata,
    )
    user_prompt = build_mapping_user_prompt(payload)
    user_parts = _mapping_user_parts(shared_context=payload, user_prompt=user_prompt)
    preflight = _agent_payload_preflight(
        system_prompt=system_prompt,
        user_prompt=user_prompt,
        user_parts=user_parts,
        enforce=True,
        phase=phase,
        agent_input_packaging=metadata,
    )
    metadata = dict(metadata)
    metadata["selected_estimated_bytes"] = preflight["estimated_request_bytes"]
    metadata["selected_estimated_chars"] = preflight["estimated_input_chars"]
    metadata["limit_bytes"] = preflight["max_request_body_bytes"]
    metadata["char_limit"] = preflight["max_input_chars"]
    metadata["request_body_exceeded"] = preflight["request_body_exceeded"]
    metadata["input_length_exceeded"] = preflight["input_length_exceeded"]
    metadata["within_limit"] = not preflight["exceeded"]
    metadata["preflight_blocked"] = preflight["preflight_blocked"]
    if preflight["preflight_blocked"]:
        metadata["mode"] = _PACKAGING_MODE_PREFLIGHT_BLOCKED
    payload["agent_input_packaging"] = metadata
    user_prompt = build_mapping_user_prompt(payload)
    user_parts = _mapping_user_parts(shared_context=payload, user_prompt=user_prompt)
    preflight = _agent_payload_preflight(
        system_prompt=system_prompt,
        user_prompt=user_prompt,
        user_parts=user_parts,
        enforce=True,
        phase=phase,
        agent_input_packaging=metadata,
    )
    return payload, user_prompt, user_parts, preflight


def _agent_input_packaging_metadata(
    *,
    mode: str,
    full_estimated_bytes: int,
    compact_estimated_bytes: int | None,
    selected_estimated_bytes: int,
    limit_bytes: int | None,
    full_estimated_chars: int,
    compact_estimated_chars: int | None,
    selected_estimated_chars: int,
    char_limit: int | None,
    request_body_exceeded: bool,
    input_length_exceeded: bool,
    slot_packets_compacted: bool,
    legacy_fields_removed: list[str],
    phase: str,
) -> dict[str, Any]:
    return {
        "mode": mode,
        "full_estimated_bytes": full_estimated_bytes,
        "compact_estimated_bytes": compact_estimated_bytes,
        "selected_estimated_bytes": selected_estimated_bytes,
        "limit_bytes": limit_bytes,
        "full_estimated_chars": full_estimated_chars,
        "compact_estimated_chars": compact_estimated_chars,
        "selected_estimated_chars": selected_estimated_chars,
        "char_limit": char_limit,
        "request_body_exceeded": request_body_exceeded,
        "input_length_exceeded": input_length_exceeded,
        "slot_packets_compacted": slot_packets_compacted,
        "full_context_compacted": slot_packets_compacted,
        "legacy_fields_removed": list(legacy_fields_removed),
        "phase": phase,
        "within_limit": not (request_body_exceeded or input_length_exceeded),
        "preflight_blocked": bool(request_body_exceeded or input_length_exceeded),
    }


def _repair_packaging_metadata(value: object) -> dict[str, Any] | None:
    if not isinstance(value, dict):
        return None
    metadata = dict(value)
    metadata["phase"] = "repair"
    return metadata


def _shared_context_for_packaging(
    shared_context: dict[str, Any],
    *,
    packaging_mode: str,
) -> dict[str, Any]:
    if packaging_mode != _PACKAGING_MODE_COMPACT_SLOT_PACKETS:
        return dict(shared_context)
    result: dict[str, Any] = {}
    if "full_agent_context" in shared_context:
        result["full_agent_context"] = _compact_full_agent_context_for_transport(shared_context["full_agent_context"])
    return result


def _compact_full_agent_context_for_transport(value: object) -> object:
    if not isinstance(value, dict):
        return value
    context = dict(value)
    full_source_context = context.get("full_source_context")
    if isinstance(full_source_context, dict):
        source_context = dict(full_source_context)
        records = source_context.get("records")
        if isinstance(records, list):
            source_context["records"] = [_compact_source_record_for_transport(record) for record in records]
        image_assets = source_context.get("image_assets")
        if isinstance(image_assets, list):
            source_context["image_assets"] = [_compact_mapping_for_transport(asset) for asset in image_assets]
        context["full_source_context"] = source_context
    return context


def _compact_source_record_for_transport(value: object) -> object:
    if not isinstance(value, dict):
        return value
    return _compact_mapping_for_transport(
        {
            key: value.get(key)
            for key in [
                "record_id",
                "document_id",
                "source_path",
                "values",
                "evidence_id",
                "cell_range_id",
                "table_id",
                "locator",
            ]
        }
    )


def _compact_mapping_for_transport(value: object) -> object:
    if isinstance(value, dict):
        compacted: dict[str, Any] = {}
        for key, item in value.items():
            if item is None or item == "" or item == () or item == []:
                continue
            if isinstance(item, dict) and not item:
                continue
            compacted[key] = _compact_mapping_for_transport(item)
        return compacted
    if isinstance(value, list):
        return [
            _compact_mapping_for_transport(item)
            for item in value
            if item is not None and item != "" and item != () and item != []
        ]
    return value


def _legacy_fields_present(shared_context: dict[str, Any] | None) -> list[str]:
    if not shared_context:
        return []
    return [key for key in _LEGACY_SHARED_CONTEXT_KEYS if key in shared_context]


def _compact_slot_packet(packet: AgentSlotPacket) -> dict[str, Any]:
    return {
        "slot_id": packet.slot_id,
        "raw_name": packet.raw_name,
        "slot_kind": packet.slot_kind,
        "expected_value_type": packet.expected_value_type,
        "context": dict(packet.context),
        "candidates": [_compact_candidate_packet(candidate) for candidate in packet.candidates],
    }


def _compact_candidate_packet(candidate: AgentCandidatePacket) -> dict[str, Any]:
    return {
        "candidate_id": candidate.candidate_id,
        "candidate_kind": candidate.candidate_kind,
        "value_label": candidate.value_label,
        "value_preview": _truncate_text(candidate.value_text, _COMPACT_VALUE_PREVIEW_MAX_CHARS),
        "value_length": len(candidate.value_text),
        "value_type": candidate.value_type,
        "relation": candidate.relation,
        "confidence_hint": candidate.confidence_hint,
        "primary_locator": dict(candidate.locator),
        "source_path": candidate.source_path,
        "provenance_refs": list(candidate.provenance_refs),
        "source_object_refs": list(candidate.source_object_refs),
        "candidate_lookup_keys": _candidate_lookup_keys(candidate),
    }


def _candidate_lookup_keys(candidate: AgentCandidatePacket) -> list[str]:
    keys = [candidate.candidate_id]
    for ref in candidate.source_object_refs:
        if isinstance(ref, dict) and isinstance(ref.get("object_key"), str):
            keys.append(ref["object_key"])
    return list(dict.fromkeys(keys))


def _agent_payload_preflight(
    *,
    system_prompt: str,
    user_prompt: str,
    user_parts: list[dict[str, Any]] | None,
    enforce: bool,
    phase: str,
    agent_input_packaging: dict[str, Any] | None = None,
) -> dict[str, Any]:
    estimated_request_bytes = estimate_chat_request_bytes(
        system_prompt=system_prompt,
        user_prompt=user_prompt,
        user_parts=user_parts,
    )
    estimated_input_chars = estimate_chat_input_chars(
        system_prompt=system_prompt,
        user_prompt=user_prompt,
        user_parts=user_parts,
    )
    max_request_body_bytes = resolve_llm_request_body_limit_bytes()
    max_input_chars = resolve_llm_input_char_limit()
    request_body_exceeded = (
        max_request_body_bytes is not None
        and estimated_request_bytes > max_request_body_bytes
    )
    input_length_exceeded = (
        max_input_chars is not None
        and estimated_input_chars > max_input_chars
    )
    exceeded = request_body_exceeded or input_length_exceeded
    return {
        "estimated_request_bytes": estimated_request_bytes,
        "max_request_body_bytes": max_request_body_bytes,
        "estimated_input_chars": estimated_input_chars,
        "max_input_chars": max_input_chars,
        "request_body_exceeded": request_body_exceeded,
        "input_length_exceeded": input_length_exceeded,
        "exceeded": exceeded,
        "preflight_blocked": bool(enforce and exceeded),
        "phase": phase,
        "agent_input_packaging": _preflight_packaging_metadata(
            agent_input_packaging,
            estimated_request_bytes=estimated_request_bytes,
            max_request_body_bytes=max_request_body_bytes,
            estimated_input_chars=estimated_input_chars,
            max_input_chars=max_input_chars,
            request_body_exceeded=request_body_exceeded,
            input_length_exceeded=input_length_exceeded,
            exceeded=exceeded,
            preflight_blocked=bool(enforce and exceeded),
            phase=phase,
        ),
    }


def _preflight_packaging_metadata(
    metadata: dict[str, Any] | None,
    *,
    estimated_request_bytes: int,
    max_request_body_bytes: int | None,
    estimated_input_chars: int,
    max_input_chars: int | None,
    request_body_exceeded: bool,
    input_length_exceeded: bool,
    exceeded: bool,
    preflight_blocked: bool,
    phase: str,
) -> dict[str, Any]:
    result = dict(metadata or {})
    if "mode" not in result:
        result["mode"] = _PACKAGING_MODE_PREFLIGHT_BLOCKED if preflight_blocked else _PACKAGING_MODE_FULL
    if preflight_blocked:
        result["mode"] = _PACKAGING_MODE_PREFLIGHT_BLOCKED
    result["selected_estimated_bytes"] = estimated_request_bytes
    result["limit_bytes"] = max_request_body_bytes
    result["selected_estimated_chars"] = estimated_input_chars
    result["char_limit"] = max_input_chars
    result["request_body_exceeded"] = request_body_exceeded
    result["input_length_exceeded"] = input_length_exceeded
    result["within_limit"] = not exceeded
    result["preflight_blocked"] = preflight_blocked
    result["phase"] = phase
    return result


def _payload_too_large_trace(
    slot_packets: list[AgentSlotPacket],
    *,
    payload_preflight: dict[str, Any],
    raw_output_text: str = "",
    repaired: bool = False,
    validation_errors: list[str] | None = None,
) -> AgentResolutionTrace:
    decisions = [_payload_too_large_decision(packet) for packet in slot_packets]
    return AgentResolutionTrace(
        decisions=decisions,
        raw_output_text=raw_output_text,
        normalized_payload={
            "schema_version": AGENT_MAPPING_SCHEMA_VERSION,
            "decisions": [asdict(decision) for decision in decisions],
        },
        repaired=repaired,
        validation_errors=validation_errors or [],
        payload_preflight=payload_preflight,
    )


def _payload_too_large_decision(packet: AgentSlotPacket) -> AgentDecision:
    return AgentDecision(
        slot_id=packet.slot_id,
        decision="unresolved",
        candidate_id=None,
        generated_value=None,
        confidence=0.0,
        reason="agent_input_payload_too_large",
    )


def _slot_context(slot: TemplateSlot, *, resolution_mode: str) -> dict[str, Any]:
    return {
        "resolution_mode": resolution_mode,
        "canonical_name": slot.canonical_name,
        "context_before": slot.context_before,
        "context_after": slot.context_after,
        "address_kind": slot.address.address_kind,
        "placement_constraints": dict(slot.placement_constraints),
    }


def _review_reason(text_fill: TextFill) -> str:
    transform = text_fill.transform or ""
    if "typed_row_match" in transform:
        return "typed_row_match_auto_fill"
    if transform.startswith("unique_"):
        return "unique_typed_row_auto_fill"
    if text_fill.confidence < 0.9:
        return "low_confidence_auto_fill"
    return "auto_fill_review"


def _candidate_limit(slot: TemplateSlot, explicit_limit: int | None, *, candidate_count: int) -> int:
    if explicit_limit is not None:
        return explicit_limit
    default_limit = _DEFAULT_MAX_IMAGE_CANDIDATES_PER_SLOT if slot.slot_kind == "image" else _DEFAULT_MAX_TEXT_CANDIDATES_PER_SLOT
    if default_limit is None:
        return candidate_count
    if slot.slot_kind == "image":
        return default_limit
    return default_limit


def _rank_candidates_for_slot(
    candidates: list[AgentCandidatePacket],
    slot: TemplateSlot,
) -> list[AgentCandidatePacket]:
    if slot.slot_kind == "image":
        relation_priority = {"same_record_row": 0, "near_primary_record_context": 1, "document_context": 2}
        return sorted(
            candidates,
            key=lambda candidate: (
                relation_priority.get(candidate.relation, 99),
                -_candidate_relevance_score(candidate, slot),
                -candidate.confidence_hint,
                candidate.candidate_id,
            ),
        )
    return sorted(
        candidates,
        key=lambda candidate: (
            -_candidate_relevance_score(candidate, slot),
            -candidate.confidence_hint,
            candidate.candidate_id,
        ),
    )


def _candidate_relevance_score(candidate: AgentCandidatePacket, slot: TemplateSlot) -> float:
    score = candidate.confidence_hint
    if candidate.relation == "same_record_row":
        score += 3.0
    elif candidate.relation == "near_primary_record_context":
        score += 2.0

    slot_text = " ".join(
        item
        for item in [
            slot.raw_name,
            slot.canonical_name,
            slot.context_before,
            slot.context_after,
        ]
        if item
    )
    candidate_label = f"{candidate.value_label} {candidate.value_type}"
    score += _text_overlap_score(slot_text, candidate_label) * 4.0
    if candidate.value_text and len(candidate.value_text) <= 80:
        score += _text_overlap_score(slot_text, candidate.value_text)
    if candidate.candidate_kind == slot.slot_kind:
        score += 0.4
    return score


def _text_overlap_score(left: str, right: str) -> float:
    left_tokens = _text_signal_tokens(left)
    right_tokens = _text_signal_tokens(right)
    if not left_tokens or not right_tokens:
        return 0.0
    overlap = left_tokens & right_tokens
    return len(overlap) / max(len(left_tokens), 1)


def _text_signal_tokens(text: str) -> set[str]:
    normalized = "".join(ch for ch in text if ch.isalnum() or "\u4e00" <= ch <= "\u9fff")
    if not normalized:
        return set()
    tokens = {normalized[index : index + 2] for index in range(max(len(normalized) - 1, 0))}
    tokens.update(ch for ch in normalized if "\u4e00" <= ch <= "\u9fff")
    return {token for token in tokens if token.strip()}


def _mapping_user_parts(
    *,
    shared_context: dict[str, Any] | None,
    user_prompt: str,
) -> list[dict[str, Any]] | None:
    if not shared_context:
        return None

    template_overview = _template_visual_context(shared_context)
    if not isinstance(template_overview, dict):
        return None

    parts: list[dict[str, Any]] = [{"type": "text", "text": user_prompt}]
    for image_path in template_overview.get("visual_pages") or []:
        data_url = _image_path_to_data_url(image_path)
        if data_url is None:
            continue
        parts.append(
            {
                "type": "image_url",
                "image_url": {
                    "url": data_url,
                    "detail": "high",
                },
            }
        )
    return parts if len(parts) > 1 else None


def _template_visual_context(shared_context: dict[str, Any]) -> dict[str, Any] | None:
    template_overview = shared_context.get("template_overview")
    if isinstance(template_overview, dict):
        return template_overview
    full_agent_context = shared_context.get("full_agent_context")
    if not isinstance(full_agent_context, dict):
        return None
    full_template_context = full_agent_context.get("full_template_context")
    return full_template_context if isinstance(full_template_context, dict) else None


def _image_path_to_data_url(raw_path: object) -> str | None:
    if not isinstance(raw_path, str):
        return None
    path = Path(raw_path).expanduser().resolve()
    if not path.exists():
        return None
    suffix = path.suffix.lower()
    mime_type = "image/png" if suffix == ".png" else "image/jpeg" if suffix in {".jpg", ".jpeg"} else None
    if mime_type is None:
        return None
    encoded = base64.b64encode(path.read_bytes()).decode("ascii")
    return f"data:{mime_type};base64,{encoded}"


def _invoke_decision_client(
    client: Callable[..., str],
    *,
    system_prompt: str,
    user_prompt: str,
    user_parts: list[dict[str, Any]] | None,
) -> str:
    try:
        signature = inspect.signature(client)
    except (TypeError, ValueError):
        signature = None

    supports_user_parts = False
    if signature is not None:
        for parameter in signature.parameters.values():
            if parameter.kind == inspect.Parameter.VAR_KEYWORD:
                supports_user_parts = True
                break
        else:
            supports_user_parts = "user_parts" in signature.parameters

    if supports_user_parts:
        return client(system_prompt=system_prompt, user_prompt=user_prompt, user_parts=user_parts)
    return client(system_prompt=system_prompt, user_prompt=user_prompt)


def _template_overview(
    template_model: TemplateModel | None,
    template_slots_by_id: dict[str, TemplateSlot],
) -> dict[str, Any]:
    if template_model is None:
        return {
            "template_path": None,
            "format_key": None,
            "document_view": "",
            "slots": [_slot_overview(slot) for slot in template_slots_by_id.values()],
        }

    return {
        "template_path": template_model.template_path,
        "format_key": template_model.format_key,
        "document_view": _template_document_view(template_model),
        "visual_pages": [str(path) for path in render_template_visual_pages(template_model.template_path, format_key=template_model.format_key)],
        "slots": [_slot_overview(slot) for slot in template_model.slots],
    }


def _slot_overview(slot: TemplateSlot) -> dict[str, Any]:
    return {
        "slot_id": slot.slot_id,
        "raw_name": slot.raw_name,
        "canonical_name": slot.canonical_name,
        "slot_kind": slot.slot_kind,
        "mode": slot.mode,
        "context_before": slot.context_before,
        "context_after": slot.context_after,
        "address": {
            "format_key": slot.address.format_key,
            "address_kind": slot.address.address_kind,
            "paragraph_index": slot.address.paragraph_index,
            "cell_paragraph_index": slot.address.cell_paragraph_index,
            "run_index": slot.address.run_index,
            "run_end_index": slot.address.run_end_index,
            "table_index": slot.address.table_index,
            "row_index": slot.address.row_index,
            "column_index": slot.address.column_index,
            "sheet_name": slot.address.sheet_name,
            "cell_ref": slot.address.cell_ref,
            "page_index": slot.address.page_index,
            "form_field_name": slot.address.form_field_name,
        },
        "placement_constraints": dict(slot.placement_constraints),
    }


def _template_document_view(template_model: TemplateModel) -> str:
    path = Path(template_model.template_path)
    if not path.exists():
        return "\n".join(
            f"{slot.raw_name} | before={slot.context_before or ''} | after={slot.context_after or ''}"
            for slot in template_model.slots
        )
    if template_model.format_key == "docx":
        return _docx_template_view(path)
    if template_model.format_key == "xlsx":
        return _xlsx_template_view(path)
    if template_model.format_key == "pdf":
        return path.read_bytes().decode("latin-1", errors="ignore")
    return path.read_text(encoding="utf-8", errors="ignore") if path.exists() else ""


def _docx_template_view(path: Path) -> str:
    document = Document(path)
    lines: list[str] = []
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if text:
            lines.append(f"P{index + 1}: {text}")
    for table_index, table in enumerate(document.tables):
        lines.append(f"[TABLE {table_index + 1}]")
        for row_index, row in enumerate(table.rows):
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip():
                lines.append(f"R{row_index + 1}: {row_text}")
    return "\n".join(lines)


def _xlsx_template_view(path: Path) -> str:
    workbook = load_workbook(path, data_only=False)
    lines: list[str] = []
    for worksheet in workbook.worksheets:
        lines.append(f"[SHEET {worksheet.title}]")
        for row in worksheet.iter_rows():
            cells = []
            for cell in row:
                value = "" if cell.value is None else str(cell.value).strip()
                if value:
                    cells.append(f"{cell.coordinate}={value}")
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _source_overview(source_model: SourceModel, primary_record: SourceRecord) -> dict[str, Any]:
    overview_records = _source_overview_records(source_model.records, primary_record)
    record_payloads = [_record_overview(record) for record in overview_records]
    diagnostics = [dict(item) for item in source_model.diagnostics[:_SOURCE_OVERVIEW_DIAGNOSTIC_LIMIT]]
    truncated_value_count = sum(
        1
        for record in record_payloads
        for value in record["values"]
        if value.get("truncated") is True
    )
    records_truncated = len(overview_records) < len(source_model.records)
    diagnostics_truncated = len(diagnostics) < len(source_model.diagnostics)
    context_status = {
        "record_count": len(source_model.records),
        "records_included_count": len(overview_records),
        "records_truncated": records_truncated,
        "truncated_value_count": truncated_value_count,
        "diagnostic_count": len(source_model.diagnostics),
        "diagnostics_included_count": len(diagnostics),
        "diagnostics_truncated": diagnostics_truncated,
    }
    return {
        "primary_record_id": primary_record.record_id,
        "record_count": len(source_model.records),
        "records_included_count": len(overview_records),
        "records_truncated": records_truncated,
        "records": record_payloads,
        "image_assets": [_image_asset_overview(asset) for asset in source_model.image_assets],
        "diagnostic_count": len(source_model.diagnostics),
        "diagnostics_truncated": diagnostics_truncated,
        "diagnostics": diagnostics,
        "context_status": context_status,
        "context_diagnostics": _source_context_diagnostics(context_status),
    }


def _source_instance_overview(
    *,
    source_instance_index: SourceInstanceIndex,
    primary_record: SourceRecord,
) -> dict[str, Any]:
    primary_instance = source_instance_index.instance_for_record(primary_record.record_id)
    primary_cluster = None
    if primary_instance is not None:
        try:
            primary_cluster = source_instance_index.cluster_for_instance(primary_instance.instance_id)
        except KeyError:
            primary_cluster = None

    cluster_instances: list[SourceInstance] = []
    cluster_instances_truncated = False
    if primary_cluster is not None:
        cluster_instances = [
            instance
            for instance_id in primary_cluster.instance_ids
            if (instance := source_instance_index.instances_by_id.get(instance_id)) is not None
        ]
        cluster_instances = sorted(
            cluster_instances,
            key=lambda instance: (
                instance.row_index if instance.row_index is not None else 10**9,
                instance.instance_id,
            ),
        )
        cluster_instances_truncated = len(cluster_instances) > _SOURCE_INSTANCE_CLUSTER_INSTANCE_LIMIT
        cluster_instances = cluster_instances[:_SOURCE_INSTANCE_CLUSTER_INSTANCE_LIMIT]

    context_status = {
        "instance_count": len(source_instance_index.instances),
        "cluster_count": len(source_instance_index.clusters),
        "primary_cluster_instance_count": len(primary_cluster.instance_ids) if primary_cluster is not None else 0,
        "primary_cluster_record_count": len(primary_cluster.record_ids) if primary_cluster is not None else 0,
        "primary_cluster_asset_count": len(primary_cluster.asset_ids) if primary_cluster is not None else 0,
        "constraint_scoped_primary": bool(
            primary_instance is not None and primary_instance.instance_kind in {"sheet_row", "table_row"}
        ),
    }

    return {
        "primary_record_id": primary_record.record_id,
        "primary_instance": _instance_overview(primary_instance),
        "primary_cluster": _cluster_overview(primary_cluster),
        "primary_cluster_instances": [_instance_overview(instance) for instance in cluster_instances],
        "primary_cluster_instances_truncated": cluster_instances_truncated,
        "context_status": context_status,
    }


def _instance_overview(instance: SourceInstance | None) -> dict[str, Any] | None:
    if instance is None:
        return None
    return {
        "instance_id": instance.instance_id,
        "instance_kind": instance.instance_kind,
        "document_id": instance.document_id,
        "source_path": instance.source_path,
        "container_id": instance.container_id,
        "row_index": instance.row_index,
        "locator": dict(instance.locator),
        "record_ids": list(instance.record_ids),
        "asset_ids": list(instance.asset_ids),
    }


def _cluster_overview(cluster: SourceCluster | None) -> dict[str, Any] | None:
    if cluster is None:
        return None
    return {
        "cluster_id": cluster.cluster_id,
        "cluster_kind": cluster.cluster_kind,
        "document_id": cluster.document_id,
        "source_path": cluster.source_path,
        "container_id": cluster.container_id,
        "locator": dict(cluster.locator),
        "instance_ids": list(cluster.instance_ids),
        "record_ids": list(cluster.record_ids),
        "asset_ids": list(cluster.asset_ids),
    }


def _source_cluster_plan_overview(cluster_plan: SourceClusterPlanningResult | None) -> dict[str, Any]:
    if cluster_plan is None:
        return {
            "clusters": [],
            "repeat_regions": [],
            "context_status": {"cluster_count": 0, "repeat_region_count": 0},
        }
    return {
        "clusters": [_cluster_plan_overview(cluster) for cluster in cluster_plan.clusters],
        "repeat_regions": [_repeat_region_cluster_plan_overview(region) for region in cluster_plan.repeat_region_plans],
        "context_status": {
            "cluster_count": len(cluster_plan.clusters),
            "repeat_region_count": len(cluster_plan.repeat_region_plans),
        },
    }


def _cluster_plan_overview(cluster: SourceClusterPlan) -> dict[str, Any]:
    return {
        "cluster_id": cluster.cluster_id,
        "cluster_kind": cluster.cluster_kind,
        "instance_ids": list(cluster.instance_ids),
        "record_count": cluster.record_count,
        "asset_count": cluster.asset_count,
        "instances": [
            {
                "instance_id": instance.instance_id,
                "instance_kind": instance.instance_kind,
                "record_ids": list(instance.record_ids),
                "asset_ids": list(instance.asset_ids),
                "row_index": instance.row_index,
            }
            for instance in cluster.instances
        ],
    }


def _repeat_region_cluster_plan_overview(region: RepeatRegionClusterPlan) -> dict[str, Any]:
    return {
        "region_id": region.region_id,
        "region_kind": region.region_kind,
        "slot_ids": list(region.slot_ids),
        "candidate_cluster_ids": list(region.candidate_cluster_ids),
        "instance_count_by_cluster_id": dict(region.instance_count_by_cluster_id),
        "coverage_by_cluster_id": {
            cluster_id: dict(coverage)
            for cluster_id, coverage in region.coverage_by_cluster_id.items()
        },
    }


def _source_document_view(record_payloads: list[dict[str, Any]]) -> str:
    lines: list[str] = []
    current_source_path: str | None = None
    for record in record_payloads:
        source_path = _optional_string(record.get("source_path"))
        if source_path != current_source_path:
            current_source_path = source_path
            lines.append(f"[SOURCE] {source_path or '<unknown>'}")
        locator = _locator_text(record.get("locator"))
        values = _record_values_text(record.get("values"))
        if not values:
            continue
        prefix = f"- {record.get('record_id')}"
        if locator:
            prefix = f"{prefix} {locator}"
        lines.append(f"{prefix}: {values}")
    return "\n".join(lines)


def _full_source_document_view(records: list[SourceRecord]) -> str:
    return _source_document_view([_record_document_view_payload(record) for record in records])


def _record_document_view_payload(record: SourceRecord) -> dict[str, Any]:
    return {
        "record_id": record.record_id,
        "source_path": record.source_path,
        "locator": dict(record.locator),
        "values": [
            {
                "canonical_key": value.canonical_key,
                "raw_label": value.raw_label,
                "value": value.value,
            }
            for value in record.values
        ],
    }


def _source_context_diagnostics(context_status: dict[str, Any]) -> list[dict[str, Any]]:
    diagnostics: list[dict[str, Any]] = []
    if (
        context_status["records_truncated"]
        or context_status["truncated_value_count"] > 0
        or context_status["diagnostics_truncated"]
    ):
        diagnostics.append(
            {
                "code": "content_truncated_for_agent_input",
                "severity": "warning",
                "message": (
                    "Agent source context was shortened for transport; review output confidence and provenance carefully."
                ),
                "record_count": context_status["record_count"],
                "records_included_count": context_status["records_included_count"],
                "truncated_value_count": context_status["truncated_value_count"],
                "diagnostic_count": context_status["diagnostic_count"],
                "diagnostics_included_count": context_status["diagnostics_included_count"],
            }
        )
    return diagnostics


def _locator_text(locator: Any) -> str:
    if not isinstance(locator, dict) or not locator:
        return ""
    parts = [f"{key}={value}" for key, value in locator.items() if value is not None and value != ""]
    return f"({' | '.join(parts)})" if parts else ""


def _record_values_text(values: Any) -> str:
    if not isinstance(values, list):
        return ""
    parts: list[str] = []
    for value in values:
        if not isinstance(value, dict):
            continue
        label = value.get("raw_label") or value.get("canonical_key") or "value"
        text = value.get("value")
        if text in {None, ""}:
            continue
        suffix = " [truncated]" if value.get("truncated") is True else ""
        parts.append(f"{label}: {text}{suffix}")
    return " | ".join(parts)


def _optional_string(value: object) -> str | None:
    if value is None:
        return None
    text = str(value)
    return text if text else None


def _source_overview_records(records: list[SourceRecord], primary_record: SourceRecord) -> list[SourceRecord]:
    if len(records) <= _SOURCE_OVERVIEW_RECORD_LIMIT:
        return list(records)

    indexed_records = list(enumerate(records))
    sorted_records = sorted(
        indexed_records,
        key=lambda item: (
            -_record_overview_relevance_score(item[1], primary_record),
            item[0],
        ),
    )
    selected = {record.record_id: record for _, record in sorted_records[:_SOURCE_OVERVIEW_RECORD_LIMIT]}
    if primary_record.record_id not in selected:
        selected.pop(next(iter(selected)))
        selected[primary_record.record_id] = primary_record
    return [record for record in records if record.record_id in selected]


def _record_overview_relevance_score(record: SourceRecord, primary_record: SourceRecord) -> float:
    if record.record_id == primary_record.record_id:
        return 1000.0

    score = 0.0
    if record.document_id == primary_record.document_id:
        score += 4.0
    if record.locator.get("row_index") is not None:
        score += 1.5
    if record.evidence_id is not None:
        score += 1.0

    primary_page = primary_record.locator.get("page_anchor")
    if primary_page is not None and primary_page == record.locator.get("page_anchor"):
        score += 2.0

    primary_table = primary_record.locator.get("table_id") or primary_record.table_id
    record_table = record.locator.get("table_id") or record.table_id
    if primary_table is not None and primary_table == record_table:
        score += 2.0

    return score


def _record_overview(record: SourceRecord) -> dict[str, Any]:
    return {
        "record_id": record.record_id,
        "source_path": record.source_path,
        "evidence_id": record.evidence_id,
        "table_id": record.table_id,
        "cell_range_id": record.cell_range_id,
        "locator": dict(record.locator),
        "values": [
            {
                "canonical_key": value.canonical_key,
                "raw_label": value.raw_label,
                "value": _truncate_text(value.value, _SOURCE_OVERVIEW_VALUE_MAX_CHARS),
                "value_length": len(value.value),
                "truncated": len(value.value) > _SOURCE_OVERVIEW_VALUE_MAX_CHARS,
                "value_type": infer_source_value_type(value),
            }
            for value in record.values
        ],
    }


def _image_asset_overview(asset) -> dict[str, Any]:
    return {
        "asset_id": asset.asset_id,
        "document_id": asset.document_id,
        "source_path": asset.source_path,
        "source_ref": asset.source_ref,
        "asset_path": asset.asset_path,
        "asset_source": asset.asset_source,
        "page_anchor": asset.page_anchor,
        "bbox": asset.bbox,
        "block_id": asset.block_id,
        "order_index": asset.order_index,
        "sheet_name": asset.sheet_name,
        "cell_anchor": asset.cell_anchor,
        "row_index": asset.row_index,
        "column_index": asset.column_index,
        "related_evidence_ids": list(asset.related_evidence_ids),
        "nearby_text_before": _truncate_text_list(asset.nearby_text_before, _IMAGE_CONTEXT_VALUE_MAX_CHARS),
        "nearby_text_after": _truncate_text_list(asset.nearby_text_after, _IMAGE_CONTEXT_VALUE_MAX_CHARS),
        "nearby_evidence_ids": list(asset.nearby_evidence_ids),
    }


def _validate_agent_payload(payload: dict[str, Any]) -> None:
    schema_path = Path("contracts/doctransai/agent_mapping_decision_v1.schema.json")
    schema = json.loads(schema_path.read_text(encoding="utf-8"))
    validator = Draft202012Validator(schema)
    errors = sorted(validator.iter_errors(payload), key=lambda error: list(error.path))
    if errors:
        raise AgentResolutionError(errors[0].message)


def _parse_and_validate_payload(
    raw_output: str,
    slot_packets: list[AgentSlotPacket],
) -> tuple[dict[str, Any], list[str]]:
    normalized_raw_output = _strip_json_code_fence(raw_output)
    try:
        payload = json.loads(normalized_raw_output)
    except json.JSONDecodeError as exc:
        return {}, [f"Agent output is not valid JSON: {exc}"]

    payload = _normalize_agent_payload(payload)
    errors: list[str] = []

    schema_path = Path("contracts/doctransai/agent_mapping_decision_v1.schema.json")
    schema = json.loads(schema_path.read_text(encoding="utf-8"))
    validator = Draft202012Validator(schema)
    schema_errors = sorted(validator.iter_errors(payload), key=lambda error: list(error.path))
    errors.extend(error.message for error in schema_errors)

    candidate_ids_by_slot = {
        packet.slot_id: {candidate.candidate_id for candidate in packet.candidates}
        for packet in slot_packets
    }
    packets_by_slot = {packet.slot_id: packet for packet in slot_packets}
    expected_slot_ids = {packet.slot_id for packet in slot_packets}
    seen_slot_ids: set[str] = set()
    for item in payload.get("decisions", []):
        if not isinstance(item, dict):
            continue
        slot_id = str(item.get("slot_id", ""))
        decision = str(item.get("decision", ""))
        candidate_id = item.get("candidate_id")
        generated_value = item.get("generated_value")

        if slot_id in seen_slot_ids:
            errors.append(f"Duplicate decision for slot `{slot_id}`.")
        seen_slot_ids.add(slot_id)
        if slot_id not in expected_slot_ids:
            errors.append(f"Unexpected slot_id `{slot_id}` in decisions.")

        if decision == "candidate":
            available_candidate_ids = candidate_ids_by_slot.get(slot_id, set())
            selected_candidate_ids = _candidate_ids_from_payload_item(item)
            if not selected_candidate_ids and candidate_id not in {None, ""}:
                selected_candidate_ids = [str(candidate_id)]
            normalized_candidate_ids: list[str] = []
            for selected_candidate_id in selected_candidate_ids:
                normalized_candidate_id = selected_candidate_id
                if normalized_candidate_id not in available_candidate_ids:
                    alias = _normalize_candidate_id_alias(normalized_candidate_id, available_candidate_ids)
                    if alias is not None:
                        normalized_candidate_id = alias
                if normalized_candidate_id not in available_candidate_ids:
                    errors.append(f"Unknown candidate `{selected_candidate_id}` for slot `{slot_id}`.")
                else:
                    normalized_candidate_ids.append(normalized_candidate_id)

            if not normalized_candidate_ids:
                errors.append(f"Slot `{slot_id}` has decision=candidate but no valid candidate_id.")
            else:
                item["candidate_id"] = normalized_candidate_ids[0]
                item["candidate_ids"] = normalized_candidate_ids
                if len(normalized_candidate_ids) > 1 and not _slot_allows_multiple_candidate_ids(packets_by_slot.get(slot_id)):
                    errors.append(f"Slot `{slot_id}` selected multiple candidates but the slot is not a repeatable image area.")
            if generated_value not in {None, ""}:
                errors.append(f"Slot `{slot_id}` has decision=candidate but generated_value is not null.")
        elif decision == "generated":
            if generated_value in {None, ""}:
                errors.append(f"Slot `{slot_id}` has decision=generated but generated_value is empty.")
            if candidate_id not in {None, ""}:
                errors.append(f"Slot `{slot_id}` has decision=generated but candidate_id is not null.")
            if _candidate_ids_from_payload_item(item):
                errors.append(f"Slot `{slot_id}` has decision=generated but candidate_ids is not empty.")
        elif decision == "unresolved":
            if candidate_id not in {None, ""}:
                errors.append(f"Slot `{slot_id}` has decision=unresolved but candidate_id is not null.")
            if generated_value not in {None, ""}:
                errors.append(f"Slot `{slot_id}` has decision=unresolved but generated_value is not null.")
            if _candidate_ids_from_payload_item(item):
                errors.append(f"Slot `{slot_id}` has decision=unresolved but candidate_ids is not empty.")

    missing_slot_ids = expected_slot_ids - seen_slot_ids
    for slot_id in sorted(missing_slot_ids):
        errors.append(f"Missing decision for slot `{slot_id}`.")

    return payload, errors


def _normalize_candidate_id_alias(candidate_id: object, candidate_ids: set[str]) -> str | None:
    if not isinstance(candidate_id, str) or not candidate_id:
        return None
    matches = [item for item in candidate_ids if item.startswith(f"{candidate_id}:")]
    if len(matches) == 1:
        return matches[0]
    return None


def _candidate_ids_from_payload_item(item: dict[str, Any]) -> list[str]:
    raw_candidate_ids = item.get("candidate_ids")
    if isinstance(raw_candidate_ids, list):
        return [str(candidate_id) for candidate_id in raw_candidate_ids if str(candidate_id)]
    return []


def _slot_allows_multiple_candidate_ids(packet: AgentSlotPacket | None) -> bool:
    if packet is None or packet.slot_kind != "image":
        return False
    placement_constraints = packet.context.get("placement_constraints")
    if not isinstance(placement_constraints, dict):
        return False
    return placement_constraints.get("multi_value_allowed") is True


def _strip_json_code_fence(raw_output: str) -> str:
    text = raw_output.strip()
    if not text.startswith("```"):
        return text
    lines = text.splitlines()
    if not lines:
        return text
    if lines[0].startswith("```"):
        lines = lines[1:]
    if lines and lines[-1].strip() == "```":
        lines = lines[:-1]
    return "\n".join(lines).strip()


def _normalize_agent_payload(payload: Any) -> dict[str, Any]:
    if isinstance(payload, dict):
        if "mappings" in payload and "decisions" not in payload and "resolutions" not in payload:
            mappings = payload.get("mappings")
            if isinstance(mappings, list):
                payload = {
                    "schema_version": payload.get("schema_version", AGENT_MAPPING_SCHEMA_VERSION),
                    "decisions": mappings,
                }
        if "resolutions" in payload and "decisions" not in payload:
            resolutions = payload.get("resolutions")
            if isinstance(resolutions, list):
                decisions: list[dict[str, Any]] = []
                for item in resolutions:
                    if not isinstance(item, dict):
                        continue
                    slot_id = item.get("slot_id")
                    if not isinstance(slot_id, str):
                        continue
                    candidate_id = item.get("candidate_id")
                    resolved = item.get("resolved")
                    if resolved is True and isinstance(candidate_id, str):
                        decisions.append(
                            {
                                "slot_id": slot_id,
                                "decision": "candidate",
                                "candidate_id": candidate_id,
                                "generated_value": None,
                                "confidence": _coerce_confidence(item.get("confidence"), default=0.75),
                                "reason": str(item.get("reason", "normalized_resolution_candidate_response")),
                            }
                        )
                    else:
                        decisions.append(
                            {
                                "slot_id": slot_id,
                                "decision": "unresolved",
                                "candidate_id": None,
                                "generated_value": None,
                                "confidence": _coerce_confidence(item.get("confidence"), default=0.0),
                                "reason": str(item.get("reason", "normalized_resolution_unresolved_response")),
                            }
                        )
                return {
                    "schema_version": AGENT_MAPPING_SCHEMA_VERSION,
                    "decisions": decisions,
                }
        if "decisions" in payload and isinstance(payload.get("decisions"), list):
            decisions: list[dict[str, Any]] = []
            for item in payload["decisions"]:
                if not isinstance(item, dict):
                    continue
                decision = item.get("decision", item.get("resolution", "unresolved"))
                normalized_decision = _normalize_decision_value(decision, candidate_id=item.get("candidate_id", item.get("resolved_candidate_id")))
                candidate_id = item.get("candidate_id", item.get("resolved_candidate_id"))
                candidate_ids = _candidate_ids_from_payload_item(item)
                generated_value = item.get("generated_value")
                normalized_item = {
                    "slot_id": str(item["slot_id"]),
                    "decision": normalized_decision,
                    "candidate_id": str(candidate_id) if candidate_id is not None else None,
                    "generated_value": str(generated_value) if generated_value is not None else None,
                    "confidence": _coerce_confidence(
                        item.get("confidence"),
                        default=0.0 if normalized_decision == "unresolved" else 0.75,
                    ),
                    "reason": str(item.get("reason", "normalized_decision_response")),
                }
                if candidate_ids:
                    normalized_item["candidate_ids"] = candidate_ids
                decisions.append(normalized_item)
            return {
                "schema_version": str(payload.get("schema_version", AGENT_MAPPING_SCHEMA_VERSION)),
                "decisions": decisions,
            }
        return payload

    if isinstance(payload, list):
        decisions: list[dict[str, Any]] = []
        for item in payload:
            if not isinstance(item, dict):
                continue
            slot_id = item.get("slot_id")
            if not isinstance(slot_id, str):
                continue
            resolved_candidate_id = item.get("resolved_candidate_id")
            if isinstance(resolved_candidate_id, str) and resolved_candidate_id != "unresolved":
                decisions.append(
                    {
                        "slot_id": slot_id,
                        "decision": "candidate",
                        "candidate_id": resolved_candidate_id,
                        "generated_value": None,
                        "confidence": float(item.get("confidence", 0.75)),
                        "reason": str(item.get("reason", "normalized_legacy_candidate_response")),
                    }
                )
            else:
                decisions.append(
                    {
                        "slot_id": slot_id,
                        "decision": "unresolved",
                        "candidate_id": None,
                        "generated_value": None,
                        "confidence": float(item.get("confidence", 0.0)),
                        "reason": str(item.get("reason", "normalized_legacy_unresolved_response")),
                    }
                )
        return {
            "schema_version": AGENT_MAPPING_SCHEMA_VERSION,
            "decisions": decisions,
        }

    raise AgentResolutionError("Agent output must be a JSON object or a normalizable legacy list.")


def _normalize_decision_value(value: object, *, candidate_id: object) -> str:
    text = str(value or "").strip().lower()
    if text in {"candidate", "generated", "unresolved"}:
        return text
    if text in {"resolved", "selected", "matched", "filled"}:
        return "candidate" if candidate_id is not None else "unresolved"
    return "unresolved"


def _coerce_confidence(value: object, *, default: float) -> float:
    if value is None:
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _text_candidate_packets(
    *,
    source_model: SourceModel,
    primary_record: SourceRecord,
    expected_value_type: str,
) -> list[AgentCandidatePacket]:
    packets: list[AgentCandidatePacket] = []
    seen: set[tuple[str, str, str, tuple[tuple[Any, ...], ...]]] = set()
    record_index_by_id = {record.record_id: index for index, record in enumerate(source_model.records)}
    ordered_records = [primary_record, *[item for item in source_model.records if item.record_id != primary_record.record_id]]
    for record in ordered_records:
        relation = "same_record_row" if record.record_id == primary_record.record_id else "document_context"
        confidence_hint = 0.75 if relation == "same_record_row" else 0.55
        for index, value in enumerate(record.values):
            if not value.value.strip():
                continue
            if _is_broad_generic_text_value(value):
                continue
            label = value.raw_label or value.canonical_key
            value_type = infer_source_value_type(value)
            if expected_value_type != "text" and value_type != expected_value_type:
                continue
            key = (label, value.value, record.record_id, _provenance_refs_key(value.provenance))
            if key in seen:
                continue
            seen.add(key)
            packets.append(
                AgentCandidatePacket(
                    candidate_id=f"{record.record_id}:{index}",
                    candidate_kind="text",
                    value_label=label,
                    value_text=value.value,
                    value_type=value_type,
                    relation=relation,
                    confidence_hint=confidence_hint,
                    locator=_source_value_locator(value, record),
                    source_path=record.source_path,
                    provenance_refs=[_provenance_ref_payload(ref) for ref in value.provenance],
                    source_object_refs=[
                        _record_source_object_ref(
                            record=record,
                            index=record_index_by_id.get(record.record_id),
                        )
                    ],
                )
            )
    return packets


def _source_value_locator(value: Any, record: SourceRecord) -> dict[str, Any]:
    locator = dict(record.locator)
    for ref in value.provenance:
        if ref.locator:
            locator.update(ref.locator)
            break
    return locator


def _provenance_refs_key(provenance_refs: list[ProvenanceRef]) -> tuple[tuple[Any, ...], ...]:
    return tuple(
        (
            ref.source_kind,
            ref.document_id,
            ref.evidence_id,
            ref.asset_id,
            ref.object_id,
            _freeze_mapping(ref.locator),
            ref.source_path,
        )
        for ref in provenance_refs
    )


def _freeze_mapping(value: Any) -> Any:
    if isinstance(value, dict):
        return tuple(sorted((key, _freeze_mapping(item)) for key, item in value.items()))
    if isinstance(value, list):
        return tuple(_freeze_mapping(item) for item in value)
    return value


def _is_broad_generic_text_value(value: Any) -> bool:
    return value.canonical_key in _GENERIC_TEXT_KEYS and len(value.value.strip()) > 80


def _image_candidate_packets(
    *,
    source_model: SourceModel,
    primary_record: SourceRecord,
    slot: TemplateSlot,
) -> list[AgentCandidatePacket]:
    packets: list[AgentCandidatePacket] = []
    slot_name = f"{slot.raw_name} {slot.canonical_name}"
    asset_index_by_id = {asset.asset_id: index for index, asset in enumerate(source_model.image_assets)}
    for asset in source_model.image_assets:
        relation = "document_context"
        confidence_hint = 0.55
        primary_sheet_name = primary_record.locator.get("sheet_name")
        primary_row_index = primary_record.locator.get("row_index")
        if (
            primary_sheet_name is not None
            and primary_row_index is not None
            and primary_sheet_name == asset.sheet_name
            and primary_row_index == asset.row_index
        ):
            relation = "same_record_row"
            confidence_hint = 0.78
        elif primary_record.evidence_id is not None and primary_record.evidence_id in asset.nearby_evidence_ids:
            relation = "near_primary_record_context"
            confidence_hint = 0.72
        packets.append(
            AgentCandidatePacket(
                candidate_id=asset.asset_id,
                candidate_kind="image",
                value_label=slot_name,
                value_text=_image_asset_candidate_text(asset),
                value_type="image",
                relation=relation,
                confidence_hint=confidence_hint,
                locator={
                    "document_id": asset.document_id,
                    "source_ref": asset.source_ref,
                    "asset_source": asset.asset_source,
                    "page_anchor": asset.page_anchor,
                    "bbox": asset.bbox,
                    "block_id": asset.block_id,
                    "order_index": asset.order_index,
                    "sheet_name": asset.sheet_name,
                    "row_index": asset.row_index,
                    "column_index": asset.column_index,
                    "cell_anchor": asset.cell_anchor,
                    "nearby_text_before": _truncate_text_list(asset.nearby_text_before, _IMAGE_CONTEXT_VALUE_MAX_CHARS),
                    "nearby_text_after": _truncate_text_list(asset.nearby_text_after, _IMAGE_CONTEXT_VALUE_MAX_CHARS),
                    "nearby_evidence_ids": list(asset.nearby_evidence_ids),
                    "related_evidence_ids": list(asset.related_evidence_ids),
                },
                source_path=asset.asset_path,
                provenance_refs=[
                    {
                        "source_kind": "image_asset",
                        "document_id": asset.document_id,
                        "asset_id": asset.asset_id,
                        "object_id": asset.asset_id,
                        "locator": {
                            "page_anchor": asset.page_anchor,
                            "bbox": asset.bbox,
                            "sheet_name": asset.sheet_name,
                            "row_index": asset.row_index,
                            "column_index": asset.column_index,
                            "cell_anchor": asset.cell_anchor,
                        },
                        "source_path": asset.asset_path,
                    }
                ],
                source_object_refs=[
                    _image_source_object_ref(
                        asset_id=asset.asset_id,
                        index=asset_index_by_id.get(asset.asset_id),
                    )
                ],
            )
        )

    return packets


def _record_source_object_ref(*, record: SourceRecord, index: int | None) -> dict[str, Any]:
    object_id = record.record_id
    return {
        "object_key": f"record:{object_id}",
        "object_kind": "record",
        "object_id": object_id,
        "collection": "records",
        "index": index,
    }


def _image_source_object_ref(*, asset_id: str, index: int | None) -> dict[str, Any]:
    return {
        "object_key": f"image_asset:{asset_id}",
        "object_kind": "image_asset",
        "object_id": asset_id,
        "collection": "image_assets",
        "index": index,
    }


def _provenance_ref_payload(ref: ProvenanceRef) -> dict[str, Any]:
    return {
        "source_kind": ref.source_kind,
        "document_id": ref.document_id,
        "evidence_id": ref.evidence_id,
        "asset_id": ref.asset_id,
        "object_id": ref.object_id,
        "locator": dict(ref.locator),
        "source_path": ref.source_path,
    }


def _image_asset_candidate_text(asset) -> str:
    pieces = [asset.source_ref or asset.asset_id]
    if asset.nearby_text_before:
        pieces.append(
            "nearby_text_before="
            + " / ".join(_truncate_text_list(asset.nearby_text_before, _IMAGE_CONTEXT_VALUE_MAX_CHARS))
        )
    if asset.nearby_text_after:
        pieces.append(
            "nearby_text_after="
            + " / ".join(_truncate_text_list(asset.nearby_text_after, _IMAGE_CONTEXT_VALUE_MAX_CHARS))
        )
    if asset.page_anchor:
        pieces.append(f"page_anchor={asset.page_anchor}")
    if asset.order_index is not None:
        pieces.append(f"order_index={asset.order_index}")
    return " | ".join(pieces)


def _truncate_text(value: str, max_chars: int) -> str:
    if len(value) <= max_chars:
        return value
    return value[: max_chars - 3].rstrip() + "..."


def _truncate_text_list(values: list[str] | tuple[str, ...], max_chars: int) -> list[str]:
    return [_truncate_text(str(value), max_chars) for value in values]
