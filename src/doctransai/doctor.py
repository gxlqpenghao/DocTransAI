"""Local dependency checks; never print credentials or contact model providers."""
from __future__ import annotations
import argparse
import importlib.metadata
import json
from pathlib import Path
import shutil
import subprocess
import sys
from .config import RuntimeConfig
from .render_document import resolve_soffice
from .settings import env_values, setting


def checks(env_file=None, require_credentials=False):
    if env_file and not Path(env_file).is_file():
        return [{"name":"Configuration file","ok":False,"detail":"Copy .env.example to .env and fill service URLs/keys."}]
    result=[]
    def add(name,ok,detail):result.append({'name':name,'ok':bool(ok),'detail':detail})
    add('Python',sys.version_info[:2]==(3,12),sys.version.split()[0])
    try:
        import pypdfium2
        add('PDFium',True,importlib.metadata.version('pypdfium2'))
    except ImportError:add('PDFium',False,'Run setup or uv sync --locked')
    try:
        cfg=RuntimeConfig.from_environment(env_file=env_file)
        add('Renderer',cfg.render_script.is_file() and bool(shutil.which(cfg.render_python)),str(cfg.render_script))
        office=resolve_soffice(cfg.soffice)
        r=subprocess.run([office,'--version'],capture_output=True,timeout=15)
        add('LibreOffice',r.returncode==0,(r.stdout+r.stderr).decode('utf-8','replace').strip()[:160] or office)
    except Exception as e:add('LibreOffice / configuration',False,str(e))
    if require_credentials:
        values=env_values(env_file)
        for name in ('DOCTRANSAI_DATAPROXY_BASE_URL','DOCTRANSAI_DATAPROXY_API_KEY','DOCTRANSAI_LLM_BASE_URL','DOCTRANSAI_LLM_API_KEY'):
            v=setting(name,values,'')
            ok=bool(v) and not any(t in v for t in ('replace-with','example.com','example/v1','<'))
            add(name,ok,'configured' if ok else 'Fill this item in the backend .env file')
    return result


def render_check(output, env_file=None):
    from docx import Document
    from docx.shared import Pt
    from docx.oxml.ns import qn
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from .documents.preview import Previewer
    output=Path(output).resolve();output.mkdir(parents=True,exist_ok=False)
    cfg=RuntimeConfig.from_environment(env_file=env_file)
    doc=Document();s=doc.styles['Normal'];s.font.name='Microsoft YaHei' if sys.platform=='win32' else 'PingFang SC' if sys.platform=='darwin' else 'Noto Sans CJK SC';s.font.size=Pt(12);s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),s.font.name)
    doc.add_paragraph('中文文档渲染自检');t=doc.add_table(rows=2,cols=2);t.style='Table Grid';t.cell(0,0).text='检查项目';t.cell(0,1).text='检查结果';t.cell(1,0).text='字体与表格';t.cell(1,1).text='中文应完整可见';doc.save(output/'中文 自检.docx')
    book=Workbook();sheet=book.active;sheet.title='中文测试';sheet.append(['检查项目','检查结果']);sheet.append(['字体与表格','中文应完整可见']);sheet.column_dimensions['A'].width=25;sheet.column_dimensions['B'].width=30;sheet.print_area='A1:B2'
    for row in sheet:
        for cell in row:cell.font=Font(name=s.font.name,size=12)
    book.save(output/'中文 自检.xlsx');pre=Previewer(output/'previews',cfg)
    records=[]
    for suffix in ('docx','xlsx'):
        source=output/f'中文 自检.{suffix}';pages=pre.render(source)
        import pypdfium2
        text=''
        with pypdfium2.PdfDocument(str(pages[0].parent/(source.stem+'.pdf'))) as pdf:
            for index in range(len(pdf)):
                page=pdf[index]
                try:
                    textpage=page.get_textpage()
                    try:text+=textpage.get_text_range()
                    finally:textpage.close()
                finally:page.close()
        if '中文' not in text or '检查' not in text:
            raise ValueError('Chinese text missing from rendered PDF; install/configure CJK fonts and inspect PNGs')
        records.append({'format':suffix,'pages':[str(p) for p in pages]})
    (output/'render-check.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
    return records


def main(argv=None):
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--env-file',type=Path);p.add_argument('--require-credentials',action='store_true');p.add_argument('--render-check',type=Path,help='New directory for DOCX/XLSX visual self-check');p.add_argument('--json',action='store_true');a=p.parse_args(argv)
    rows=checks(a.env_file,a.require_credentials)
    if a.render_check and all(x['ok'] for x in rows):
        try:render_check(a.render_check,a.env_file);rows.append({'name':'Render smoke test','ok':True,'detail':str(a.render_check)+' (inspect PNGs visually; text extraction alone is not visual proof)'})
        except Exception as e:rows.append({'name':'Render smoke test','ok':False,'detail':str(e)})
    print(json.dumps(rows,ensure_ascii=False,indent=2) if a.json else '\n'.join(('OK  ' if r['ok'] else 'FAIL ')+r['name']+': '+r['detail'] for r in rows))
    return 0 if all(r['ok'] for r in rows) else 1
if __name__=='__main__':raise SystemExit(main())
