from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from doctransai.formats.registry import detect_format
from doctransai.models.template import (
    HeaderDrivenRegion,
    RepeatRegion,
    SlotAddress,
    TemplateDiagnostic,
    TemplateModel,
    TemplateSlot,
)

_PLACEHOLDER_PATTERN = re.compile(r"\{([^{}]+)\}")
_DOCX_CHOICE_PATTERN = re.compile(r"□([^□]+)")
_PDF_DICT_PATTERN = re.compile(r"<<(?P<body>.*?)>>", re.DOTALL)
_PDF_FIELD_NAME_PATTERN = re.compile(r"/T\s*\((?P<name>(?:\\.|[^\\)])*)\)")
_PDF_FIELD_TYPE_PATTERN = re.compile(r"/FT\s*/(?P<field_type>[A-Za-z]+)")
_PHOTO_HEADER_HINTS = ("照片", "图片", "头像", "photo", "image", "img", "pic")
_PHOTO_TEXT_HINTS = ("说明", "描述", "备注", "caption", "description", "desc")
_MEDIA_DESCRIPTION_HINTS = ("描述", "说明", "caption", "content")


def analyze_template(template_path: str | Path) -> TemplateModel:
    path = Path(template_path).expanduser().resolve()
    format_spec = detect_format(path)
    analyzer_key = format_spec.template_analyzer_key
    if analyzer_key is None:
        raise ValueError(f"Format does not support template analysis: {path.suffix.lower() or '<none>'}")

    if format_spec.key == "docx":
        return _analyze_docx(path, analyzer_key)
    if format_spec.key == "xlsx":
        return _analyze_xlsx(path, analyzer_key)
    if format_spec.key == "pdf":
        return _analyze_pdf(path, analyzer_key)

    raise ValueError(f"Unsupported template analysis format: {format_spec.key}")


def _analyze_docx(path: Path, analyzer_key: str) -> TemplateModel:
    document = Document(path)
    slots: list[TemplateSlot] = []
    seen_merged_cell_slots: set[tuple[int, str, int, str]] = set()
    occupied_cell_tokens: set[tuple[int, str]] = set()

    for paragraph_index, paragraph in enumerate(document.paragraphs):
        address = SlotAddress(
            format_key="docx",
            address_kind="paragraph",
            paragraph_index=paragraph_index,
        )
        slots.extend(_extract_docx_writable_slots(paragraph=paragraph, address=address, slot_offset=len(slots)))

    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            for column_index, cell in enumerate(row.cells):
                cell_token = _docx_cell_token(cell, table_index=table_index, row_index=row_index, column_index=column_index)
                for paragraph_index_in_cell, paragraph in enumerate(cell.paragraphs):
                    address = SlotAddress(
                        format_key="docx",
                        address_kind="table_cell",
                        cell_paragraph_index=paragraph_index_in_cell,
                        table_index=table_index,
                        row_index=row_index,
                        column_index=column_index,
                    )
                    extracted_slots = _extract_docx_writable_slots(
                        paragraph=paragraph,
                        address=address,
                        slot_offset=len(slots),
                        slot_prefix=f"docx-table-{table_index}",
                        blank_slot_prefix=f"docx-table-{table_index}-blank",
                    )
                    for slot in extracted_slots:
                        if slot.context_before is None:
                            slot.context_before = _neighbor_cell_text(row, column_index, direction=-1)
                        if slot.context_after is None:
                            slot.context_after = _neighbor_cell_text(row, column_index, direction=1)
                        dedupe_key = _docx_cell_slot_dedupe_key(
                            table_index=table_index,
                            cell_token=cell_token,
                            paragraph_index_in_cell=paragraph_index_in_cell,
                            slot=slot,
                        )
                        if dedupe_key in seen_merged_cell_slots:
                            continue
                        seen_merged_cell_slots.add(dedupe_key)
                        occupied_cell_tokens.add((table_index, cell_token))
                        slots.append(slot)
                if (table_index, cell_token) not in occupied_cell_tokens:
                    blank_slot = _extract_docx_blank_table_cell_slot(
                        table=table,
                        table_index=table_index,
                        row_index=row_index,
                        column_index=column_index,
                        slot_offset=len(slots),
                    )
                    if blank_slot is not None:
                        occupied_cell_tokens.add((table_index, cell_token))
                        slots.append(blank_slot)
                if (table_index, cell_token) not in occupied_cell_tokens:
                    instruction_slot = _extract_docx_instruction_table_cell_slot(
                        table=table,
                        table_index=table_index,
                        row_index=row_index,
                        column_index=column_index,
                        slot_offset=len(slots),
                    )
                    if instruction_slot is not None:
                        occupied_cell_tokens.add((table_index, cell_token))
                        slots.append(instruction_slot)
                if (table_index, cell_token) not in occupied_cell_tokens:
                    choice_slot = _extract_docx_choice_table_cell_slot(
                        table=table,
                        table_index=table_index,
                        row_index=row_index,
                        column_index=column_index,
                        slot_offset=len(slots),
                    )
                    if choice_slot is not None:
                        occupied_cell_tokens.add((table_index, cell_token))
                        slots.append(choice_slot)

    slots.extend(_extract_docx_header_footer_slots(document=document, slot_offset=len(slots)))

    diagnostics: list[TemplateDiagnostic] = []
    if not slots:
        diagnostics.append(
            TemplateDiagnostic(
                severity="info",
                code="docx_no_placeholders",
                message="No {field} placeholders were found in the DOCX template.",
            )
        )

    _annotate_docx_media_group_slots(slots)
    repeat_regions = _docx_paragraph_repeat_regions(slots)

    return TemplateModel(
        template_id=_template_id_for(path),
        template_path=str(path),
        format_key="docx",
        analyzer_key=analyzer_key,
        slots=slots,
        repeat_regions=repeat_regions,
        diagnostics=diagnostics,
    )


def _extract_docx_header_footer_slots(*, document, slot_offset: int) -> list[TemplateSlot]:
    slots: list[TemplateSlot] = []
    for section_index, section in enumerate(document.sections):
        for story_name, story in (("header", section.header), ("footer", section.footer)):
            if section_index > 0 and getattr(story, "is_linked_to_previous", False):
                continue
            story_prefix = f"docx-{story_name}-{section_index}"
            for paragraph_index, paragraph in enumerate(story.paragraphs):
                address = SlotAddress(
                    format_key="docx",
                    address_kind=f"{story_name}_paragraph",
                    paragraph_index=paragraph_index,
                    page_index=section_index,
                )
                slots.extend(
                    _extract_docx_writable_slots(
                        paragraph=paragraph,
                        address=address,
                        slot_offset=slot_offset + len(slots),
                        slot_prefix=f"{story_prefix}-paragraph",
                        blank_slot_prefix=f"{story_prefix}-blank",
                    )
                )

            for table_index, table in enumerate(story.tables):
                for row_index, row in enumerate(table.rows):
                    for column_index, cell in enumerate(row.cells):
                        for paragraph_index_in_cell, paragraph in enumerate(cell.paragraphs):
                            address = SlotAddress(
                                format_key="docx",
                                address_kind=f"{story_name}_table_cell",
                                cell_paragraph_index=paragraph_index_in_cell,
                                table_index=table_index,
                                row_index=row_index,
                                column_index=column_index,
                                page_index=section_index,
                            )
                            extracted_slots = _extract_docx_writable_slots(
                                paragraph=paragraph,
                                address=address,
                                slot_offset=slot_offset + len(slots),
                                slot_prefix=f"{story_prefix}-table-{table_index}",
                                blank_slot_prefix=f"{story_prefix}-table-{table_index}-blank",
                            )
                            for slot in extracted_slots:
                                if slot.context_before is None:
                                    slot.context_before = _neighbor_cell_text(row, column_index, direction=-1)
                                if slot.context_after is None:
                                    slot.context_after = _neighbor_cell_text(row, column_index, direction=1)
                                slots.append(slot)
    return slots


def _extract_docx_writable_slots(
    *,
    paragraph,
    address: SlotAddress,
    slot_offset: int,
    slot_prefix: str = "docx-paragraph",
    blank_slot_prefix: str = "docx-paragraph-blank",
) -> list[TemplateSlot]:
    """Return only explicit placeholders and styled fill blanks from DOCX text."""
    slots = _extract_placeholder_slots(
        text=paragraph.text,
        slot_prefix=slot_prefix,
        slot_offset=slot_offset,
        address=address,
        paragraph=paragraph,
    )
    slots.extend(
        _extract_implicit_blank_slots(
            paragraph=paragraph,
            slot_prefix=blank_slot_prefix,
            slot_offset=slot_offset + len(slots),
            address=address,
        )
    )
    return slots


def _annotate_docx_media_group_slots(slots: list[TemplateSlot]) -> None:
    slots_by_paragraph: dict[tuple[object, ...], list[TemplateSlot]] = {}
    slots_by_table_row: dict[tuple[int, int], list[TemplateSlot]] = {}
    for slot in slots:
        if slot.address.format_key != "docx":
            continue
        key = (
            slot.address.address_kind,
            slot.address.paragraph_index,
            slot.address.table_index,
            slot.address.row_index,
            slot.address.column_index,
            slot.address.cell_paragraph_index,
        )
        slots_by_paragraph.setdefault(key, []).append(slot)
        if slot.address.table_index is not None and slot.address.row_index is not None:
            slots_by_table_row.setdefault((slot.address.table_index, slot.address.row_index), []).append(slot)

    for paragraph_slots in slots_by_paragraph.values():
        image_slots = [slot for slot in paragraph_slots if slot.slot_kind == "image"]
        caption_slots = [
            slot
            for slot in paragraph_slots
            if slot.slot_kind == "text" and _is_media_caption_slot_name(slot.raw_name)
        ]
        if not image_slots or not caption_slots:
            continue

        image_slot = image_slots[0]
        caption_slot = caption_slots[0]
        image_slot.placement_constraints["multi_value_allowed"] = True
        image_slot.placement_constraints["media_group_role"] = "image"
        image_slot.placement_constraints["paired_caption_slot_id"] = caption_slot.slot_id
        caption_slot.placement_constraints["multi_value_allowed"] = True
        caption_slot.placement_constraints["media_group_role"] = "caption"
        caption_slot.placement_constraints["paired_image_slot_id"] = image_slot.slot_id

    for row_slots in slots_by_table_row.values():
        image_slots = [slot for slot in row_slots if slot.slot_kind == "image"]
        text_slots = [slot for slot in row_slots if slot.slot_kind == "text"]
        for image_slot in image_slots:
            partner = _paired_row_text_slot_for_image(image_slot, text_slots)
            if partner is None:
                continue
            image_slot.placement_constraints.setdefault("multi_value_allowed", True)
            image_slot.placement_constraints.setdefault("media_group_role", "image")
            image_slot.placement_constraints.setdefault("paired_text_slot_id", partner.slot_id)
            partner.placement_constraints.setdefault("multi_value_allowed", True)
            partner.placement_constraints.setdefault("media_group_role", "description")
            partner.placement_constraints.setdefault("paired_image_slot_id", image_slot.slot_id)


def _paired_row_text_slot_for_image(image_slot: TemplateSlot, text_slots: list[TemplateSlot]) -> TemplateSlot | None:
    image_base = _media_slot_base_name(image_slot.raw_name)
    if not image_base:
        return None
    for text_slot in text_slots:
        if not _is_media_description_slot(text_slot):
            continue
        if _media_slot_base_name(text_slot.raw_name) == image_base:
            return text_slot
    return None


def _media_slot_base_name(name: str) -> str:
    normalized = re.sub(r"\s+", "", name.casefold())
    normalized = re.sub(r"(照片|图片|photo|image|img|pic)$", "", normalized)
    normalized = re.sub(r"(描述|说明|备注|caption|description|desc|情况|结果|content)$", "", normalized)
    return normalized


def _extract_docx_blank_table_cell_slot(
    *,
    table,
    table_index: int,
    row_index: int,
    column_index: int,
    slot_offset: int,
) -> TemplateSlot | None:
    row = table.rows[row_index]
    cell = row.cells[column_index]
    if not _docx_cell_is_blank(cell):
        return None

    label = _left_label_for_blank_cell(row, column_index)
    label_source = "left_cell"
    placement_constraints: dict[str, str] = {}

    if label is None:
        label = _header_label_for_blank_cell(table, row_index, column_index)
        label_source = "header_cell"
        if label is not None:
            placement_constraints["header_name"] = label

    if label is None:
        return None

    placement_constraints.update(_blank_label_constraints(label))
    placement_constraints["label_source"] = label_source

    return TemplateSlot(
        slot_id=f"docx-table-{table_index}-cell-blank-{slot_offset}",
        raw_name=label,
        canonical_name=_canonicalize_name(label),
        slot_kind=_slot_kind_for_name(label),
        address=SlotAddress(
            format_key="docx",
            address_kind="table_cell_blank",
            cell_paragraph_index=0,
            table_index=table_index,
            row_index=row_index,
            column_index=column_index,
        ),
        mode="table_cell_blank",
        context_before=label,
        placement_constraints=placement_constraints,
    )


def _extract_docx_instruction_table_cell_slot(
    *,
    table,
    table_index: int,
    row_index: int,
    column_index: int,
    slot_offset: int,
) -> TemplateSlot | None:
    row = table.rows[row_index]
    cell = row.cells[column_index]
    instruction_text = _instruction_placeholder_text(cell.text)
    if instruction_text is None:
        return None

    label = _left_label_for_blank_cell(row, column_index)
    if label is None:
        return None

    placement_constraints = _blank_label_constraints(label)
    placement_constraints["label_source"] = "left_cell"
    placement_constraints["instruction_text"] = instruction_text

    return TemplateSlot(
        slot_id=f"docx-table-{table_index}-cell-instruction-{slot_offset}",
        raw_name=label,
        canonical_name=_canonicalize_name(label),
        slot_kind=_slot_kind_for_name(label),
        address=SlotAddress(
            format_key="docx",
            address_kind="table_cell_instruction",
            cell_paragraph_index=0,
            table_index=table_index,
            row_index=row_index,
            column_index=column_index,
        ),
        mode="table_cell_instruction",
        context_before=label,
        context_after=instruction_text,
        placement_constraints=placement_constraints,
    )


def _extract_docx_choice_table_cell_slot(
    *,
    table,
    table_index: int,
    row_index: int,
    column_index: int,
    slot_offset: int,
) -> TemplateSlot | None:
    row = table.rows[row_index]
    cell = row.cells[column_index]
    choice_options = _docx_checkbox_options(cell.text)
    if not choice_options:
        return None
    choices = [option["label"] for option in choice_options]

    label = _left_label_for_blank_cell(row, column_index)
    if label is None:
        return None

    placement_constraints = _blank_label_constraints(label)
    placement_constraints["inferred_value_type"] = "choice"
    placement_constraints["label_source"] = "left_cell"
    placement_constraints["choices"] = choices
    placement_constraints["choice_options"] = choice_options

    return TemplateSlot(
        slot_id=f"docx-table-{table_index}-cell-choice-{slot_offset}",
        raw_name=label,
        canonical_name=_canonicalize_name(label),
        slot_kind="text",
        address=SlotAddress(
            format_key="docx",
            address_kind="table_cell_choice",
            cell_paragraph_index=0,
            table_index=table_index,
            row_index=row_index,
            column_index=column_index,
        ),
        mode="table_cell_choice",
        context_before=label,
        context_after=cell.text.strip(),
        placement_constraints=placement_constraints,
    )


def _docx_checkbox_options(text: str | None) -> list[dict[str, object]]:
    if text is None or "□" not in text:
        return []
    options: list[dict[str, object]] = []
    for match in _DOCX_CHOICE_PATTERN.finditer(text):
        raw_choice = re.sub(r"\s+", "", match.group(1))
        label = _docx_choice_label(raw_choice)
        if not label:
            continue
        option: dict[str, object] = {
            "label": label,
            "raw_text": raw_choice,
        }
        detail_prompt = _docx_choice_detail_prompt(raw_choice)
        if detail_prompt is not None:
            option["has_detail_blank"] = True
            option["detail_prompt"] = detail_prompt
        options.append(option)
    return options


def _docx_choice_label(raw_choice: str) -> str:
    return re.sub(r"[（(].*?[）)]", "", raw_choice).strip("：:；;，,。.、 ")


def _docx_choice_detail_prompt(raw_choice: str) -> str | None:
    parenthetical = re.search(r"[（(](?P<body>.*?)[）)]", raw_choice)
    if parenthetical is None:
        return None
    body = parenthetical.group("body")
    if re.search(r"[_＿]{2,}", body) is None:
        return None
    prompt = re.split(r"[:：_＿]+", body, maxsplit=1)[0].strip("：:；;，,。.、 ")
    return prompt or "detail"


def _instruction_placeholder_text(text: str | None) -> str | None:
    if text is None or _PLACEHOLDER_PATTERN.search(text):
        return None
    normalized = re.sub(r"\s+", "", text.strip())
    if not normalized:
        return None
    if "□" in normalized:
        return None
    if not _looks_like_instruction_placeholder(normalized):
        return None
    return text.strip()


def _looks_like_instruction_placeholder(text: str) -> bool:
    wrapped = (
        (text.startswith("（") and text.endswith("）"))
        or (text.startswith("(") and text.endswith(")"))
        or (text.startswith("【") and text.endswith("】"))
        or (text.startswith("[") and text.endswith("]"))
    )
    if wrapped:
        return True
    return any(token in text for token in ("字以内", "简要说明", "详细说明", "注明", "名称、", "编号、"))


def _docx_cell_is_blank(cell) -> bool:
    return not (cell.text or "").strip()


def _left_label_for_blank_cell(row, column_index: int) -> str | None:
    for candidate_index in range(column_index - 1, -1, -1):
        label = _clean_blank_cell_label(row.cells[candidate_index].text)
        if label is not None:
            return label
    return None


def _header_label_for_blank_cell(table, row_index: int, column_index: int) -> str | None:
    if row_index <= 0:
        return None

    current_row_labels = [_clean_blank_cell_label(cell.text) for cell in table.rows[row_index].cells]
    if any(label is not None for label in current_row_labels):
        return None

    previous_row_labels = [_clean_blank_cell_label(cell.text) for cell in table.rows[row_index - 1].cells]
    if not _looks_like_header_row(previous_row_labels):
        return None
    if column_index >= len(previous_row_labels):
        return None
    return previous_row_labels[column_index]


def _looks_like_header_row(labels: list[str | None]) -> bool:
    return sum(1 for label in labels if label is not None) >= 2


def _clean_blank_cell_label(text: str | None) -> str | None:
    if text is None:
        return None
    if _PLACEHOLDER_PATTERN.search(text):
        return None
    normalized = re.sub(r"\s+", "", text)
    normalized = _strip_placeholder_x_markers(normalized)
    normalized = normalized.strip("：:；;，,。.、 ")
    if not normalized:
        return None
    if len(normalized) > 40:
        return None
    return normalized


def _blank_label_constraints(label: str) -> dict[str, str]:
    compact = re.sub(r"\s+", "", label)
    constraints: dict[str, str] = {}
    if _is_photo_header(compact):
        constraints["inferred_value_type"] = "image"
        return constraints
    if any(token in compact for token in ("电话", "手机号", "手机", "联系方式", "联系电话")):
        constraints["inferred_value_type"] = "phone"
        return constraints
    if compact in {"姓名", "人员姓名"} or compact.endswith(("人员姓名", "联系人", "负责人", "责任人", "检查人")):
        constraints["inferred_value_type"] = "person_name"
        return constraints
    if any(token in compact for token in ("日期", "时间", "期限", "年月日", "截止")):
        constraints["inferred_value_type"] = "datetime"
        if any(token in compact for token in ("期限", "截止")):
            constraints["temporal_granularity"] = "deadline_phrase"
        else:
            constraints["temporal_granularity"] = "date_or_datetime"
        return constraints
    if any(token in compact for token in ("简介", "情况", "创新", "说明", "事迹", "业绩")):
        constraints["inferred_value_type"] = "long_text"
        return constraints
    constraints["inferred_value_type"] = "text"
    return constraints


def _docx_cell_slot_dedupe_key(
    *,
    table_index: int,
    cell_token: str,
    paragraph_index_in_cell: int,
    slot: TemplateSlot,
) -> tuple[int, str, int, str]:
    if slot.mode == "implicit_blank":
        slot_key = f"{slot.raw_name}:{slot.address.run_index}:{slot.address.run_end_index}"
    else:
        slot_key = slot.raw_name
    return (table_index, cell_token, paragraph_index_in_cell, slot_key)


def _docx_cell_token(cell, *, table_index: int, row_index: int, column_index: int) -> str:
    try:
        return cell._tc.getroottree().getpath(cell._tc)
    except Exception:
        return f"table:{table_index}:row:{row_index}:column:{column_index}"


def _neighbor_cell_text(row, column_index: int, *, direction: int) -> str | None:
    if direction not in {-1, 1}:
        return None

    candidate_indices = range(column_index - 1, -1, -1) if direction < 0 else range(column_index + 1, len(row.cells))
    for candidate_index in candidate_indices:
        text = row.cells[candidate_index].text.strip()
        if text:
            return text
    return None


def _docx_paragraph_repeat_regions(slots: list[TemplateSlot]) -> list[RepeatRegion]:
    paragraph_slots_by_index: dict[int, list[TemplateSlot]] = {}
    for slot in slots:
        if slot.address.address_kind != "paragraph" or slot.address.paragraph_index is None:
            continue
        paragraph_slots_by_index.setdefault(slot.address.paragraph_index, []).append(slot)

    repeat_regions: list[RepeatRegion] = []
    consumed_paragraph_indexes: set[int] = set()
    for paragraph_index in sorted(paragraph_slots_by_index):
        if paragraph_index in consumed_paragraph_indexes:
            continue

        paragraph_slots = paragraph_slots_by_index[paragraph_index]
        image_slots = [slot for slot in paragraph_slots if slot.slot_kind == "image"]
        if len(image_slots) != 1:
            continue

        region_slots = list(paragraph_slots)
        consumed_paragraph_indexes.add(paragraph_index)

        next_index = paragraph_index + 1
        while next_index in paragraph_slots_by_index:
            next_slots = paragraph_slots_by_index[next_index]
            if any(slot.slot_kind == "image" for slot in next_slots):
                break
            if not all(slot.slot_kind == "text" for slot in next_slots):
                break
            if not all(_is_media_description_slot(slot) for slot in next_slots):
                break
            region_slots.extend(next_slots)
            consumed_paragraph_indexes.add(next_index)
            next_index += 1

        region_id = f"docx-repeat-paragraph-{len(repeat_regions)}"
        description_slots = [slot for slot in region_slots if slot.slot_kind == "text"]
        image_slot = image_slots[0]
        for slot in region_slots:
            slot.repeat_region_id = region_id
            slot.placement_constraints.setdefault("multi_value_allowed", True)
            if slot.slot_kind == "image":
                slot.placement_constraints.setdefault("media_group_role", "image")
                if description_slots:
                    slot.placement_constraints.setdefault("paired_caption_slot_id", description_slots[0].slot_id)
            elif slot.slot_kind == "text":
                slot.placement_constraints.setdefault("media_group_role", "caption")
                slot.placement_constraints.setdefault("paired_image_slot_id", image_slot.slot_id)

        repeat_regions.append(
            RepeatRegion(
                region_id=region_id,
                format_key="docx",
                region_kind="docx_paragraph_media_region",
                address=SlotAddress(
                    format_key="docx",
                    address_kind="paragraph",
                    paragraph_index=paragraph_index,
                ),
            )
        )

    return repeat_regions


def _is_media_description_slot(slot: TemplateSlot) -> bool:
    lowered = f"{slot.raw_name} {slot.canonical_name}".casefold()
    return any(token in lowered for token in _MEDIA_DESCRIPTION_HINTS)


def _is_media_caption_slot_name(name: str) -> bool:
    lowered = re.sub(r"\s+", "", name.casefold())
    has_photo_hint = any(token in lowered for token in _PHOTO_HEADER_HINTS)
    has_text_hint = any(token in lowered for token in _PHOTO_TEXT_HINTS)
    return has_text_hint or (has_photo_hint and any(token in lowered for token in _MEDIA_DESCRIPTION_HINTS))


def _extract_implicit_blank_slots(
    *,
    paragraph,
    slot_prefix: str,
    slot_offset: int,
    address: SlotAddress,
) -> list[TemplateSlot]:
    groups: list[tuple[int, int]] = []
    index = 0
    runs = paragraph.runs
    while index < len(runs):
        run = runs[index]
        if not _is_fill_blank_run(run):
            index += 1
            continue
        start = index
        while index + 1 < len(runs) and _is_fill_blank_run(runs[index + 1]):
            index += 1
        groups.append((start, index))
        index += 1

    slots: list[TemplateSlot] = []
    for group_index, (start, end) in enumerate(groups):
        context_before = "".join(run.text for run in runs[:start])
        context_after = "".join(run.text for run in runs[end + 1 :])
        raw_name, placement_constraints = _implicit_blank_metadata(
            context_before=context_before,
            context_after=context_after,
            index=group_index,
        )
        slots.append(
            TemplateSlot(
                slot_id=f"{slot_prefix}-{slot_offset + len(slots)}",
                raw_name=raw_name,
                canonical_name=_canonicalize_name(raw_name),
                slot_kind="text",
                address=SlotAddress(
                    format_key=address.format_key,
                    address_kind=_implicit_blank_address_kind(address.address_kind),
                    paragraph_index=address.paragraph_index,
                    cell_paragraph_index=address.cell_paragraph_index,
                    run_index=start,
                    run_end_index=end,
                    table_index=address.table_index,
                    row_index=address.row_index,
                    column_index=address.column_index,
                    sheet_name=address.sheet_name,
                    cell_ref=address.cell_ref,
                    page_index=address.page_index,
                    form_field_name=address.form_field_name,
                    bbox=address.bbox,
                ),
                mode="implicit_blank",
                context_before=context_before[-120:].strip() or None,
                context_after=context_after[:120].strip() or None,
                placement_constraints=placement_constraints,
            )
        )
    return slots


def _implicit_blank_address_kind(address_kind: str) -> str:
    if address_kind in {"header_paragraph", "footer_paragraph", "header_table_cell", "footer_table_cell"}:
        return f"{address_kind}_implicit_blank"
    return "implicit_blank"


def _is_fill_blank_run(run) -> bool:
    text = run.text or ""
    return bool(text) and text.isspace() and bool(run.underline)


def _implicit_blank_metadata(*, context_before: str, context_after: str, index: int) -> tuple[str, dict[str, str]]:
    before = _strip_placeholder_x_markers(re.sub(r"\s+", "", context_before))
    after = _strip_placeholder_x_markers(re.sub(r"\s+", "", context_after))
    constraints: dict[str, str] = {}
    if after.startswith("年"):
        constraints["inferred_value_type"] = "datetime"
        constraints["temporal_granularity"] = "year"
        return f"datetime_year_blank_{index + 1}", constraints
    if after.startswith("月"):
        constraints["inferred_value_type"] = "datetime"
        constraints["temporal_granularity"] = "month"
        return f"datetime_month_blank_{index + 1}", constraints
    if after.startswith(("日内", "天内", "个工作日内", "个日历日内")):
        constraints["inferred_value_type"] = "duration"
        constraints["temporal_granularity"] = "days"
        return f"duration_days_blank_{index + 1}", constraints
    if after.startswith("日"):
        constraints["inferred_value_type"] = "datetime"
        constraints["temporal_granularity"] = "day"
        return f"datetime_day_blank_{index + 1}", constraints
    if after.startswith("前") or "之前" in after[:8] or before.endswith("截至"):
        constraints["inferred_value_type"] = "datetime"
        constraints["temporal_granularity"] = "deadline_phrase"
        return f"datetime_deadline_blank_{index + 1}", constraints
    if before.endswith("于"):
        constraints["inferred_value_type"] = "datetime"
        constraints["temporal_granularity"] = "date_or_datetime"
        return f"datetime_blank_{index + 1}", constraints
    contextual_label = _contextual_blank_label(context_before)
    if contextual_label is not None:
        constraints = _blank_label_constraints(contextual_label)
        constraints["label_source"] = "context_before"
        return contextual_label, constraints
    constraints["inferred_value_type"] = "text"
    return f"blank_{index + 1}", constraints


def _contextual_blank_label(context_before: str) -> str | None:
    compact = re.sub(r"\s+", "", context_before or "")
    if not compact:
        return None
    parts = [part for part in re.split(r"[:：；;，,。.\n\r]", compact) if part]
    label = parts[-1] if parts else compact
    label = re.sub(r"[（(][^）)]{0,12}[）)]$", "", label)
    label = _strip_placeholder_x_markers(label)
    label = label.strip("：:；;，,。.、 ")
    if not label or "{" in label or "}" in label:
        return None
    if len(label) < 2 or len(label) > 30:
        return None
    return label


def _strip_placeholder_x_markers(value: str) -> str:
    if not any("\u4e00" <= character <= "\u9fff" for character in value):
        return value
    return re.sub(r"X+$", "", re.sub(r"^X+", "", value))


def _analyze_xlsx(path: Path, analyzer_key: str) -> TemplateModel:
    workbook = load_workbook(path, data_only=False)
    slots: list[TemplateSlot] = []
    header_regions: list[HeaderDrivenRegion] = []

    for worksheet in workbook.worksheets:
        for header_row_index in range(worksheet.max_row):
            header_cells = _header_cells_for_row(worksheet, header_row_index)
            if len(header_cells) < 2:
                continue

            sample_row_index = _find_sample_row_index(worksheet, header_row_index, [column for column, _ in header_cells])
            is_blank_template_row = sample_row_index is None
            if sample_row_index is None:
                sample_row_index = header_row_index + 1

            has_placeholders = False
            for column_index, _ in header_cells:
                sample_value = _string_value(worksheet.cell(row=sample_row_index + 1, column=column_index + 1).value)
                if sample_value and _PLACEHOLDER_PATTERN.search(sample_value):
                    has_placeholders = True
                    break

            photo_columns = [(column_index, name) for column_index, name in header_cells if _is_photo_header(name)]
            if not has_placeholders and not photo_columns and not is_blank_template_row:
                continue

            region = HeaderDrivenRegion(
                region_id=f"xlsx-header-{worksheet.title}-{header_row_index}",
                format_key="xlsx",
                sheet_name=worksheet.title,
                header_row_index=header_row_index,
                column_names=[name for _, name in header_cells],
                sample_row_index=sample_row_index,
            )
            header_regions.append(region)

            image_cell_refs: set[str] = set()
            for column_index, header_name in header_cells:
                sample_cell = worksheet.cell(row=sample_row_index + 1, column=column_index + 1)
                sample_text = _string_value(sample_cell.value)
                if sample_text:
                    extracted_slots = _extract_placeholder_slots(
                        text=sample_text,
                        slot_prefix=f"xlsx-{worksheet.title}",
                        slot_offset=len(slots),
                        address=SlotAddress(
                            format_key="xlsx",
                            address_kind="cell",
                            sheet_name=worksheet.title,
                            cell_ref=sample_cell.coordinate,
                            row_index=sample_row_index,
                            column_index=column_index,
                        ),
                        repeat_region_id=region.region_id,
                    )
                    for slot in extracted_slots:
                        if slot.slot_kind == "image":
                            image_cell_refs.add(sample_cell.coordinate)
                    slots.extend(extracted_slots)
                elif not _is_photo_header(header_name):
                    slots.append(
                        TemplateSlot(
                            slot_id=f"xlsx-text-{len(slots)}",
                            raw_name=header_name,
                            canonical_name=_canonicalize_name(header_name),
                            slot_kind="text",
                            address=SlotAddress(
                                format_key="xlsx",
                                address_kind="cell",
                                sheet_name=worksheet.title,
                                cell_ref=sample_cell.coordinate,
                                row_index=sample_row_index,
                                column_index=column_index,
                            ),
                            mode="header_column",
                            repeat_region_id=region.region_id,
                            placement_constraints={"header_name": header_name},
                        )
                    )

                if _is_photo_header(header_name) and sample_cell.coordinate not in image_cell_refs:
                    slots.append(
                        TemplateSlot(
                            slot_id=f"xlsx-image-{len(slots)}",
                            raw_name=header_name,
                            canonical_name=_canonicalize_name(header_name),
                            slot_kind="image",
                            address=SlotAddress(
                                format_key="xlsx",
                                address_kind="cell",
                                sheet_name=worksheet.title,
                                cell_ref=sample_cell.coordinate,
                                row_index=sample_row_index,
                                column_index=column_index,
                            ),
                            mode="header_column",
                            repeat_region_id=region.region_id,
                            placement_constraints={"header_name": header_name},
                        )
                    )

    diagnostics: list[TemplateDiagnostic] = []
    if not header_regions and not slots:
        diagnostics.append(
            TemplateDiagnostic(
                severity="info",
                code="xlsx_no_template_regions",
                message="No header-driven regions or placeholder cells were detected in the XLSX template.",
            )
        )

    return TemplateModel(
        template_id=_template_id_for(path),
        template_path=str(path),
        format_key="xlsx",
        analyzer_key=analyzer_key,
        slots=slots,
        header_regions=header_regions,
        diagnostics=diagnostics,
    )


def _analyze_pdf(path: Path, analyzer_key: str) -> TemplateModel:
    payload = path.read_bytes().decode("latin-1", errors="ignore")
    slots: list[TemplateSlot] = []

    for dictionary_match in _PDF_DICT_PATTERN.finditer(payload):
        body = dictionary_match.group("body")
        field_name_match = _PDF_FIELD_NAME_PATTERN.search(body)
        if field_name_match is None:
            continue

        raw_name = _unescape_pdf_literal(field_name_match.group("name")).strip()
        if not raw_name:
            continue

        field_type_match = _PDF_FIELD_TYPE_PATTERN.search(body)
        field_type = field_type_match.group("field_type") if field_type_match is not None else "Unknown"

        slots.append(
            TemplateSlot(
                slot_id=f"pdf-field-{len(slots)}",
                raw_name=raw_name,
                canonical_name=_canonicalize_name(raw_name),
                slot_kind="text",
                address=SlotAddress(
                    format_key="pdf",
                    address_kind="form_field",
                    form_field_name=raw_name,
                ),
                mode="form_field",
                placement_constraints={"field_type": field_type},
            )
        )

    diagnostics: list[TemplateDiagnostic] = []
    if not slots:
        diagnostics.append(
            TemplateDiagnostic(
                severity="info",
                code="pdf_analysis_limited",
                message="PDF template analysis is currently limited to AcroForm field discovery; no form fields were found.",
            )
        )

    return TemplateModel(
        template_id=_template_id_for(path),
        template_path=str(path),
        format_key="pdf",
        analyzer_key=analyzer_key,
        slots=slots,
        diagnostics=diagnostics,
    )


def _extract_placeholder_slots(
    *,
    text: str,
    slot_prefix: str,
    slot_offset: int,
    address: SlotAddress,
    repeat_region_id: str | None = None,
    paragraph=None,
) -> list[TemplateSlot]:
    slots: list[TemplateSlot] = []
    placement_constraints = _paragraph_list_constraints(paragraph)
    for match_index, match in enumerate(_PLACEHOLDER_PATTERN.finditer(text)):
        raw_name = match.group(1).strip()
        if not raw_name:
            continue

        slots.append(
            TemplateSlot(
                slot_id=f"{slot_prefix}-{slot_offset + match_index}",
                raw_name=raw_name,
                canonical_name=_canonicalize_name(raw_name),
                slot_kind=_slot_kind_for_name(raw_name),
                address=address,
                context_before=text[: match.start()] or None,
                context_after=text[match.end() :] or None,
                repeat_region_id=repeat_region_id,
                placement_constraints=dict(placement_constraints),
            )
        )

    return slots


def _paragraph_list_constraints(paragraph) -> dict[str, object]:
    if paragraph is None or not _paragraph_is_list_item(paragraph):
        return {}
    return {
        "list_item": True,
        "list_style": _paragraph_list_style(paragraph),
        "multi_value_allowed": True,
    }


def _paragraph_is_list_item(paragraph) -> bool:
    ppr = paragraph._p.pPr
    if ppr is not None and ppr.numPr is not None:
        return True
    style_name = paragraph.style.name.casefold() if paragraph.style is not None and paragraph.style.name else ""
    return "list" in style_name


def _paragraph_list_style(paragraph) -> str:
    style_name = paragraph.style.name.casefold() if paragraph.style is not None and paragraph.style.name else ""
    if "bullet" in style_name:
        return "bulleted"
    return "numbered"


def _header_cells_for_row(worksheet, row_index: int) -> list[tuple[int, str]]:
    header_cells: list[tuple[int, str]] = []
    for column_index in range(worksheet.max_column):
        text = _string_value(worksheet.cell(row=row_index + 1, column=column_index + 1).value)
        if not text:
            continue
        if _PLACEHOLDER_PATTERN.search(text):
            return []
        header_cells.append((column_index, text))
    return header_cells


def _find_sample_row_index(worksheet, header_row_index: int, column_indices: list[int]) -> int | None:
    for row_index in range(header_row_index + 1, worksheet.max_row):
        for column_index in column_indices:
            text = _string_value(worksheet.cell(row=row_index + 1, column=column_index + 1).value)
            if text:
                return row_index
    return None


def _string_value(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def _slot_kind_for_name(name: str) -> str:
    return "image" if _is_photo_header(name) else "text"


def _is_photo_header(name: str) -> bool:
    lowered = re.sub(r"\s+", "", name.casefold())
    if any(token in lowered for token in _PHOTO_HEADER_HINTS) and any(token in lowered for token in _PHOTO_TEXT_HINTS):
        return False
    return any(token in lowered for token in _PHOTO_HEADER_HINTS)


def _canonicalize_name(name: str) -> str:
    normalized = re.sub(r"\s+", " ", name).strip()
    canonical = re.sub(r"[^\w]+", "_", normalized, flags=re.UNICODE).strip("_")
    return canonical.casefold() or normalized


def _template_id_for(path: Path) -> str:
    return _canonicalize_name(path.stem) or path.stem


def _unescape_pdf_literal(value: str) -> str:
    return value.replace(r"\(", "(").replace(r"\)", ")").replace(r"\\", "\\")
