from __future__ import annotations

import hashlib
import re
import shutil
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook

from doctransai.renderers.docx_package import prune_unreferenced_docx_media_parts


@dataclass(slots=True)
class TemplateNormalizationResult:
    template_path: Path
    analysis_template_path: Path
    diagnostics: list[dict[str, Any]] = field(default_factory=list)
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass(slots=True)
class _DocxCellNormalization:
    text: str
    image_placeholder_label: str | None = None


def normalize_sample_template(
    template_path: str | Path,
    output_path: str | Path,
    *,
    enabled: bool = False,
    cache_root: str | Path | None = None,
) -> TemplateNormalizationResult:
    resolved_template_path = Path(template_path).expanduser().resolve()
    resolved_output_path = Path(output_path).expanduser().resolve()
    suffix = resolved_template_path.suffix.lower()

    if not enabled:
        return TemplateNormalizationResult(
            template_path=resolved_template_path,
            analysis_template_path=resolved_template_path,
            metadata={
                "enabled": False,
                "strategy": "sample_normalization_disabled",
                "template_path": str(resolved_template_path),
                "normalized_template_path": str(resolved_template_path),
            },
        )

    if suffix not in {".docx", ".xlsx"}:
        diagnostic = {
            "severity": "warning",
            "code": "sample_template_normalization_unsupported_format",
            "message": f"Sample template normalization is not supported for {suffix or '<none>'}: {resolved_template_path}",
            "field": None,
        }
        return TemplateNormalizationResult(
            template_path=resolved_template_path,
            analysis_template_path=resolved_template_path,
            diagnostics=[diagnostic],
            metadata={
                "enabled": True,
                "strategy": "unsupported_format",
                "template_path": str(resolved_template_path),
                "normalized_template_path": str(resolved_template_path),
                "requested_template_format": suffix.lstrip("."),
            },
        )

    normalization_root = (
        Path(cache_root).expanduser().resolve()
        if cache_root is not None
        else (resolved_output_path.parent / ".doctransai_template_normalization").resolve()
    )
    cache_dir = normalization_root / _cache_key(resolved_template_path)
    cache_dir.mkdir(parents=True, exist_ok=True)
    normalized_path = (cache_dir / f"{resolved_template_path.stem}.sample_normalized{suffix}").resolve()
    shutil.copy2(resolved_template_path, normalized_path)

    if suffix == ".xlsx":
        stats = _normalize_xlsx(normalized_path)
    else:
        stats = _normalize_docx(normalized_path)

    diagnostic = {
        "severity": "info",
        "code": "sample_template_normalized",
        "message": f"Normalized filled sample template for analysis/rendering: {resolved_template_path} -> {normalized_path}",
        "field": None,
    }
    if not stats.get("changed"):
        diagnostic = {
            "severity": "warning",
            "code": "sample_template_normalization_no_changes",
            "message": f"Sample template normalization did not find clearable sample content: {resolved_template_path}",
            "field": None,
        }

    return TemplateNormalizationResult(
        template_path=resolved_template_path,
        analysis_template_path=normalized_path,
        diagnostics=[diagnostic],
        metadata={
            "enabled": True,
            "strategy": "sample_template_normalization",
            "template_path": str(resolved_template_path),
            "normalized_template_path": str(normalized_path),
            "requested_template_format": suffix.lstrip("."),
            **stats,
        },
    )


def _normalize_xlsx(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path)
    placeholderized_cell_count = 0
    removed_image_count = 0
    normalized_regions: list[dict[str, Any]] = []

    for worksheet in workbook.worksheets:
        header_row = _first_xlsx_header_row(worksheet)
        if header_row is None:
            continue

        data_start_row = header_row + 1
        if not _xlsx_row_has_sample_evidence(worksheet, header_row=header_row, row_number=data_start_row):
            continue

        for column in range(1, worksheet.max_column + 1):
            header_label = _clean_header_text(worksheet.cell(row=header_row, column=column).value)
            if header_label is None:
                continue
            cell = worksheet.cell(row=data_start_row, column=column)
            placeholder = f"{{{header_label}}}"
            if cell.value != placeholder:
                cell.value = placeholder
                placeholderized_cell_count += 1

        original_images = list(getattr(worksheet, "_images", []))
        kept_images = []
        for image in original_images:
            anchor_row, anchor_column = _xlsx_image_anchor_cell(image)
            if (
                anchor_row == data_start_row
                and anchor_column is not None
                and _is_xlsx_photo_header(worksheet.cell(row=header_row, column=anchor_column).value)
            ):
                removed_image_count += 1
                continue
            kept_images.append(image)
        worksheet._images = kept_images

        normalized_regions.append(
            {
                "sheet_name": worksheet.title,
                "header_row": header_row,
                "sample_row": data_start_row,
            }
        )

    workbook.save(path)
    return {
        "changed": bool(placeholderized_cell_count or removed_image_count),
        "normalized_region_count": len(normalized_regions),
        "normalized_regions": normalized_regions,
        "cleared_cell_count": 0,
        "placeholderized_cell_count": placeholderized_cell_count,
        "removed_image_count": removed_image_count,
    }


def _first_xlsx_header_row(worksheet) -> int | None:
    for row_number in range(1, worksheet.max_row + 1):
        values = [_clean_header_text(worksheet.cell(row=row_number, column=column).value) for column in range(1, worksheet.max_column + 1)]
        labels = [value for value in values if value]
        if len(labels) < 2:
            continue
        if any("{" in value and "}" in value for value in labels):
            continue
        short_label_count = sum(1 for value in labels if len(value) <= 30)
        if short_label_count >= 2:
            return row_number
    return None


def _clean_header_text(value: object) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", "", str(value))
    return text.strip("：:；;，,。.、 ") or None


def _xlsx_row_has_sample_evidence(worksheet, *, header_row: int, row_number: int) -> bool:
    if row_number > worksheet.max_row:
        return False

    data_cell_count = 0
    for column in range(1, worksheet.max_column + 1):
        header_label = _clean_header_text(worksheet.cell(row=header_row, column=column).value)
        if header_label is None:
            continue
        value = worksheet.cell(row=row_number, column=column).value
        if value is not None and _contains_sample_marker(str(value)):
            return True
        if _looks_like_filled_xlsx_sample_value(value, header_label=header_label):
            data_cell_count += 1

    if not any(_is_xlsx_photo_header(worksheet.cell(row=header_row, column=column).value) for column in range(1, worksheet.max_column + 1)):
        return data_cell_count >= 2

    for image in getattr(worksheet, "_images", []):
        anchor_row, anchor_column = _xlsx_image_anchor_cell(image)
        if (
            anchor_row == row_number
            and anchor_column is not None
            and _is_xlsx_photo_header(worksheet.cell(row=header_row, column=anchor_column).value)
        ):
            return True
    return data_cell_count >= 2


def _looks_like_filled_xlsx_sample_value(value: object, *, header_label: str) -> bool:
    if value is None:
        return False
    text = str(value).strip()
    if not text or text.startswith("{") and text.endswith("}"):
        return False
    if _clean_header_text(text) == _clean_header_text(header_label):
        return False
    return True


def _is_xlsx_photo_header(value: object) -> bool:
    label = _clean_header_text(value)
    if label is None:
        return False
    lowered = label.casefold()
    return any(token in lowered for token in ("照片", "图片", "图像", "头像", "photo", "image", "img", "pic"))


def _xlsx_image_anchor_row(image) -> int | None:
    row, _ = _xlsx_image_anchor_cell(image)
    return row


def _xlsx_image_anchor_cell(image) -> tuple[int | None, int | None]:
    anchor = getattr(image, "anchor", None)
    if isinstance(anchor, str):
        match = re.match(r"^\$?(?P<column>[A-Za-z]+)\$?(?P<row>\d+)$", anchor)
        if match is None:
            row_match = re.search(r"\d+", anchor)
            return (int(row_match.group(0)) if row_match else None), None
        return int(match.group("row")), _xlsx_column_number(match.group("column"))
    marker = getattr(anchor, "_from", None)
    row = getattr(marker, "row", None)
    column = getattr(marker, "col", None)
    return (row + 1 if isinstance(row, int) else None), (column + 1 if isinstance(column, int) else None)


def _xlsx_column_number(column_letters: str) -> int:
    number = 0
    for char in column_letters.upper():
        if not char.isalpha():
            continue
        number = number * 26 + (ord(char) - ord("A") + 1)
    return number


def _normalize_docx(path: Path) -> dict[str, Any]:
    document = Document(path)
    changed_cells = 0
    removed_drawing_count = 0

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                normalization = _normalize_docx_cell_text(cell.text)
                if normalization is not None:
                    cell_drawing_count = _count_drawings_in_paragraphs(cell.paragraphs)
                    normalized_text = normalization.text
                    if cell_drawing_count:
                        normalized_text = _append_image_placeholder(normalized_text, normalization.image_placeholder_label)
                        removed_drawing_count += cell_drawing_count
                    _replace_cell_text(cell, normalized_text)
                    changed_cells += 1

    document.save(path)
    removed_media_part_count = prune_unreferenced_docx_media_parts(path)
    return {
        "changed": bool(changed_cells or removed_drawing_count or removed_media_part_count),
        "normalized_docx_cell_count": changed_cells,
        "removed_image_count": removed_drawing_count,
        "removed_media_part_count": removed_media_part_count,
    }


def _count_drawings_in_paragraphs(paragraphs) -> int:
    count = 0
    for paragraph in paragraphs:
        for tag in ("w:drawing", "w:pict"):
            count += len(paragraph._p.xpath(f".//{tag}"))
    return count


def _normalize_docx_cell_text(text: str) -> _DocxCellNormalization | None:
    if not text.strip() or "{" in text or "}" in text:
        return None

    lines = text.splitlines()
    normalized_lines: list[str] = []
    changed = False
    current_section_label: str | None = None
    current_section_has_slot = False
    image_placeholder_label: str | None = None

    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            normalized_lines.append(raw_line)
            continue

        section_label = _section_heading_label(line)
        if section_label is not None:
            normalized_lines.append(raw_line)
            current_section_label = section_label
            current_section_has_slot = False
            continue

        label_value = _label_value_line(line)
        if label_value is not None:
            label, separator = label_value
            normalized_lines.append(f"{label}{separator}{{{label}}}")
            changed = True
            image_placeholder_label = _image_placeholder_label(label)
            continue

        if current_section_label is not None and _looks_like_sample_body_line(line):
            if not current_section_has_slot:
                normalized_lines.append(f"{{{current_section_label}}}")
                current_section_has_slot = True
                image_placeholder_label = _image_placeholder_label(current_section_label)
            changed = True
            continue

        normalized_lines.append(raw_line)

    if not changed:
        return None
    return _DocxCellNormalization(
        text="\n".join(normalized_lines),
        image_placeholder_label=image_placeholder_label,
    )


def _section_heading_label(text: str) -> str | None:
    match = re.match(r"^\s*(?:[一二三四五六七八九十]+[、.．]|\d+[、.)．])\s*(?P<label>[^:：]{2,40})\s*$", text)
    if match is None:
        return None
    raw_label = match.group("label")
    if re.search(r"[。.!！?？；;，,]", raw_label):
        return None
    label = _clean_slot_label(raw_label)
    if label is None:
        return None
    if len(label) > 24:
        return None
    return label


def _label_value_line(text: str) -> tuple[str, str] | None:
    match = re.match(r"^\s*(?P<label>[^:：]{2,30})(?P<separator>[:：])\s*(?P<value>.+?)\s*$", text)
    if match is None:
        return None
    label = _clean_slot_label(match.group("label"))
    if label is None:
        return None
    value = match.group("value").strip()
    if not value or value.startswith("{"):
        return None
    return label, match.group("separator")


def _looks_like_sample_body_line(text: str) -> bool:
    compact = re.sub(r"\s+", "", text)
    if len(compact) < 2:
        return False
    if compact.startswith("{") and compact.endswith("}"):
        return False
    if _section_heading_label(text) is not None:
        return False
    if len(compact) > 80:
        return False
    if _contains_sample_marker(compact):
        return True
    if _looks_like_instruction_or_static_clause(compact):
        return False
    return _looks_like_list_item_body(compact)


def _looks_like_list_item_body(text: str) -> bool:
    return bool(re.match(r"^(?:[（(]?[一二三四五六七八九十]+[）)、.]|\d+[、.)．])[^:：]{2,}$", text))


def _looks_like_instruction_or_static_clause(text: str) -> bool:
    return bool(re.search(r"(填写|说明|注意|须|应|请|不得|禁止|本表|本单|本文件|按.*要求)", text))


def _contains_sample_marker(text: str) -> bool:
    lowered = text.casefold()
    return any(token in lowered for token in ("样例", "示例", "例子", "sample", "example", "dummy"))


def _append_image_placeholder(text: str, label: str | None) -> str:
    image_label = _image_placeholder_label(label)
    placeholder = f"{{{image_label}}}"
    if placeholder in text:
        return text
    lines = text.splitlines()
    if lines and lines[-1].strip():
        lines.append("")
    lines.append(placeholder)
    return "\n".join(lines)


def _image_placeholder_label(section_label: str | None) -> str:
    if section_label is None:
        return "图片"
    normalized = _clean_slot_label(section_label) or "图片"
    if any(token in normalized for token in ("图片", "照片", "图像", "image", "photo")):
        return normalized
    return f"{normalized}图片"


def _clean_slot_label(value: str | None) -> str | None:
    if value is None:
        return None
    label = re.sub(r"\s+", "", value)
    label = label.strip("：:；;，,。.、 ")
    if not label or len(label) > 40:
        return None
    return label


def _replace_cell_text(cell, text: str) -> None:
    cell.text = text


def _cache_key(template_path: Path) -> str:
    stat = template_path.stat()
    payload = f"{template_path}|{stat.st_mtime_ns}|{stat.st_size}|sample-normalization-v3"
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()[:16]
