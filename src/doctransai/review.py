"""Fresh-context, read-only review followed by bounded, scoped repair."""
from dataclasses import replace
import hashlib
import json
from pathlib import Path
import shutil
import time

from jsonschema import Draft202012Validator

from doctransai.agent import run_agent, append_event
from doctransai.documents.preview import image_part
from doctransai.tools import ToolRuntime, tool, obj, TEXT, REFS

CATEGORIES = ('factual_support', 'completeness', 'template_preservation', 'images', 'layout', 'data_types_and_formulas')
ISSUE_KINDS = ('factual', 'omission', 'template_residue', 'image', 'layout', 'data_type', 'formula', 'ambiguity', 'unsupported')
READ_TOOLS = {'read_source', 'view_source_page', 'view_source_asset', 'inspect_document', 'preview_document'}
REVIEW_TOOL = tool('submit_review',
    'Submit independent findings after viewing all current output pages. Coverage must honestly assess each category; uncertain prevents passing. '
    'Use exact current-document targets (DOCX IDs or XLSX ranges). Empty targets means manual-only. '
    'Use source refs as evidence where available; source absence is not a license to invent. '
    'Only mark repairable when the action is supported and unambiguous. No verdict field: the system derives it.',
    obj(dict(summary=TEXT, coverage=obj({k: {'enum': ['pass', 'uncertain'] if k in CATEGORIES[:3] or k == 'layout' else ['pass', 'not_applicable', 'uncertain']} for k in CATEGORIES}, CATEGORIES),
        issues={'type': 'array', 'maxItems': 30, 'items': obj(dict(
            id={'type': 'string', 'minLength': 1}, kind={'enum': list(ISSUE_KINDS)}, targets=REFS,
            evidence_refs=REFS, explanation={'type': 'string', 'minLength': 1},
            suggested_action=TEXT, repairable={'type': 'boolean'}),
            ['id', 'kind', 'targets', 'evidence_refs', 'explanation', 'suggested_action', 'repairable'])},
        grounding={'type':'array','maxItems':40,'items':obj(dict(targets={**REFS,'minItems':1}, output_claim=TEXT,
            target_role=TEXT, source_claim=TEXT, source_role=TEXT, source_refs=REFS,
            judgment={'enum':['supported','unsupported','uncertain']}),
            ['targets','output_claim','target_role','source_claim','source_role','source_refs','judgment'])},
        notes={'type': 'array', 'items': TEXT}), ['summary', 'coverage', 'issues', 'grounding', 'notes']))


def _hash(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


class ReviewRuntime(ToolRuntime):
    result_key = 'independent_review'

    def __init__(self, source, editor, previewer):
        super().__init__(source, editor, previewer)
        self.tools = [t for t in self.tools if t['function']['name'] in READ_TOOLS] + [REVIEW_TOOL]
        self.validators = {t['function']['name']: Draft202012Validator(t['function']['parameters']) for t in self.tools}

    def invoke(self, name, args):
        if name not in self.validators:
            raise ValueError('Independent reviewer is read-only; that tool is unavailable')
        if name != 'submit_review':
            return super().invoke(name, args)
        self.validators[name].validate(args)
        missing = set(range(1, len(self.output_pages)+1)) - self.coverage
        if self.preview_version != self.editor.version or self.preview_round >= self.round or missing:
            raise ValueError(f'Review requires all current output pages from an earlier round; missing: {sorted(missing)}')
        ids = [i['id'] for i in args['issues']]
        if len(set(ids)) != len(ids):
            raise ValueError('Issue IDs must be unique')
        self.editor.inspect()
        for issue in args['issues']:
            self.source.check_refs(issue['evidence_refs'])
            for target in issue['targets']:
                if self.format_key == 'xlsx':
                    self.editor.resolve(target)
                elif target not in self.editor.nodes:
                    raise ValueError(f'Unknown review target: {target}; use current document IDs')
            if issue['repairable'] and (not issue['targets'] or issue['kind'] in ('ambiguity', 'unsupported')):
                raise ValueError('Ambiguous, unsupported or unlocated issues cannot be auto-repaired')
        if self.source.evidence and not args['grounding']:
            raise ValueError('Provide explicit grounding assessments for material output claims, not only category verdicts')
        for check in args['grounding']:
            self.source.check_refs(check['source_refs'])
            for target in check['targets']:
                if self.format_key == 'xlsx':self.editor.resolve(target)
                elif target not in self.editor.nodes:raise ValueError(f'Unknown grounding target: {target}')
            if check['judgment'] != 'supported' and not any(set(check['targets']) & set(i['targets']) for i in args['issues']):
                raise ValueError('Unsupported/uncertain grounding must have a located issue; it cannot be hidden in coverage or notes')
        passed = not args['issues'] and all(v != 'uncertain' for v in args['coverage'].values())
        self.finished = dict(args, status='passed' if passed else 'needs_review',
            document_version=self.editor.version, document_sha256=_hash(self.editor.output),
            pages_checked=sorted(self.coverage))
        return {'submitted': True, 'status': self.finished['status']}, []


class RepairRuntime(ToolRuntime):
    """Permission is bounded by findings, not merely by instructions to the model."""
    def __init__(self, source, editor, previewer, issues):
        super().__init__(source, editor, previewer)
        self.expected_hash = _hash(editor.output)
        self.scopes = []
        for issue in issues:
            if not issue['repairable']:continue
            for target in issue['targets']:
                if self.format_key == 'xlsx':
                    ws, bounds = editor.resolve(target)
                    self.scopes.append((ws.title, bounds, issue['kind']))
                else:
                    element = editor.nodes[target]._element
                    descendants = set(element.iter())
                    keys = {key for key, node in editor.nodes.items() if node._element in descendants}
                    self.scopes.append((keys, None, issue['kind']))

    @staticmethod
    def _operation_allowed(kind, op):
        if kind == 'layout':return op in ('format_cells', 'set_print_area', 'format_paragraph', 'format_row')
        if kind == 'image':return op in ('set_images', 'set_cell_images', 'format_cells', 'format_paragraph')
        if kind == 'data_type':return op in ('set_cells', 'format_cells')
        if kind == 'formula':return op == 'set_formula'
        return True

    def _allowed(self, target, op):
        if self.format_key == 'xlsx':
            ws, (a,b,c,d) = self.editor.resolve(target)
            ranges = [bounds for name,bounds,kind in self.scopes if name == ws.title and self._operation_allowed(kind,op)]
            return all(any(x[0] <= col <= x[2] and x[1] <= row <= x[3] for x in ranges)
                       for row in range(b,d+1) for col in range(a,c+1))
        return any(target in keys and self._operation_allowed(kind, op) for keys,_,kind in self.scopes)

    def invoke(self, name, args):
        if name == 'edit_document':
            if _hash(self.editor.output) != self.expected_hash:
                raise ValueError('Repair document changed outside this reviewed version')
            before_nodes = set(self.editor.nodes)
            self.validators[name].validate(args)
            if args['batch_id'] not in self.editor.applied_batches:
                for edit in args['edits']:
                    targets = edit.get('targets', [edit.get('target')])
                    if not all(self._allowed(t, edit['op']) for t in targets):
                        raise ValueError('Repair outside reviewed targets/operation scope is forbidden')
                    if edit['op'] == 'set_print_area':
                        ws,(a,b,c,d) = self.editor.resolve(edit['target'])
                        if any(cell.value is not None and not (a <= cell.column <= c and b <= cell.row <= d) for cell in ws._cells.values()):
                            raise ValueError('Repair print area cannot hide existing content')
                    if edit['op'] == 'copy_rows':
                        from openpyxl.utils import get_column_letter
                        _,(a,b,c,d) = self.editor.resolve(edit['target'])
                        prefix = edit['target'].split('!')[0]
                        for row in edit['destination_rows']:
                            dest=f'{prefix}!{get_column_letter(a)}{row}:{get_column_letter(c)}{row+d-b}'
                            if not self._allowed(dest, edit['op']):
                                raise ValueError('Copy destination outside reviewed repair scope')
            result = super().invoke(name, args)
            self.expected_hash = _hash(self.editor.output)
            if self.format_key != 'xlsx':
                added = set(self.editor.nodes) - before_nodes
                kinds = {kind for keys,_,kind in self.scopes if any(t in keys for e in args['edits'] for t in e.get('targets',[e.get('target')]))}
                for kind in kinds:
                    if added:self.scopes.append((added,None,kind))
            return result
        return super().invoke(name, args)


def _seed_pages(runtime, pages):
    if not pages:raise ValueError("No rendered output pages available for review")
    selected = pages[:4]
    runtime.output_pages = pages
    runtime.preview_version = runtime.editor.version
    runtime.preview_round = 0
    runtime.coverage = set(range(1,len(selected)+1))
    return [{'type':'text','text':f'Current output version {runtime.editor.version}, page {i} of {len(pages)}'}
            if j == 0 else image_part(path)
            for i,path in enumerate(selected,1) for j in range(2)]


def run_review(source, editor, previewer, initial, client, config, trace, deadline, progress=None):
    runtime = ReviewRuntime(source,editor,previewer)
    before, template_before = _hash(editor.output), _hash(editor.template)
    prompt = (Path(__file__).parent/'prompts/document_review_policy.md').read_text()
    pages = previewer.render(editor.output)
    append_event(trace, dict(kind="review_start", read_only=True, document_sha256=before, prompt_sha256=hashlib.sha256(prompt.encode()).hexdigest()))
    snapshot = editor.inspect()
    context = dict(step_budget=config.review_max_steps, time_remaining_seconds=max(0,int(deadline-time.monotonic())), task_instruction=initial['instruction'], primary_source_refs=initial['primary_source_refs'],
        source=initial['source'], original_template=initial['document'], current_document=snapshot,
        application_requirements=(Path(__file__).parent/'prompts/document_editing_policy.md').read_text(),
        repair_permissions={kind: [op for op in snapshot["supported_operations"] if RepairRuntime._operation_allowed(kind,op)] for kind in ISSUE_KINDS if kind not in ("ambiguity","unsupported")},
        output_page_count=len(pages), missing_output_pages=list(range(5,len(pages)+1)),
        address_note='Original template IDs are reference only; findings must use current_document IDs.')
    messages=[dict(role='system',content=prompt),dict(role='user',content=[
        {'type':'text','text':json.dumps(context,ensure_ascii=False)},*_seed_pages(runtime,pages)])]
    try:
        result=run_agent(client,messages,runtime,replace(config,max_steps=config.review_max_steps),trace,deadline,progress)
    finally:
        if _hash(editor.output)!=before or _hash(editor.template)!=template_before:
            raise RuntimeError('Read-only review changed a document')
    return dict(result['independent_review'], output_pages=[str(p) for p in runtime.output_pages])


def review_and_repair(source, editor, previewer, initial, client, config, output_dir, deadline, progress=None):
    history=[]
    for round_number in range(config.max_repair_rounds+1):
        phase=f'review-{round_number}'
        diagnostics=getattr(client,'diagnostics',None)
        if diagnostics is not None:diagnostics.current_phase=phase
        if progress:progress(f'独立复核 {round_number + 1}')
        try:
            review=run_review(source,editor,previewer,initial,client,config,output_dir/f'{phase}.jsonl',deadline,progress)
        except Exception as exc:
            secret=getattr(client,'api_key',None)
            message=str(exc).replace(secret,'<redacted>') if secret else str(exc)
            history.append(dict(phase=phase,status='incomplete',error=message[:1500]))
            return dict(status='incomplete',rounds=history,issues=[],reason='Independent review did not complete')
        history.append(dict(phase=phase,**review))
        (output_dir/f'{phase}.json').write_text(json.dumps(review,ensure_ascii=False,indent=2))
        if review['status']=='passed':return dict(status='passed',rounds=history,issues=[])
        repairable=[i for i in review['issues'] if i['repairable']]
        if round_number == config.max_repair_rounds or not repairable:
            return dict(status='needs_review',rounds=history,issues=review['issues'])
        if _hash(editor.output) != review['document_sha256']:
            return dict(status='incomplete',rounds=history,issues=review['issues'],reason='Document changed after review; no repair applied')
        checkpoint=output_dir/'revisions'/f'before-repair-{round_number+1}{editor.output.suffix}'
        checkpoint.parent.mkdir(exist_ok=True)
        shutil.copy2(editor.output,checkpoint)
        if progress:progress(f'有限修正 {round_number + 1}：仅处理复核指出的位置')
        runtime=RepairRuntime(source,editor,previewer,repairable)
        prompt=(Path(__file__).parent/'prompts/document_editing_policy.md').read_text()
        instruction=dict(step_budget=config.repair_max_steps, time_remaining_seconds=max(0,int(deadline-time.monotonic())), task_instruction=initial['instruction'],primary_source_refs=initial['primary_source_refs'],
            source=initial['source'],current_document=editor.inspect(),review_findings=repairable,
            repair_rule='Only repair these findings at authorized targets. Preserve other facts and structure. Do not invent missing evidence. Preview all output pages and finish; another independent review decides acceptance.')
        messages=[dict(role='system',content=prompt),dict(role='user',content=json.dumps(instruction,ensure_ascii=False))]
        version=editor.version
        if diagnostics is not None:diagnostics.current_phase=f'repair-{round_number+1}'
        try:
            run_agent(client,messages,runtime,replace(config,max_steps=config.repair_max_steps),
                      output_dir/f'repair-{round_number+1}.jsonl',deadline,progress)
        except Exception as exc:
            secret=getattr(client,'api_key',None)
            message=str(exc).replace(secret,'<redacted>') if secret else str(exc)
            history.append(dict(phase=f'repair-{round_number+1}',status='incomplete',error=message[:1500],checkpoint=str(checkpoint)))
            return dict(status='incomplete',rounds=history,issues=review['issues'],reason='Repair did not complete; pre-repair checkpoint retained')
        history.append(dict(phase=f'repair-{round_number+1}',status='completed',checkpoint=str(checkpoint),document_version=editor.version))
        if editor.version == version:
            return dict(status='needs_review',rounds=history,issues=review['issues'],reason='Repair made no changes')
    raise AssertionError('Bounded review loop must return')
