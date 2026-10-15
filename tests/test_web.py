import io
import json
from pathlib import Path
import threading
from types import SimpleNamespace

from docx import Document
import httpx
import pytest

from doctransai.web import LocalApp, make_server, run_worker, settings_config, safe_message
from doctransai.upstream.dataproxy.service import save_receipt


@pytest.fixture
def app(tmp_path):
    renderer=tmp_path/'render_docx.py';renderer.write_text('# fixture')
    env=tmp_path/'private.env'
    env.write_text(f'DOCTRANSAI_DATAPROXY_BASE_URL=https://service.invalid\nDOCTRANSAI_DATAPROXY_API_KEY=private-upstream-key\nDOCTRANSAI_LLM_BASE_URL=https://provider.invalid/v1\nDOCTRANSAI_LLM_API_KEY=private-llm-key\nDOCTRANSAI_DOCX_RENDERER={renderer}\n')
    return LocalApp(tmp_path/'jobs',env,launcher=lambda path:SimpleNamespace(poll=lambda:None))


@pytest.fixture
def server(app):
    server=make_server(app,0)
    thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    with httpx.Client(base_url=f'http://127.0.0.1:{server.server_port}') as client:
        yield client, app
    server.shutdown();server.server_close();thread.join()


def template():
    doc=Document();doc.add_paragraph('Record');out=io.BytesIO();doc.save(out);return out.getvalue()


def submit(client, app, **extra):
    kwargs=dict(headers={'X-DocTransAI-Token':app.token},files={'source':('source.txt',b'Example source'),'template':('template.docx',template())},data={'instruction':'Fill from source','settings':'{}'})
    kwargs.update(extra)
    return client.post('/api/jobs',**kwargs)


def test_frontend_and_config_do_not_disclose_keys(server):
    client,app=server
    assert client.get('/').status_code==200
    r=client.get('/api/config')
    assert r.json()['ready']
    assert 'private-llm-key' not in r.text and 'private-upstream-key' not in r.text


def test_only_local_same_origin_token_requests_can_submit(server):
    client,app=server
    assert submit(client,app,headers={}).status_code==403
    assert submit(client,app,headers={'X-DocTransAI-Token':app.token,'Origin':'https://untrusted.invalid'}).status_code==403
    assert client.get('/api/config',headers={'Host':'rebound.invalid'}).status_code==403


def test_files_saved_separately_and_busy_prevents_second_job(server):
    client,app=server
    r=submit(client,app);assert r.status_code==202,r.text
    job=r.json()['id'];root=app.root/job
    assert (root/'source/source.txt').read_bytes()==b'Example source'
    assert (root/'template/template.docx').is_file()
    assert submit(client,app).status_code==409
    assert len(list(app.root.iterdir()))==1
    assert client.get('/api/jobs').json()['jobs'][0]['id']==job


@pytest.mark.parametrize('name,data',[('template.pdf',b'%PDF'),('template.docx',b'not a zip'),('template.doc',b'legacy')])
def test_invalid_template_rejected_before_worker(server,name,data):
    client,app=server
    r=submit(client,app,files={'source':('source.txt',b'data'),'template':(name,data)})
    assert r.status_code==400 and not list(app.root.iterdir())


def test_upload_size_limit_is_checked_before_parsing(server,monkeypatch):
    import doctransai.web
    monkeypatch.setattr(doctransai.web,'MAX_REQUEST',1024)
    client,app=server
    assert submit(client,app).status_code==413
    assert not list(app.root.iterdir())


def test_xlsx_template_is_accepted(server):
    from openpyxl import Workbook
    out=io.BytesIO();book=Workbook();book.active['A1']='Field';book.save(out)
    client,app=server
    r=submit(client,app,files={'source':('source.txt',b'test'),'template':('template.xlsx',out.getvalue())})
    assert r.status_code==202,r.text
    assert r.json()['template']=='template.xlsx'


def test_missing_configuration_returns_actionable_error(server):
    client,app=server
    app.env_file.unlink()
    assert not client.get('/api/config').json()['ready']
    assert submit(client,app).status_code==400


def test_unsafe_filename_cannot_escape_job_directory(server):
    client,app=server
    r=submit(client,app,files={'source':('../../escape.txt',b'test'),'template':('template.docx',template())})
    assert r.status_code==202
    assert (app.root/r.json()['id']/'source/escape.txt').is_file()
    assert not (app.root.parent/'escape.txt').exists()


def test_download_allowlist_and_audit_sandbox(server):
    client,app=server
    job=submit(client,app).json()['id'];root=app.root/job/'run/conversion';root.mkdir(parents=True)
    (root/'result.docx').write_bytes(b'result')
    (root/'audit.html').write_text('<html>audit</html>')
    assert client.get(f'/api/jobs/{job}/files/result.docx').content==b'result'
    assert 'attachment' in client.get(f'/api/jobs/{job}/files/result.docx').headers['content-disposition']
    assert 'sandbox' in client.get(f'/api/jobs/{job}/files/audit.html').headers['content-security-policy']
    assert client.get(f'/api/jobs/{job}/files/worker.log').status_code==404
    assert client.get(f'/api/jobs/{job}/files/../../private.env').status_code==404


def test_worker_calls_existing_workflow_and_preserves_template_boundary(server,monkeypatch):
    client,app=server
    job=submit(client,app).json()['id'];called={}
    def run(**kwargs):
        called.update(kwargs)
        kwargs['event_callback']({'phase':'auditing','message':'Audit running'})
        kwargs['output_dir'].mkdir()
        return {'status':'completed'}
    monkeypatch.setattr('doctransai.upstream.dataproxy.workflow.run_remote',run)
    run_worker(app.root/job,app.env_file)
    state=client.get('/api/jobs/'+job).json()
    assert state['status']=='completed'
    assert called['source'].parent.name=='source' and called['template'].parent.name=='template'
    assert called['api_key']=='private-upstream-key'
    assert any(e['message']=='Audit running' for e in state['events'])


def test_failure_is_reported_and_secrets_redacted(server,monkeypatch):
    client,app=server;job=submit(client,app).json()['id']
    def fail(**kwargs):raise ValueError('provider echoed private-llm-key')
    monkeypatch.setattr('doctransai.upstream.dataproxy.workflow.run_remote',fail)
    run_worker(app.root/job,app.env_file)
    r=client.get('/api/jobs/'+job)
    assert r.json()['status']=='failed'
    assert 'private-llm-key' not in r.text and '已隐藏' in r.text


def test_restart_marks_unfinished_job_without_rerun(app):
    path=app.root/('a'*32);path.mkdir()
    save_receipt(path/'web-job.json',{'id':'a'*32,'status':'running'})
    new=LocalApp(app.root,app.env_file)
    assert new.status('a'*32)['status']=='interrupted'


def test_server_rejects_unknown_or_invalid_model_settings(app):
    with pytest.raises(ValueError):settings_config({'api_key':'cannot-send-from-ui'},app.env_file)
    with pytest.raises(ValueError):settings_config({'max_output_tokens':-10},app.env_file)
    cfg=settings_config({'model':'custom-vision-model'},app.env_file)
    assert cfg.reasoning_effort=='omit'
