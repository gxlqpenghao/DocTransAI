from pathlib import Path
import hashlib
import json

from docx import Document
from PIL import Image
import httpx
import pytest

from doctransai.config import RuntimeConfig
from doctransai.conversion import convert
from doctransai.data_layer.models import SourceCorpus, ComponentFile
from doctransai.data_layer.view import SourceView
from doctransai.documents.docx import DocxEditor
from doctransai.documents.preview import Previewer
from doctransai.tools import ToolRuntime
from doctransai.types import TaskSpec


def material(tmp_path):
    source = tmp_path / "source.docx"
    d = Document(); d.add_paragraph("Example project"); d.save(source)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    doc = dict(document_id="doc1", source_path=str(source), source_checksum=digest)
    rows = dict(source_manifest=[doc], evidence_units=[dict(document_id="doc1", evidence_id="ev1", text="Example project", locator={"section_title": "Award 2024"})],
                table_rows=[dict(document_id="doc1", table_row_id="row1", row_text="row content")],
                table_cells=[dict(document_id="doc1", table_cell_id="cell1", value="cell content")],
                visual_asset_inventory=[dict(document_id="doc1", asset_id="img1", source_ref="images/one.png", asset_type="scanned_page", page_anchor="page_1")])
    image = tmp_path / "parsed" / digest / "images/one.png"
    image.parent.mkdir(parents=True); Image.new("RGB", (30, 30), "white").save(image)
    corpus = SourceCorpus(tmp_path / "bundle.json", "data_asset_pack_v0", {}, {}, [doc], [],
        {k: ComponentFile(k, tmp_path / f"{k}.jsonl", len(v), v) for k, v in rows.items()}, {}, {})
    return corpus, source, image


def test_native_context_refs_and_asset_path_are_retained(tmp_path):
    corpus, _, image = material(tmp_path)
    view = SourceView(corpus)
    assert view.asset_path("img1") == image
    assert "Award 2024" in view.read(refs=["ev1"])["objects"][0]["json_text"]
    assert "row1" in view.objects and "cell1" in view.objects
    assert view.initial(1)["evidence_omitted"] == 1
    with pytest.raises(ValueError, match="out-of-document"):
        view.check_refs(["other-doc-ref"])


def test_primary_source_refs_are_first_class_initial_context(tmp_path):
    corpus, _, _ = material(tmp_path)
    view = SourceView(corpus)
    initial = view.initial(10000, primary_refs=['row1'])
    assert initial['primary_objects'][0]['ref'] == 'row1'
    assert initial['primary_objects'][0]['kind'] == 'table_rows'
    assert initial['evidence'][0]['evidence_id'] == 'ev1'


def test_primary_asset_preparation_does_not_mix_workbooks_or_sheets(tmp_path):
    corpus, _, _ = material(tmp_path)
    corpus.evidence_units.append(dict(
        document_id='doc1', evidence_id='row-primary', text='Sheet=Sheet0 | Range=A6:X6',
        locator={'cell_range': 'A6:X6', 'sheet_name': 'Sheet0'}))
    corpus.visual_assets.extend([
        dict(corpus.visual_assets[0], asset_id='same-row-sheet0', cell_anchor='K6', sheet_name='Sheet0'),
        dict(corpus.visual_assets[0], asset_id='same-row-sheet1', cell_anchor='K6', sheet_name='Sheet1'),
    ])
    view = SourceView(corpus)
    assert view.initial_asset_refs(['row-primary']) == ['same-row-sheet0']


def test_source_document_selection_is_not_implicit_for_multiple_documents(tmp_path):
    corpus, _, _ = material(tmp_path)
    corpus.documents.append(dict(document_id="doc2", source_path="/other.docx"))
    with pytest.raises(ValueError, match="exactly one"):
        SourceView(corpus)


def test_native_scanned_page_does_not_require_rendering(tmp_path):
    corpus, _, image = material(tmp_path)
    preview = Previewer(tmp_path / "cache", RuntimeConfig())
    info, path = preview.source_page(SourceView(corpus), 1)
    assert info["origin"] == "dataproxy_scanned_page" and path == image


def test_asset_content_must_match_declared_checksum(tmp_path):
    corpus, _, image = material(tmp_path)
    corpus.visual_assets[0]["checksum"] = hashlib.sha256(image.read_bytes()).hexdigest()
    view = SourceView(corpus)
    image.write_bytes(b"different")
    with pytest.raises(ValueError, match="checksum"):
        view.asset_path("img1")


def test_repeated_asset_id_preserves_bbox_variants_without_blocking_text(tmp_path):
    corpus, _, image = material(tmp_path)
    original = corpus.visual_assets[0]
    corpus.visual_assets.append(dict(original, bbox=[1, 2, 3, 4]))
    view = SourceView(corpus)
    assert view.asset_path("img1") == image
    assert len(view.objects["img1"]["variants"]) == 2
    assert view.initial(10000)["repeated_native_ids"] == ["img1"]
    assert "ev1" in view.objects
    assert "variants" not in original


def test_local_source_preview_reports_unmapped_pagination_and_changed_source(tmp_path, monkeypatch):
    corpus, source, image = material(tmp_path)
    view = SourceView(corpus)
    view.assets.clear()
    preview = Previewer(tmp_path / "cache", RuntimeConfig())
    seen = []
    monkeypatch.setattr(preview, "render", lambda path, page_number: seen.append((path, page_number)) or [image])
    info, path = preview.source_page(view, 2)
    assert info["native_anchor"] is None and "unverified" in info["mapping"]
    assert seen == [(source, 2)]
    source.write_text("changed")
    with pytest.raises(ValueError, match="changed"):
        preview.source_page(view, 2)


class FakePreview:
    def __init__(self, image, count=1):
        self.image, self.count = image, count

    def render(self, path, **kwargs):
        return [self.image] * self.count


class RecordingPreview(FakePreview):
    def __init__(self, image, count=1):
        super().__init__(image, count)
        self.calls = []

    def render(self, path, **kwargs):
        self.calls.append((path, kwargs))
        return super().render(path, **kwargs)


def test_finish_requires_all_pages_and_a_later_model_round(tmp_path):
    corpus, template, image = material(tmp_path)
    view = SourceView(corpus)
    editor = DocxEditor(template, tmp_path / "working.docx", view.asset_path)
    runtime = ToolRuntime(view, editor, FakePreview(image, 2))
    finish = dict(summary="done", checks=["checked"], missing=[], issues=[])
    runtime.round = 1
    runtime.invoke("preview_document", dict(max_pages=1))
    with pytest.raises(ValueError, match="earlier"):
        runtime.invoke("finish", finish)
    runtime.round = 2
    with pytest.raises(ValueError, match="Not all"):
        runtime.invoke("finish", finish)
    runtime.invoke("preview_document", dict(start_page=2, max_pages=1))
    runtime.round = 3
    runtime.invoke("edit_document", dict(batch_id="edit", version=0, edits=[dict(op="set_text", target="n0", text="", source_refs=[], reason="clear")]))
    with pytest.raises(ValueError, match="current"):
        runtime.invoke("finish", finish)


def test_finish_accepts_minimal_delivery_contract(tmp_path):
    corpus, template, image = material(tmp_path)
    runtime = ToolRuntime(SourceView(corpus), DocxEditor(template, tmp_path / 'working.docx', SourceView(corpus).asset_path), FakePreview(image))
    runtime.round = 1
    runtime.invoke('preview_document', {})
    runtime.round = 2
    runtime.invoke('finish', {'summary': 'delivered'})
    assert runtime.finished == {'summary': 'delivered', 'issues': [], 'missing': []}


def test_invalid_reference_is_rejected_before_any_edit(tmp_path):
    corpus, template, image = material(tmp_path)
    view = SourceView(corpus)
    editor = DocxEditor(template, tmp_path / "working.docx", view.asset_path)
    runtime = ToolRuntime(view, editor, FakePreview(image))
    before = editor.output.read_bytes()
    with pytest.raises(ValueError, match="Unknown"):
        runtime.invoke("edit_document", dict(batch_id="bad", version=0, edits=[dict(op="set_text", target="n0", text="new", source_refs=["missing"], reason="source")]))
    assert editor.output.read_bytes() == before


def test_marker_cleanup_preserves_date_labels_without_fabricating_refs(tmp_path):
    corpus, _, image = material(tmp_path)
    template = tmp_path / "template.docx"
    doc = Document(); doc.add_paragraph("X 年X 月X 日"); doc.save(template)
    view = SourceView(corpus)
    editor = DocxEditor(template, tmp_path / "working.docx", view.asset_path)
    runtime = ToolRuntime(view, editor, FakePreview(image))
    runtime.invoke("edit_document", dict(batch_id="clear", version=0, edits=[dict(
        op="set_text", target="n0", text="年 月 日", source_refs=[], reason="remove unsupported placeholders; keep units")]))
    assert Document(editor.output).paragraphs[0].text == "年 月 日"


def test_retry_of_applied_batch_succeeds_even_after_target_was_replaced(tmp_path):
    corpus, template, image = material(tmp_path)
    view = SourceView(corpus)
    editor = DocxEditor(template, tmp_path / "working.docx", view.asset_path)
    runtime = ToolRuntime(view, editor, FakePreview(image))
    args = dict(batch_id="one", version=0, edits=[dict(op="set_text", target="n0", text="Example project", source_refs=["ev1"], reason="native")])
    runtime.invoke("edit_document", args)
    runtime.invoke("edit_document", dict(batch_id="two", version=1, edits=[dict(op="repeat_paragraphs", targets=["n0"], items=[dict(contents=[dict(text="Example project")], source_refs=["ev1"])], reason="repeat")]))
    result, _ = runtime.invoke("edit_document", args)
    assert result["already_applied"] and editor.version == 2


class FakeClient:
    model_name = "test-client"

    def __init__(self):
        self.index = 0
        self.saw_output_image = False
        self.review_image_count = 0

    def chat(self, messages, tools, **kwargs):
        calls = [("edit_document", dict(batch_id="1", version=0, edits=[dict(op="set_text", target="n0", text="Example project", source_refs=["ev1"], reason="native fact")])),
                 ("preview_document", {}),
                 ("finish", dict(summary="done", checks=["seen"], missing=[], issues=[]))]
        name, args = calls[self.index]
        if self.index == 2:
            self.saw_output_image = any(p.get("type") == "image_url" for p in messages[-1]["content"])
            self.review_image_count = sum(p.get("type") == "image_url" for m in messages if isinstance(m.get("content"), list) for p in m["content"])
            assert "Example project" in messages[-1]["content"][0]["text"]
            assert '"current_document"' in messages[-1]["content"][0]["text"]
        self.index += 1
        return dict(message=dict(role="assistant", content="", tool_calls=[dict(id=str(self.index), type="function", function=dict(name=name, arguments=json.dumps(args)))]), usage=None, model=self.model_name)


def test_conversion_uses_new_tools_and_publishes_only_after_finish(tmp_path, monkeypatch):
    corpus, template, image = material(tmp_path)
    monkeypatch.setattr("doctransai.conversion.read_parsed_document_bundle", lambda _: corpus)
    client = FakeClient()
    task = TaskSpec(corpus.bundle_path, template, tmp_path / "run", "Fill the project")
    result = convert(task, config=RuntimeConfig(review_enabled=False), client=client, previewer=FakePreview(image))
    assert result.execution_status == "completed", result.issues
    assert result.review_status == "not_run"  # no independent semantic reviewer yet
    assert client.saw_output_image
    assert client.review_image_count == 2  # prepared source asset + current output page; old template page is not resent
    assert Document(result.output_path).paragraphs[0].text == "Example project"
    report = json.loads(Path(result.report_path).read_text())
    assert report["delivered_output_pages"] == [1]
    assert report["editor_self_check_completed"]
    assert report["template_unchanged"]
    assert not (task.output_dir / "working.docx").exists()


def test_successful_edit_is_automatically_previewed_before_finish(tmp_path, monkeypatch):
    corpus, template, image = material(tmp_path)
    monkeypatch.setattr("doctransai.conversion.read_parsed_document_bundle", lambda _: corpus)
    preview = RecordingPreview(image)

    class AutoClient:
        model_name = "test-client"

        def __init__(self):
            self.index = 0
            self.seen_result_check = False

        def chat(self, messages, tools, **kwargs):
            if self.index == 1:
                parts = messages[-1]["content"]
                text = parts[0]["text"]
                self.seen_result_check = ("程序已经对当前版本真实渲染" in text and
                                          any("Current output page" in part.get("text", "") for part in parts if isinstance(part, dict)))
            calls = [
                ("edit_document", dict(batch_id="auto", version=0, edits=[dict(
                    op="set_text", target="n0", text="Example project", source_refs=["ev1"], reason="native fact")])),
                ("finish", dict(summary="delivered", missing=[], issues=[])),
            ]
            name, args = calls[self.index]
            self.index += 1
            return dict(message=dict(role="assistant", content="", tool_calls=[dict(
                id=str(self.index), type="function", function=dict(name=name, arguments=json.dumps(args)))]),
                        usage=None, model=self.model_name)

    client = AutoClient()
    result = convert(TaskSpec(corpus.bundle_path, template, tmp_path / "auto-run", "Fill"),
                     config=RuntimeConfig(review_enabled=False), client=client, previewer=preview)
    assert result.execution_status == "completed", result.issues
    assert client.index == 2 and client.seen_result_check
    assert len(preview.calls) == 3  # source prep + template prep + automatic current-output render
    assert json.loads(Path(result.report_path).read_text())["auto_preview_count"] == 1


def test_budget_exhaustion_retains_draft_but_never_publishes_success(tmp_path, monkeypatch):
    corpus, template, image = material(tmp_path)
    monkeypatch.setattr("doctransai.conversion.read_parsed_document_bundle", lambda _: corpus)
    result = convert(TaskSpec(corpus.bundle_path, template, tmp_path / "run", "Fill"),
                     config=RuntimeConfig(max_steps=1,review_enabled=False), client=FakeClient(), previewer=FakePreview(image))
    assert result.execution_status == "failed" and result.output_path is None
    report = json.loads(Path(result.report_path).read_text())
    assert Path(report["draft_path"]).exists()
    assert not (tmp_path / "run/result.docx").exists()


def test_transient_model_retry_does_not_replay_successful_edit(tmp_path, monkeypatch):
    corpus, template, image = material(tmp_path)
    monkeypatch.setattr("doctransai.conversion.read_parsed_document_bundle", lambda _: corpus)
    monkeypatch.setattr("doctransai.agent.time.sleep", lambda _: None)

    class RetryClient(FakeClient):
        failed = False

        def chat(self, *args, **kwargs):
            if self.index == 1 and not self.failed:
                self.failed = True
                raise httpx.ReadError("temporary")
            return super().chat(*args, **kwargs)

    result = convert(TaskSpec(corpus.bundle_path, template, tmp_path / "retry-run", "Fill"),
                     config=RuntimeConfig(review_enabled=False),
                     client=RetryClient(), previewer=FakePreview(image))
    report = json.loads(Path(result.report_path).read_text())
    assert result.execution_status == "completed"
    assert report["document_version"] == 1
    assert report["model_request_count"] == 4


def test_template_preview_does_not_satisfy_output_review(tmp_path):
    corpus, template, image = material(tmp_path)
    view = SourceView(corpus)
    editor = DocxEditor(template, tmp_path / "working.docx", view.asset_path)
    runtime = ToolRuntime(view, editor, FakePreview(image))
    runtime.round = 1
    runtime.invoke("preview_document", {"target": "template"})
    runtime.round = 2
    with pytest.raises(ValueError, match="current output"):
        runtime.invoke("finish", dict(summary="done", checks=["seen"], missing=[], issues=[]))


def test_navigation_keeps_native_anchors_and_search_can_reach_later_hits(tmp_path):
    corpus, _, _ = material(tmp_path)
    view = SourceView(corpus)
    view.objects['ev1']['locator']['page_anchor'] = 'page_1'
    assert view.read(refs=['ev1'])['objects'][0]['navigation']['scanned_page_asset_refs'] == ['img1']
    for i in range(25):
        view.objects[f'hit{i}'] = dict(text='unique needle')
    first = view.read(query='unique needle')
    second = view.read(query='unique needle', offset=first['next_offset'])
    assert len(first['matches']) == 20 and len(second['matches']) == 5
    assert second['next_offset'] is None


def test_runtime_accepts_ranges_lists_rows_and_requires_evidence(tmp_path):
    corpus, _, image = material(tmp_path)
    doc = Document(); doc.add_paragraph('____ and ____')
    template = tmp_path / 'target.docx'; doc.save(template)
    e = DocxEditor(template, tmp_path / 'out.docx', lambda _: image)
    rt = ToolRuntime(SourceView(corpus), e, FakePreview(image))
    edit = dict(op='replace_range', target='n0', start=9, end=13, expected='____', text='fact', source_refs=[], reason='fill')
    with pytest.raises(ValueError, match='require source_refs'):
        rt.invoke('edit_document', dict(batch_id='a', version=0, edits=[edit]))
    edit['source_refs'] = ['ev1']
    rt.invoke('edit_document', dict(batch_id='a', version=0, edits=[edit]))
    assert e.doc.paragraphs[0].text == '____ and fact'
    rt.invoke('edit_document', dict(batch_id='b', version=1, edits=[dict(op='set_list', target='n0',
        items=['one', 'two'], style='decimal', source_refs=['ev1'], reason='organize')]))
    assert [p.text for p in Document(e.output).paragraphs] == ['one', 'two']


def test_edit_delta_does_not_resend_unchanged_nodes_and_detail_is_on_demand(tmp_path):
    corpus, _, image = material(tmp_path)
    doc = Document(); doc.add_paragraph('first'); doc.add_paragraph('fixed')
    template = tmp_path / 'template.docx'; doc.save(template)
    e = DocxEditor(template, tmp_path / 'out.docx', lambda _: image)
    rt = ToolRuntime(SourceView(corpus), e, FakePreview(image))
    summary, _ = rt.invoke('inspect_document', {})
    assert 'runs' not in summary['nodes'][0]
    detail, _ = rt.invoke('inspect_document', dict(node_ids=['n0']))
    assert detail['nodes'][0]['runs'][0]['start'] == 0
    delta, _ = rt.invoke('edit_document', dict(batch_id='a', version=0, edits=[dict(
        op='set_text', target='n0', text='new', source_refs=['ev1'], reason='fill')]))
    assert [n['id'] for n in delta['nodes']] == ['n0']
    assert delta['removed_node_ids'] == [] and 'n1' in e.nodes


def test_xlsx_conversion_selects_workbook_tools_and_output_extension(tmp_path, monkeypatch):
    from openpyxl import Workbook, load_workbook
    corpus, _, image = material(tmp_path)
    monkeypatch.setattr('doctransai.conversion.read_parsed_document_bundle', lambda _: corpus)
    template=tmp_path/'template.xlsx';w=Workbook();w.active['A1']='sample';w.save(template)
    class Client:
        index=0
        def chat(self,messages,tools,**kwargs):
            assert 'set_cells' in json.dumps(tools) and 'repeat_paragraphs' not in json.dumps(tools)
            calls=[('edit_document',dict(batch_id='a',version=0,edits=[dict(op='set_cells',target='s0!A1',values=[['Example project']],source_refs=['ev1'],reason='native fact')])),
                   ('preview_document',{}),('finish',dict(summary='done',checks=['seen'],missing=[],issues=[]))]
            name,args=calls[self.index];self.index+=1
            return dict(message=dict(role='assistant',content='',tool_calls=[dict(id=str(self.index),type='function',function=dict(name=name,arguments=json.dumps(args)))]),usage=None)
    result=convert(TaskSpec(corpus.bundle_path,template,tmp_path/'run','Fill'),config=RuntimeConfig(review_enabled=False),client=Client(),previewer=FakePreview(image))
    assert result.execution_status=='completed',result.issues
    assert result.output_path.endswith('.xlsx') and load_workbook(result.output_path).active['A1'].value=='Example project'
    assert load_workbook(template).active['A1'].value=='sample'


def test_selected_source_ref_does_not_implicitly_clear_existing_xlsx_content(tmp_path, monkeypatch):
    from openpyxl import Workbook, load_workbook
    corpus, _, image = material(tmp_path)
    monkeypatch.setattr('doctransai.conversion.read_parsed_document_bundle', lambda _: corpus)
    template = tmp_path / 'filled-template.xlsx'
    workbook = Workbook()
    workbook.active['A1'] = 'header'
    workbook.active['A2'] = 'existing content'
    workbook.save(template)

    class Client:
        index = 0

        def chat(self, messages, tools, **kwargs):
            calls = [
                ('edit_document', dict(batch_id='retain', version=0, edits=[dict(
                    op='set_cells', target='s0!A3', values=[['new fact']], source_refs=['ev1'], reason='explicit fill')])),
                ('finish', dict(summary='done')),
            ]
            name, args = calls[self.index]
            self.index += 1
            return dict(message=dict(role='assistant', content='', tool_calls=[dict(
                id=str(self.index), type='function', function=dict(name=name, arguments=json.dumps(args)))]),
                        usage=None, model='test-client')

    result = convert(TaskSpec(corpus.bundle_path, template, tmp_path / 'retain-run', 'Fill selected fact'),
                     config=RuntimeConfig(review_enabled=False), client=Client(), previewer=FakePreview(image))
    assert result.execution_status == 'completed', result.issues
    output = load_workbook(result.output_path).active
    assert output['A2'].value == 'existing content' and output['A3'].value == 'new fact'


def test_post_write_verification_error_never_reports_completed(tmp_path, monkeypatch):
    corpus, template, image = material(tmp_path)
    monkeypatch.setattr('doctransai.conversion.read_parsed_document_bundle', lambda _: corpus)
    original = Path.read_bytes
    def fail_final_check(path):
        if path.name == 'result.docx':
            raise OSError('simulated final verification failure')
        return original(path)
    monkeypatch.setattr(Path, 'read_bytes', fail_final_check)
    task = TaskSpec(corpus.bundle_path, template, tmp_path/'run', 'Fill')
    result = convert(task, client=FakeClient(), previewer=FakePreview(image))
    assert result.execution_status == 'failed' and result.output_path is None
    report = json.loads(Path(result.report_path).read_text())
    assert Path(report['draft_path']).exists()
    assert not (task.output_dir/'result.docx').exists()
    assert Document(report['draft_path']).paragraphs[0].text == 'Example project'


def test_formal_visual_refs_are_prepared_and_navigable(tmp_path):
    corpus, _, _ = material(tmp_path)
    corpus.evidence_units[0]['visual_asset_refs'] = ['img1']
    view = SourceView(corpus)
    assert view.initial_asset_refs(['ev1']) == ['img1']
    assert view.navigation('ev1')['asset_refs'] == ['img1']
