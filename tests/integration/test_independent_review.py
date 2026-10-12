import hashlib
from copy import deepcopy
import json
from pathlib import Path
import time

from docx import Document
from PIL import Image
import pytest

from doctransai.config import RuntimeConfig
from doctransai.conversion import review_document
from doctransai.documents.docx import DocxEditor
from doctransai.review import CATEGORIES, ReviewRuntime, RepairRuntime, run_review, review_and_repair
from doctransai.types import TaskSpec
from doctransai.data_layer.models import SourceCorpus, ComponentFile
from doctransai.data_layer.view import SourceView


def fixture(tmp_path):
    template=tmp_path/'template.docx';d=Document();d.add_paragraph('placeholder');d.add_paragraph('fixed');d.save(template)
    image=tmp_path/'page.png';Image.new('RGB',(20,20),'white').save(image)
    doc=dict(document_id='d',source_path=str(template))
    evidence=[dict(document_id='d',evidence_id='ev',text='supported fact')]
    corpus=SourceCorpus(tmp_path/'bundle.json','data_asset_pack_v0',{}, {},[doc],[],
        {'evidence_units':ComponentFile('evidence_units',tmp_path/'evidence.jsonl',1,evidence)}, {},{})
    source=SourceView(corpus)
    editor=DocxEditor(template,tmp_path/'working.docx',lambda _:image)
    class Preview:
        def render(self,path):return [image]
    initial=dict(instruction='Fill supported facts',primary_source_refs=[],source=source.initial(10000),document=editor.inspect(),
                 editor_self_check='UNTRUSTED_EDITOR_PASS')
    return source,editor,Preview(),initial


def issue(kind='factual',repairable=True):
    return dict(id='i1',kind=kind,targets=['n0'],evidence_refs=['ev'],explanation='Text is unsupported',suggested_action='Use supported fact',repairable=repairable)


def finding(issues=()):
    return dict(summary='Independent check',coverage={k:'pass' for k in CATEGORIES},issues=list(issues),grounding=[dict(targets=['n0'],output_claim='supported fact',target_role='fact',source_claim='supported fact',source_role='fact',source_refs=['ev'],judgment='supported')],notes=[])


class Client:
    def __init__(self,calls):self.calls=iter(calls);self.contexts=[]
    def chat(self,messages,tools,**kwargs):
        self.contexts.append((deepcopy(messages),deepcopy(tools)))
        name,args=next(self.calls)
        return dict(message=dict(role='assistant',content='',tool_calls=[dict(id=str(len(self.contexts)),type='function',function=dict(name=name,arguments=json.dumps(args)))]),usage=None)


def test_review_has_no_write_capability_and_context_excludes_editor_assurance(tmp_path):
    source,e,p,initial=fixture(tmp_path);before=e.output.read_bytes()
    runtime=ReviewRuntime(source,e,p)
    names={t['function']['name'] for t in runtime.tools}
    assert 'edit_document' not in names and 'finish' not in names
    with pytest.raises(ValueError,match='read-only'):runtime.invoke('edit_document',{})
    c=Client([('submit_review',finding())])
    result=run_review(source,e,p,initial,c,RuntimeConfig(),tmp_path/'review.jsonl',time.monotonic()+10)
    assert result['status']=='passed' and e.output.read_bytes()==before
    assert 'UNTRUSTED_EDITOR_PASS' not in json.dumps(c.contexts[0][0])
    assert len(c.contexts[0][0])==2


def test_issue_and_uncertain_coverage_cannot_pass(tmp_path):
    source,e,p,initial=fixture(tmp_path)
    report=finding();report['coverage']['factual_support']='uncertain'
    c=Client([('submit_review',report)])
    r=run_review(source,e,p,initial,c,RuntimeConfig(),tmp_path/'review.jsonl',time.monotonic()+10)
    assert r['status']=='needs_review'
    c=Client([('submit_review',finding([issue()]))])
    assert run_review(source,e,p,initial,c,RuntimeConfig(),tmp_path/'r2.jsonl',time.monotonic()+10)['status']=='needs_review'


def test_repair_permissions_reject_other_nodes_and_layout_content_rewrites(tmp_path):
    source,e,p,_=fixture(tmp_path);runtime=RepairRuntime(source,e,p,[issue('layout')]);before=e.output.read_bytes()
    for edit in [dict(op='set_text',target='n0',text='bad',source_refs=['ev'],reason='bad'),
                 dict(op='format_paragraph',target='n1',format={'font_size_pt':10},source_refs=[],reason='outside')]:
        with pytest.raises(ValueError,match='scope'):
            runtime.invoke('edit_document',dict(batch_id='bad',version=0,edits=[edit]))
    assert e.output.read_bytes()==before
    runtime.invoke('edit_document',dict(batch_id='ok',version=0,edits=[dict(op='format_paragraph',target='n0',format={'font_size_pt':10},source_refs=[],reason='layout')]))
    assert e.doc.paragraphs[0].text=='placeholder'


def repair_calls(final):
    return [('submit_review',finding([issue()])),
        ('edit_document',dict(batch_id='repair',version=0,edits=[dict(op='set_text',target='n0',text='supported fact',source_refs=['ev'],reason='review finding')])),
        ('preview_document',{}),('finish',dict(summary='repaired',checks=['seen'],issues=[],missing=[])),
        ('submit_review',final)]


def test_repair_is_followed_by_fresh_independent_review_and_checkpoint(tmp_path):
    source,e,p,initial=fixture(tmp_path);c=Client(repair_calls(finding()))
    result=review_and_repair(source,e,p,initial,c,RuntimeConfig(),tmp_path,time.monotonic()+10)
    assert result['status']=='passed'
    assert len(result['rounds'])==3 and e.doc.paragraphs[0].text=='supported fact'
    assert Document(tmp_path/'revisions/before-repair-1.docx').paragraphs[0].text=='placeholder'
    last=json.dumps(c.contexts[-1][0]);assert 'review_findings' not in last and 'UNTRUSTED_EDITOR_PASS' not in last
    assert len(c.contexts[-1][0])==2


def test_repair_limit_stops_with_unresolved_findings(tmp_path):
    source,e,p,initial=fixture(tmp_path);c=Client(repair_calls(finding([issue()])))
    result=review_and_repair(source,e,p,initial,c,RuntimeConfig(max_repair_rounds=1),tmp_path,time.monotonic()+10)
    assert result['status']=='needs_review' and len(c.contexts)==5


def test_ambiguity_never_triggers_automatic_repair(tmp_path):
    source,e,p,initial=fixture(tmp_path);c=Client([('submit_review',finding([issue('ambiguity',False)]))])
    result=review_and_repair(source,e,p,initial,c,RuntimeConfig(),tmp_path,time.monotonic()+10)
    assert result['status']=='needs_review' and len(c.contexts)==1


def test_missing_page_blocks_review_even_if_model_claims_complete(tmp_path):
    source,e,p,initial=fixture(tmp_path);image=p.render(e.output)[0]
    p.render=lambda path:[image]*5
    c=Client([('submit_review',finding()),('preview_document',dict(start_page=5)),('submit_review',finding())])
    r=run_review(source,e,p,initial,c,RuntimeConfig(),tmp_path/'review.jsonl',time.monotonic()+10)
    assert r['status']=='passed' and r['pages_checked']==[1,2,3,4,5]
    events=[json.loads(l) for l in (tmp_path/'review.jsonl').read_text().splitlines()]
    assert any(x.get('name')=='submit_review' and x['success'] is False for x in events)


def test_review_timeout_publishes_only_candidate_and_preserves_inputs(tmp_path,monkeypatch):
    source,e,p,initial=fixture(tmp_path)
    monkeypatch.setattr('doctransai.conversion.read_parsed_document_bundle',lambda _:source.corpus)
    class TimeoutClient:
        def chat(self,*a,**k):raise TimeoutError('provider timeout')
    before=e.output.read_bytes()
    result=review_document(TaskSpec(source.corpus.bundle_path,e.template,tmp_path/'run','Check'),e.output,client=TimeoutClient(),previewer=p)
    assert result.execution_status=='completed' and result.review_status=='incomplete'
    assert Path(result.output_path).name=='candidate.docx' and not (tmp_path/'run/result.docx').exists()
    assert e.output.read_bytes()==before


def test_review_only_preserves_original_template_context(tmp_path,monkeypatch):
    source,e,p,initial=fixture(tmp_path)
    e.apply(batch_id='original-edit',version=0,edits=[dict(op='set_text',target='n0',text='supported fact')])
    monkeypatch.setattr('doctransai.conversion.read_parsed_document_bundle',lambda _:source.corpus)
    client=Client([('submit_review',finding())])
    result=review_document(TaskSpec(source.corpus.bundle_path,e.template,tmp_path/'run','Check'),e.output,client=client,previewer=p)
    assert result.review_status=='model_passed' and Path(result.output_path).name=='candidate.docx'
    text=client.contexts[0][0][1]['content'][0]['text']
    context=json.loads(text)
    assert context['original_template']['nodes'][0]['text']=='placeholder'
    assert context['current_document']['nodes'][0]['text']=='supported fact'


def test_xlsx_repair_is_bounded_to_ranges_and_cannot_hide_other_content(tmp_path):
    from openpyxl import Workbook
    from doctransai.documents.xlsx import XlsxEditor
    source,_,p,_=fixture(tmp_path)
    path=tmp_path/'template.xlsx';w=Workbook();w.active['A1']='2030-01-02';w.active['B1']='fixed';w.save(path)
    e=XlsxEditor(path,tmp_path/'working.xlsx',lambda _:None)
    i=issue('data_type');i['targets']=['s0!A1']
    runtime=RepairRuntime(source,e,p,[i])
    with pytest.raises(ValueError,match='scope'):
        runtime.invoke('edit_document',dict(batch_id='bad',version=0,edits=[dict(op='set_cells',target='s0!A1:B1',values=[['x','y']],source_refs=['ev'],reason='bad')]))
    runtime.invoke('edit_document',dict(batch_id='ok',version=0,edits=[dict(op='set_cells',target='s0!A1',values=[[{'date':'2030-01-02'}]],source_refs=['ev'],reason='typed date')]))
    assert e.book.active['A1'].data_type=='d' and e.book.active['B1'].value=='fixed'
    i=issue('layout');i['targets']=['s0!A1']
    runtime=RepairRuntime(source,e,p,[i])
    with pytest.raises(ValueError,match='hide'):
        runtime.invoke('edit_document',dict(batch_id='crop',version=1,edits=[dict(op='set_print_area',target='s0!A1',source_refs=[],reason='bad crop')]))


def test_failed_second_review_never_inherits_first_repair_success(tmp_path):
    source,e,p,initial=fixture(tmp_path)
    class C(Client):
        def chat(self,*a,**k):
            if len(self.contexts)==4:raise TimeoutError('verification timeout')
            return super().chat(*a,**k)
    c=C(repair_calls(finding()))
    r=review_and_repair(source,e,p,initial,c,RuntimeConfig(),tmp_path,time.monotonic()+10)
    assert e.doc.paragraphs[0].text=='supported fact'
    assert r['status']=='incomplete' and r['rounds'][-1]['phase']=='review-1'


def test_cli_requires_attention_exit_code_for_unresolved_review(tmp_path,monkeypatch):
    from doctransai.cli import main
    from doctransai.types import ConversionResult
    monkeypatch.setattr('doctransai.conversion.convert',lambda *a,**k:ConversionResult('completed','needs_review','candidate.docx','review.json'))
    assert main(['review','--document','candidate.docx','--bundle','bundle.json','--template','template.docx','--output-dir',str(tmp_path/'run'),'--instruction','check'])==3


def test_repair_rejects_disk_change_after_findings(tmp_path):
    source,e,p,_=fixture(tmp_path);runtime=RepairRuntime(source,e,p,[issue()])
    d=Document(e.output);d.paragraphs[0].text='external change';d.save(e.output)
    before=e.output.read_bytes()
    with pytest.raises(ValueError,match='changed outside'):
        runtime.invoke('edit_document',dict(batch_id='old',version=0,edits=[dict(op='set_text',target='n0',text='supported fact',source_refs=['ev'],reason='old review')]))
    assert e.output.read_bytes()==before


def test_unsupported_grounding_cannot_be_hidden_by_empty_issues(tmp_path):
    source,e,p,initial=fixture(tmp_path);bad=finding();bad['grounding'][0]['judgment']='unsupported'
    client=Client([('submit_review',bad),('submit_review',finding([issue()]))])
    r=run_review(source,e,p,initial,client,RuntimeConfig(),tmp_path/'review.jsonl',time.monotonic()+10)
    assert r['status']=='needs_review' and len(client.contexts)==2


def test_acceptance_is_bound_to_the_exact_document_hash(tmp_path,monkeypatch):
    source,e,p,_=fixture(tmp_path)
    monkeypatch.setattr('doctransai.conversion.read_parsed_document_bundle',lambda _:source.corpus)
    monkeypatch.setattr('doctransai.conversion.review_and_repair',lambda *a,**k:dict(status='passed',rounds=[dict(document_version=0,document_sha256='wrong')],issues=[]))
    result=review_document(TaskSpec(source.corpus.bundle_path,e.template,tmp_path/'run','Check'),e.output,client=Client([]),previewer=p)
    assert result.review_status=='incomplete' and Path(result.output_path).name=='candidate.docx'


@pytest.mark.parametrize('kwargs',[dict(max_repair_rounds=3),dict(max_seconds=float('inf')),dict(max_seconds=float('nan'))])
def test_review_budgets_cannot_be_unbounded(kwargs):
    with pytest.raises(ValueError):RuntimeConfig(**kwargs)


def test_ordinary_runtime_defaults_to_no_independent_review():
    config = RuntimeConfig()
    assert config.review_enabled is False
    assert config.model == 'kimi-k3'


def test_adjacent_reviewed_cells_may_be_repaired_as_one_exact_union(tmp_path):
    from openpyxl import Workbook
    from doctransai.documents.xlsx import XlsxEditor
    source,_,p,_=fixture(tmp_path)
    path=tmp_path/'union.xlsx';w=Workbook();w.active.append(['2030-01-02','2030-01-03','fixed']);w.save(path)
    e=XlsxEditor(path,tmp_path/'working.xlsx',lambda _:None)
    i=issue('data_type');i['targets']=['s0!A1','s0!B1']
    runtime=RepairRuntime(source,e,p,[i])
    runtime.invoke('edit_document',dict(batch_id='union',version=0,edits=[dict(op='set_cells',target='s0!A1:B1',values=[[{'date':'2030-01-02'},{'date':'2030-01-03'}]],source_refs=['ev'],reason='date types')]))
    assert e.book.active['A1'].data_type==e.book.active['B1'].data_type=='d'
    assert e.book.active['C1'].value=='fixed'


def test_explicit_conversion_review_runs_after_editor_finish(tmp_path,monkeypatch):
    from doctransai.conversion import convert
    source,e,p,_=fixture(tmp_path)
    monkeypatch.setattr('doctransai.conversion.read_parsed_document_bundle',lambda _:source.corpus)
    client=Client([
        ('edit_document',dict(batch_id='initial',version=0,edits=[dict(op='set_text',target='n0',text='supported fact',source_refs=['ev'],reason='source')])),
        ('preview_document',{}),
        ('finish',dict(summary='EDITOR_CERTAINLY_APPROVED',checks=['seen'],issues=[],missing=[])),
        ('submit_review',finding()),
    ])
    result=convert(TaskSpec(source.corpus.bundle_path,e.template,tmp_path/'convert-run','Fill'),
                   config=RuntimeConfig(review_enabled=True), client=client,previewer=p)
    assert result.review_status=='model_passed' and Path(result.output_path).name=='candidate.docx'
    assert len(client.contexts)==4 and 'EDITOR_CERTAINLY_APPROVED' not in json.dumps(client.contexts[-1][0])
