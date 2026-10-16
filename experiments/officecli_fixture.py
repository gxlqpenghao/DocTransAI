"""Synthetic fixture only; no business documents or credentials."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from PIL import Image, ImageDraw


def create(root):
    root=Path(root);root.mkdir(parents=True,exist_ok=True)
    image=Image.new('RGB',(640,180),'#edf5ee');ImageDraw.Draw(image).rectangle((15,15,625,165),outline='#235b45',width=6);image.save(root/'image.png')
    d=Document();style=d.styles['Normal'];style.font.name='Microsoft YaHei';style.font.size=Pt(11);style.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    d.add_heading('模板填写工具验证',0)
    d.sections[0].header.paragraphs[0].text='项目：{{project}}'
    d.add_paragraph('固定说明：保留原版式和图片，仅填写指定位置。')
    p=d.add_paragraph();p.add_run('项目名称：');p.add_run('{{pro').bold=True;p.add_run('ject}}').bold=True
    table=d.add_table(rows=3,cols=2);table.style='Table Grid'
    for row,(key,val) in zip(table.rows,[('检查日期','{{date}}'),('整改期限','{{deadline}}'),('检查结果','{{result}}')]):row.cells[0].text=key;row.cells[1].text=val
    d.add_picture(str(root/'image.png'),width=Inches(4));d.add_paragraph('图片说明：验证图像保持不变。')
    d.sections[0].footer.paragraphs[0].text='本地工具对照测试'
    d.save(root/'fixture.docx')
    w=Workbook();s=w.active;s.title='检查记录';s.merge_cells('A1:D1');s['A1']='{{project}}';s['A1'].font=Font(name='Microsoft YaHei',bold=True,size=16)
    for row in [('项目','日期','数量','小计'),('{{project}}','{{date}}',3,'=C3*2')]:s.append(row)
    for col in 'ABCD':s.column_dimensions[col].width=24
    s.print_area='A1:D8';w.save(root/'fixture.xlsx')

if __name__=='__main__':
    import sys
    create(sys.argv[1])
