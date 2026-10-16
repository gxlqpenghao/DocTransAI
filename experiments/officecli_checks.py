"""Check the synthetic probe outputs; records limitations rather than hiding them."""
import argparse,hashlib,json,os,shutil,subprocess
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from doctransai.documents.docx import DocxEditor

p=argparse.ArgumentParser();p.add_argument('--binary',required=True);p.add_argument('--root',type=Path,required=True);a=p.parse_args();root=a.root.resolve()
env=dict(os.environ,OFFICECLI_SKIP_UPDATE='1',OFFICECLI_RESIDENT_FLUSH='each')
original=root/'fixture.docx';copy=root/'atomic-check.docx';shutil.copy2(original,copy);before=hashlib.sha256(copy.read_bytes()).hexdigest()
ops=[{'command':'set','path':'/body/p[2]','props':{'text':'MUST ROLLBACK'}},{'command':'set','path':'/body/p[9999]','props':{'text':'invalid'}}]
r=subprocess.run([a.binary,'batch',str(copy),'--commands',json.dumps(ops),'--json'],env=env,capture_output=True,text=True,timeout=30)
subprocess.run([a.binary,'close',str(copy)],env=env,capture_output=True,timeout=30)
merged=Document(root/'merged.docx');w=load_workbook(root/'merged.xlsx');cache=load_workbook(root/'merged.xlsx',data_only=True)
ed=DocxEditor(original,root/'current-editor-check.docx',lambda _:root/'image.png')
node=next(k for k,v in ed.nodes.items() if getattr(v,'text',None)=='项目名称：{{project}}')
ed.apply(batch_id='probe',version=0,edits=[{'op':'replace_text','target':node,'old':'{{project}}','text':'合成测试项目'}]);ed.save()
checks={'atomic_failure_exit':r.returncode,'atomic_rollback_preserved_bytes':hashlib.sha256(copy.read_bytes()).hexdigest()==before,
 'officecli_split_placeholder_remaining':any('{{project}}' in x.text for x in merged.paragraphs),
 'officecli_header_filled':'合成测试项目' in merged.sections[0].header.paragraphs[0].text,
 'current_editor_split_placeholder_filled':any(x.text=='项目名称：合成测试项目' for x in Document(root/'current-editor-check.docx').paragraphs),
 'xlsx_value':w.active['A3'].value,'xlsx_formula':w.active['D3'].value,'xlsx_formula_cache':cache.active['D3'].value}
(root/'synthetic-checks.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2));print(json.dumps(checks,ensure_ascii=False,indent=2))
